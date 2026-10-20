from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "lib"))
import word_core as WORD


class BibliographyLanguageTests(unittest.TestCase):
    def test_reference_visibility_options_are_independent(self):
        pandoc = WORD.find_pandoc(None, ROOT)
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            bibliography = directory / "references.bib"
            bibliography.write_text(
                "@article{sample, author={Alpha, A}, title={A verifiable title}, "
                "journal={Journal}, year={2020}, volume={1}, number={2}, "
                "pages={1-4}, doi={10.1234/test}, url={https://example.org}, "
                "urldate={2025-01-02}, langid={english}}",
                encoding="utf-8",
            )
            configurations = (
                ("default", True, True, True, True),
                ("no-doi", False, True, False, True),
                ("no-url", True, False, True, False),
            )
            for name, include_doi, include_url, expect_doi, expect_url in configurations:
                with self.subTest(name=name):
                    job_directory = directory / name
                    job_directory.mkdir()
                    prepared = WORD.prepare_bibliography(
                        bibliography,
                        pandoc=pandoc,
                        cwd=directory,
                        temp_dir=job_directory,
                        include_doi=include_doi,
                        include_url=include_url,
                    )
                    output = job_directory / "references.docx"
                    markdown = job_directory / "input.md"
                    markdown.write_text("[@sample]", encoding="utf-8")
                    subprocess.run(
                        [
                            str(pandoc), "--citeproc", "--from=markdown", "--to=docx",
                            f"--bibliography={prepared}",
                            f"--csl={ROOT / 'word/china-national-standard-gb-t-7714-2015-numeric.csl'}",
                            f"--output={output}", str(markdown),
                        ],
                        cwd=directory,
                        check=True,
                    )
                    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
                    self.assertEqual("10.1234/test" in text, expect_doi)
                    self.assertEqual("https://example.org" in text, expect_url)
                    self.assertEqual("2025-01-02" in text, expect_url)
                    source_items = json.loads(
                        WORD.bibliography_source_items_path(prepared).read_text(encoding="utf-8")
                    )
                    self.assertEqual(source_items[0]["DOI"], "10.1234/test")
                    self.assertEqual(source_items[0]["URL"], "https://example.org")

    def test_document_class_options_include_reference_visibility(self):
        options = WORD.parse_document_class_options(
            r"\documentclass[proposal,humanities,no-doi,no-url]{yibinthesis}"
        )
        self.assertEqual(options, {"proposal", "humanities", "no-doi", "no-url"})

    def test_citation_clusters_use_gb_style_ranges_without_spaces(self):
        numbers = {f"k{number}": number for number in range(1, 9)}
        self.assertEqual(
            WORD._citation_cluster_parts(["k1", "k2", "k3", "k5", "k7", "k8"], numbers),
            [("k1", "k3"), ("k5", None), ("k7", None), ("k8", None)],
        )
        self.assertEqual(
            WORD._citation_cluster_parts(["k1", "k2"], numbers),
            [("k1", None), ("k2", None)],
        )

    def test_real_citeproc_edition_language(self):
        pandoc = WORD.find_pandoc(None, ROOT)
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            bibliography = directory / "references.bib"
            bibliography.write_text(
                "@book{en, author={Krippendorff, Klaus}, title={Content Analysis}, edition={4}, year={2019}}\n"
                "@book{zh, author={张甲}, title={中文书名}, edition={4}, year={2019}, langid={chinese}}",
                encoding="utf-8",
            )
            prepared = WORD.prepare_bibliography(bibliography, pandoc=pandoc, cwd=directory, temp_dir=directory)
            items = json.loads(prepared.with_suffix(".json").read_text(encoding="utf-8"))
            output = directory / "references.docx"
            subprocess.run([
                str(pandoc), "--citeproc", "--from=markdown", "--to=docx",
                f"--bibliography={prepared}",
                f"--csl={ROOT / 'word/china-national-standard-gb-t-7714-2015-numeric.csl'}",
                f"--output={output}",
            ], input="[@en; @zh]", text=True, encoding="utf-8", check=True)
            document = Document(output)
            WORD._localize_bibliography_terms(document, WORD.CitationRegistry(order=["en", "zh"]), items)
            entries = [p.text for p in document.paragraphs if p.style.name == "Bibliography"]
            self.assertIn("[M]. 4th ed.", entries[0])
            self.assertIn("[M]. 4 版.", entries[1])
            self.assertEqual(str(items[0]["edition"]), "4")

    def test_edition_ordinals_split_runs_and_idempotence(self):
        for edition, ordinal in ((1, "1st"), (2, "2nd"), (3, "3rd"), (4, "4th"), (11, "11th"), (12, "12th"), (13, "13th"), (21, "21st")):
            with self.subTest(edition=edition):
                document = Document()
                document.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
                paragraph = document.add_paragraph(style="Bibliography")
                paragraph.add_run("[1] ALPHA A, BETA B, GAMMA G, 等. Title[M]. ")
                paragraph.add_run(str(edition)).bold = True
                paragraph.add_run(" 版. Publisher.")
                items = [{"id": "one", "title": "Title", "edition": edition}]
                citations = WORD.CitationRegistry(order=["one"])
                WORD._localize_bibliography_terms(document, citations, items)
                expected = f"[1] ALPHA A, BETA B, GAMMA G, et al. Title[M]. {ordinal} ed. Publisher."
                self.assertEqual(paragraph.text, expected)
                self.assertTrue(paragraph.runs[1].bold)
                WORD._localize_bibliography_terms(document, citations, items)
                self.assertEqual(paragraph.text, expected)

    def test_real_citeproc_author_threshold_and_languages(self):
        pandoc = WORD.find_pandoc(None, ROOT)
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            bibliography = directory / "references.bib"
            bibliography.write_text("\n".join(
                f"@article{{{language}{count}, author={{{' and '.join(names[:count])}}}, "
                f"title={{{title}}}, year={{2026}}, langid={{{language}}}}}"
                for language, title, names in (
                    ("english", "English title", ["Alpha, Adam", "Beta, Ben", "Gamma, Gail", "Delta, Dan", "Epsilon, Eve"]),
                    ("chinese", "中文题名", ["张甲", "李乙", "王丙", "赵丁", "钱戊"]),
                )
                for count in (3, 4, 5)
            ), encoding="utf-8")
            prepared = WORD.prepare_bibliography(
                bibliography, pandoc=pandoc, cwd=directory, temp_dir=directory,
            )
            items = json.loads(prepared.with_suffix(".json").read_text(encoding="utf-8"))
            keys = [f"{language}{count}" for language in ("english", "chinese") for count in (3, 4, 5)]
            output = directory / "references.docx"
            subprocess.run([
                str(pandoc), "--citeproc", "--from=markdown", "--to=docx",
                f"--bibliography={prepared}",
                f"--csl={ROOT / 'word/china-national-standard-gb-t-7714-2015-numeric.csl'}",
                f"--output={output}",
            ], input="; ".join(f"[@{key}]" for key in keys), text=True, encoding="utf-8", check=True)
            document = Document(output)
            WORD._localize_bibliography_terms(document, WORD.CitationRegistry(order=keys), items)
            entries = [p.text for p in document.paragraphs if p.style.name == "Bibliography"]
            self.assertEqual(len(entries), 6)
            for index, entry in enumerate(entries):
                if index % 3 == 0:
                    self.assertNotIn(", 等.", entry)
                    self.assertNotIn(", et al.", entry)
                elif index < 3:
                    self.assertIn("ALPHA A, BETA B, GAMMA G, et al.", entry)
                    self.assertNotIn("等", entry)
                else:
                    self.assertIn("张甲, 李乙, 王丙, 等.", entry)
            self.assertEqual(len(next(item for item in items if item["id"] == "english5")["author"]), 5)

    def test_only_author_term_changes_and_runs_survive(self):
        document = Document()
        document.styles.add_style(WORD.STYLE_REVIEW_BIBLIOGRAPHY, WD_STYLE_TYPE.PARAGRAPH)
        paragraph = document.add_paragraph(style=WORD.STYLE_REVIEW_BIBLIOGRAPHY)
        paragraph.add_run("[1] ALPHA A, BETA B, GAMMA G, ")
        paragraph.add_run("等").italic = True
        paragraph.add_run(". 中文题名, 等. 保留正文与机构名称等.")
        body = document.add_paragraph("[1] ALPHA A, BETA B, GAMMA G, 等. 正文不改")
        WORD._localize_bibliography_terms(document, WORD.CitationRegistry(order=["one"]), [
            {"id": "one", "title": "中文题名", "language": "en-US"},
        ])
        self.assertIn(", et al. 中文题名, 等.", paragraph.text)
        self.assertTrue(paragraph.runs[1].italic)
        self.assertIn(", 等.", body.text)


if __name__ == "__main__":
    unittest.main()
