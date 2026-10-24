from __future__ import annotations

import importlib.util
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING


ROOT = Path(__file__).resolve().parents[2]


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


BUILD_WORD = _load_module(
    "yibinthesis_build_word_profiles",
    ROOT / "lib" / "build_word.py",
)
BUILD_REFERENCE = _load_module(
    "yibinthesis_build_reference_profiles",
    ROOT / "lib" / "build_reference_docx.py",
)


class DocumentProfileParsingTests(unittest.TestCase):
    def test_repeated_profile_binding_preserves_module_identity(self) -> None:
        first = BUILD_WORD._profiles_module()
        second = BUILD_WORD._profiles_module()
        self.assertIs(first, second)
        self.assertEqual(first.__name__, "word_profiles")
        self.assertEqual(first.__file__, str(ROOT / "lib" / "word_profiles.py"))
        self.assertEqual(BUILD_WORD._bind_oxml_module().__name__, "word_oxml")
        self.assertIsNot(first._postprocess_review_docx, BUILD_WORD._postprocess_review_docx)

    def test_document_type_defaults_and_explicit_options(self) -> None:
        cases = (
            (r"\documentclass{yibinthesis}", "thesis", "humanities", "2024"),
            (
                r"\documentclass[humanities,proposal]{yibinthesis}",
                "proposal",
                "humanities",
                "2024",
            ),
            (
                r"\documentclass[literature-review,science]{yibinthesis}",
                "literature-review",
                "science",
                "2024",
            ),
        )
        for source, document_type, discipline, template_year in cases:
            with self.subTest(source=source):
                profile = BUILD_WORD.resolve_profile(
                    source,
                    {"template-year": template_year},
                )
                self.assertEqual(profile.document_type, document_type)
                self.assertEqual(profile.discipline, discipline)
                self.assertEqual(profile.template_year, template_year)

    def test_document_type_options_are_mutually_exclusive(self) -> None:
        for options in (
            "thesis,proposal",
            "proposal,literature-review",
            "thesis,proposal,literature-review",
        ):
            with self.subTest(options=options):
                with self.assertRaisesRegex(BUILD_WORD.BuildError, "文档类型类选项互斥"):
                    BUILD_WORD.parse_document_type(
                        rf"\documentclass[{options}]{{yibinthesis}}"
                    )

    def test_template_year_is_bound_to_document_type(self) -> None:
        cases = (
            (r"\documentclass{yibinthesis}", "2024"),
            (r"\documentclass[proposal]{yibinthesis}", "2024"),
            (r"\documentclass[literature-review]{yibinthesis}", "2024"),
        )
        for main, expected in cases:
            with self.subTest(main=main):
                self.assertEqual(BUILD_WORD.resolve_profile(main, {}).template_year, expected)
                self.assertEqual(
                    BUILD_WORD.resolve_profile(main, {"template-year": "  "}).template_year,
                    expected,
                )
                self.assertEqual(
                    BUILD_WORD.resolve_profile(main, {"template-year": expected}).template_year,
                    expected,
                )

    def test_mismatched_template_year_is_rejected(self) -> None:
        cases = (
            (r"\documentclass{yibinthesis}", "2022", "thesis", "2024"),
            (r"\documentclass[proposal]{yibinthesis}", "2022", "proposal", "2024"),
            (
                r"\documentclass[literature-review]{yibinthesis}",
                "2022",
                "literature-review",
                "2024",
            ),
        )
        for main, actual, document_type, expected in cases:
            with self.subTest(main=main, actual=actual):
                with self.assertRaisesRegex(
                    BUILD_WORD.BuildError,
                    rf"文档类型 {re.escape(document_type)} 仅支持 template-year={expected}",
                ):
                    BUILD_WORD.resolve_profile(main, {"template-year": actual})


class ProposalFieldContractTests(unittest.TestCase):
    @staticmethod
    def _source(keys: tuple[str, ...]) -> str:
        return "\n".join(
            rf"\yibinproposalfield{{{key}}}{{{key} 的内容 {{含嵌套参数}}}}"
            for key in keys
        )

    def test_seven_fields_are_extracted_in_the_official_order(self) -> None:
        fields = BUILD_WORD.extract_proposal_fields(
            self._source(BUILD_WORD.PROPOSAL_FIELD_ORDER)
        )
        self.assertEqual(
            tuple(key for key, _ in fields),
            BUILD_WORD.PROPOSAL_FIELD_ORDER,
        )
        self.assertEqual(len(fields), 7)
        self.assertIn("{含嵌套参数}", fields[0][1])

    def test_missing_field_is_rejected(self) -> None:
        keys = BUILD_WORD.PROPOSAL_FIELD_ORDER[:-1]
        with self.assertRaisesRegex(
            BUILD_WORD.BuildError,
            "开题报告字段缺失：advisor-opinion",
        ):
            BUILD_WORD.extract_proposal_fields(self._source(keys))

    def test_duplicate_field_is_rejected(self) -> None:
        keys = BUILD_WORD.PROPOSAL_FIELD_ORDER + ("schedule",)
        with self.assertRaisesRegex(
            BUILD_WORD.BuildError,
            "开题报告字段重复：schedule",
        ):
            BUILD_WORD.extract_proposal_fields(self._source(keys))

    def test_out_of_order_fields_are_rejected(self) -> None:
        keys = list(BUILD_WORD.PROPOSAL_FIELD_ORDER)
        keys[1], keys[2] = keys[2], keys[1]
        with self.assertRaisesRegex(
            BUILD_WORD.BuildError,
            "开题报告字段顺序必须为",
        ):
            BUILD_WORD.extract_proposal_fields(self._source(tuple(keys)))

    def test_unnumbered_heading_uses_reusable_proposal_style(self) -> None:
        converted = BUILD_WORD.normalize_proposal_headings(
            "正文\n# 编号标题\n内容\n# 可行性分析 {.unnumbered}\n后文"
        )
        self.assertIn("\n# 编号标题\n", converted)
        self.assertIn(
            f'custom-style="{BUILD_WORD.STYLE_PROPOSAL_UNNUMBERED_HEADING}"',
            converted,
        )
        self.assertNotIn("## 可行性分析", converted)

    def test_figure_numbering_preserves_following_custom_style_block(self) -> None:
        source = (
            "![研究路线](route.png){#fig:route}\n\n"
            f'::: {{custom-style="{BUILD_WORD.STYLE_PROPOSAL_UNNUMBERED_HEADING}"}}\n'
            "可行性分析\n:::"
        )
        converted = BUILD_WORD.apply_float_numbering(
            source,
            "humanities",
            0,
            BUILD_WORD.FloatCounters(),
            BUILD_WORD.LabelRegistry(),
        )
        self.assertIn(
            f'}}\n\n::: {{custom-style="{BUILD_WORD.STYLE_PROPOSAL_UNNUMBERED_HEADING}"}}',
            converted,
        )


class WordProfileStructureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._temporary = tempfile.TemporaryDirectory()
        cls.root = Path(cls._temporary.name)
        cls.reference_docx = cls.root / "reference.docx"
        BUILD_REFERENCE.build_reference_docx(cls.reference_docx)

    @classmethod
    def tearDownClass(cls) -> None:
        cls._temporary.cleanup()

    @staticmethod
    def _metadata() -> dict[str, str]:
        return {
            "template-year": "2024",
            "title": "公共服务数据治理方法研究",
            "author": "张同学",
            "student-id": "202400000",
            "college": "示例学部",
            "affiliation-city": "示例城市",
            "affiliation-postcode": "000000",
            "major": "示例专业",
            "grade": "2024",
            "class-name": "1班",
            "advisor": "李老师",
            "advisor-title": "讲师",
            "date": "2026年7月",
        }

    @staticmethod
    def _field_instructions(element) -> list[str]:
        instructions = [
            node.text or ""
            for node in element.iter(BUILD_WORD.qn("w:instrText"))
        ]
        instructions.extend(
            node.get(BUILD_WORD.qn("w:instr"), "")
            for node in element.iter(BUILD_WORD.qn("w:fldSimple"))
        )
        return instructions

    @staticmethod
    def _page_number_attributes(section) -> dict[str, str] | None:
        node = section._sectPr.find(BUILD_WORD.qn("w:pgNumType"))
        if node is None:
            return None
        return {
            "format": node.get(BUILD_WORD.qn("w:fmt"), ""),
            "start": node.get(BUILD_WORD.qn("w:start"), ""),
        }

    def test_reference_document_defines_reusable_profile_styles(self) -> None:
        document = Document(self.reference_docx)
        expected = {
            BUILD_WORD.STYLE_PROPOSAL_TITLE,
            BUILD_WORD.STYLE_PROPOSAL_SUBTITLE,
            BUILD_WORD.STYLE_PROPOSAL_LABEL,
            BUILD_WORD.STYLE_PROPOSAL_BODY,
            BUILD_WORD.STYLE_PROPOSAL_PROMPT,
            BUILD_WORD.STYLE_PROPOSAL_UNNUMBERED_HEADING,
            BUILD_WORD.STYLE_PROPOSAL_SIGNATURE,
            BUILD_WORD.STYLE_REVIEW_DOCUMENT_TITLE,
            BUILD_WORD.STYLE_REVIEW_THESIS_TITLE,
            BUILD_WORD.STYLE_REVIEW_INFO_LABEL,
            BUILD_WORD.STYLE_REVIEW_INFO_VALUE,
            BUILD_WORD.STYLE_REVIEW_DATE,
            BUILD_WORD.STYLE_REVIEW_COVER_FIELD,
            BUILD_WORD.STYLE_REVIEW_AUTHOR,
            BUILD_WORD.STYLE_REVIEW_AFFILIATION,
            BUILD_WORD.STYLE_REVIEW_ABSTRACT_TITLE,
            BUILD_WORD.STYLE_REVIEW_ENGLISH_ABSTRACT_TITLE,
            BUILD_WORD.STYLE_REVIEW_CHINESE_ABSTRACT,
            BUILD_WORD.STYLE_REVIEW_ENGLISH_ABSTRACT,
            BUILD_WORD.STYLE_REVIEW_KEYWORDS,
            BUILD_WORD.STYLE_REVIEW_ENGLISH_KEYWORDS,
            BUILD_WORD.STYLE_REVIEW_BODY,
            BUILD_WORD.STYLE_REVIEW_HEADING_1,
            BUILD_WORD.STYLE_REVIEW_BIBLIOGRAPHY,
            BUILD_WORD.STYLE_FIGURE_CAPTION,
            BUILD_WORD.STYLE_TABLE_CAPTION,
            BUILD_WORD.STYLE_CITATION,
            BUILD_WORD.STYLE_LIST_BODY,
            BUILD_WORD.STYLE_APPENDIX_SECTION,
        }
        self.assertTrue(expected.issubset({style.name for style in document.styles}))

    def test_review_styles_keep_abstracts_together_and_headings_in_flow(self) -> None:
        document = Document(self.reference_docx)
        for name in (
            BUILD_WORD.STYLE_REVIEW_ABSTRACT_TITLE,
            BUILD_WORD.STYLE_REVIEW_ENGLISH_ABSTRACT_TITLE,
            BUILD_WORD.STYLE_REVIEW_HEADING_1,
        ):
            with self.subTest(style=name):
                self.assertFalse(document.styles[name].paragraph_format.page_break_before)
        self.assertTrue(document.styles[BUILD_WORD.STYLE_REVIEW_CHINESE_ABSTRACT].paragraph_format.keep_together)
        self.assertTrue(document.styles[BUILD_WORD.STYLE_REVIEW_ENGLISH_ABSTRACT].paragraph_format.keep_together)
        self.assertTrue(document.styles[BUILD_WORD.STYLE_REVIEW_KEYWORDS].paragraph_format.keep_with_next)
        self.assertFalse(document.styles[BUILD_WORD.STYLE_REVIEW_ENGLISH_KEYWORDS].paragraph_format.keep_with_next)
        self.assertTrue(document.styles[BUILD_WORD.STYLE_FRONT_TITLE].paragraph_format.page_break_before)
        self.assertEqual(document.styles[BUILD_WORD.STYLE_REVIEW_BODY].font.size.pt, 12)
        self.assertEqual(document.styles[BUILD_WORD.STYLE_REVIEW_BIBLIOGRAPHY].font.size.pt, 10.5)

    def test_review_heading_list_has_its_own_word_identity(self) -> None:
        document = Document(self.reference_docx)
        thesis_id = BUILD_WORD._configure_heading_numbering(document, "humanities")
        review_id = BUILD_WORD._configure_heading_numbering(
            document, "humanities", style_names=(
                BUILD_WORD.STYLE_REVIEW_HEADING_1, BUILD_WORD.STYLE_REVIEW_HEADING_2,
                BUILD_WORD.STYLE_REVIEW_HEADING_3, BUILD_WORD.STYLE_REVIEW_HEADING_4,
            ),
        )
        identities = {
            int(node.get(BUILD_WORD.qn("w:abstractNumId"))): node.find(BUILD_WORD.qn("w:nsid")).get(BUILD_WORD.qn("w:val"))
            for node in document.part.numbering_part.element.findall(BUILD_WORD.qn("w:abstractNum"))
        }
        self.assertNotEqual(identities[thesis_id], identities[review_id])

    def test_reference_caption_directory_matches_first_level_toc(self) -> None:
        document = Document(self.reference_docx)
        style = document.styles[BUILD_REFERENCE.STYLE_TABLE_OF_FIGURES]
        paragraph = style.paragraph_format
        properties = style.element.get_or_add_pPr()
        indent = properties.find(BUILD_REFERENCE.qn("w:ind"))
        tabs = properties.find(BUILD_REFERENCE.qn("w:tabs"))
        tab = tabs.find(BUILD_REFERENCE.qn("w:tab"))
        fonts = style.element.get_or_add_rPr().find(BUILD_REFERENCE.qn("w:rFonts"))

        self.assertEqual(style.font.name, "SimSun")
        self.assertEqual(style.font.size.pt, 12)
        self.assertFalse(style.font.bold)
        self.assertIsNone(style.element.get(BUILD_REFERENCE.qn("w:customStyle")))
        self.assertEqual(
            fonts.get(BUILD_REFERENCE.qn("w:eastAsia")),
            "SimSun",
        )
        self.assertEqual(paragraph.alignment, BUILD_REFERENCE.WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertEqual(paragraph.line_spacing, 1.5)
        self.assertEqual(paragraph.space_before.pt, 0)
        self.assertEqual(paragraph.space_after.pt, 0)
        self.assertEqual(indent.get(BUILD_REFERENCE.qn("w:left")), "0")
        self.assertEqual(indent.get(BUILD_REFERENCE.qn("w:leftChars")), "0")
        self.assertEqual(indent.get(BUILD_REFERENCE.qn("w:firstLine")), "0")
        self.assertEqual(tab.get(BUILD_REFERENCE.qn("w:val")), "right")
        self.assertEqual(tab.get(BUILD_REFERENCE.qn("w:leader")), "dot")
        self.assertEqual(tab.get(BUILD_REFERENCE.qn("w:pos")), "8777")

    def test_existing_pandoc_caption_is_not_duplicated(self) -> None:
        document = Document(self.reference_docx)
        drawing = document.paragraphs[0]
        drawing.text = ""
        drawing.add_run().add_picture(
            str(ROOT / "examples" / "full-featured" / "assets" / "document-pipeline.png")
        )
        for properties in drawing._p.iter(BUILD_WORD.qn("wp:docPr")):
            properties.set("descr", "图1 测试图题")
        document.add_paragraph("图1 测试图题", style=BUILD_WORD.STYLE_FIGURE_CAPTION)

        BUILD_WORD._split_pandoc_captioned_figures(document)

        captions = [paragraph for paragraph in document.paragraphs if paragraph.text == "图1 测试图题"]
        self.assertEqual(len(captions), 1)
        self.assertEqual(captions[0].style.name, BUILD_WORD.STYLE_FIGURE_CAPTION)

    def test_custom_thesis_headings_do_not_replace_builtin_headings(self) -> None:
        document = Document(self.reference_docx)
        custom_names = (
            BUILD_WORD.STYLE_HEADING_1,
            BUILD_WORD.STYLE_HEADING_2,
            BUILD_WORD.STYLE_HEADING_3,
            BUILD_WORD.STYLE_HEADING_4,
        )
        for level, custom_name in enumerate(custom_names, start=1):
            with self.subTest(level=level):
                builtin = document.styles[f"Heading {level}"]
                custom = document.styles[custom_name]
                self.assertEqual(builtin.style_id, f"Heading{level}")
                self.assertEqual(builtin.name, f"Heading {level}")
                self.assertNotEqual(custom.style_id, builtin.style_id)
                self.assertEqual(custom.name, custom_name)

        self.assertTrue(
            document.styles[BUILD_WORD.STYLE_HEADING_1]
            .paragraph_format.page_break_before
        )

    def test_common_cover_optional_styles_follow_school_type_scale(self) -> None:
        document = Document(self.reference_docx)
        version = document.styles["CoverVersion"].font
        date = document.styles["CoverDate"].font

        self.assertEqual(version.size.pt, 14)
        self.assertFalse(version.bold)
        self.assertEqual(
            document.styles["CoverVersion"].paragraph_format.space_after.pt,
            23,
        )
        self.assertEqual(date.size.pt, 15)
        self.assertTrue(date.bold)

    def test_proposal_has_2024_cover_and_form_sections_starting_at_page_one(self) -> None:
        source = Document(self.reference_docx)
        source.paragraphs[0].text = ""
        for key in BUILD_WORD.PROPOSAL_FIELD_ORDER:
            source.add_paragraph(BUILD_WORD.PROPOSAL_FIELD_MARKER_PREFIX + key.upper())
            source.add_paragraph(f"{key} 内容")

        input_docx = self.root / "proposal-input.docx"
        output_docx = self.root / "proposal-output.docx"
        source.save(input_docx)
        long_proposal_title = "基于政策文本分析的低空经济领域标准需求识别与供需匹配研究"
        BUILD_WORD._postprocess_proposal_docx(
            input_docx,
            output_docx,
            metadata={
                **self._metadata(),
                "template-year": "2024",
                "title": long_proposal_title,
            },
            project_root=ROOT,
            main_dir=self.root,
            metadata_dir=self.root,
            discipline="humanities",
            citation_mode="linked",
            labels=BUILD_WORD.LabelRegistry(),
            citations=BUILD_WORD.CitationRegistry(),
            bibliography_items=[],
            table_layouts=[],
            word_figure_sequence="图",
            word_table_sequence="表",
        )

        document = Document(output_docx)
        self.assertIn("document-type=proposal", document.core_properties.keywords)
        self.assertIn("template-year=2024", document.core_properties.keywords)
        self.assertEqual(len(document.sections), 2)
        self.assertIsNone(self._page_number_attributes(document.sections[0]))
        self.assertEqual(
            self._page_number_attributes(document.sections[1]),
            {"format": "decimal", "start": "1"},
        )
        self.assertEqual([len(table.rows) for table in document.tables], [1, 1, 1, 1, 1, 1, 1, 7])
        title_table = document.tables[0]
        self.assertEqual(title_table.alignment, BUILD_WORD.WD_TABLE_ALIGNMENT.RIGHT)
        self.assertEqual(
            [int(grid.get(BUILD_WORD.qn("w:w"))) for grid in title_table._tbl.tblGrid],
            [1480, 6094],
        )
        title_row_height = title_table.rows[0]._tr.trPr.find(BUILD_WORD.qn("w:trHeight"))
        self.assertEqual(title_row_height.get(BUILD_WORD.qn("w:hRule")), "atLeast")
        self.assertEqual(title_row_height.get(BUILD_WORD.qn("w:val")), "960")
        title_paragraphs = title_table.cell(0, 1).paragraphs
        self.assertEqual(len(title_paragraphs), 2)
        self.assertEqual("".join(paragraph.text for paragraph in title_paragraphs), long_proposal_title)
        for title_paragraph in title_paragraphs:
            self.assertEqual(title_paragraph.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertEqual(title_paragraph.runs[0].font.size.pt, 14)
            self.assertIsNotNone(
                title_paragraph._p.find(
                    BUILD_WORD.qn("w:pPr")
                    + "/"
                    + BUILD_WORD.qn("w:pBdr")
                    + "/"
                    + BUILD_WORD.qn("w:bottom")
                )
            )
        for title_paragraph in title_paragraphs:
            self.assertIsNotNone(
                title_paragraph._p.find(
                    BUILD_WORD.qn("w:pPr")
                    + "/"
                    + BUILD_WORD.qn("w:pBdr")
                    + "/"
                    + BUILD_WORD.qn("w:between")
                )
            )
        form_table_indent = document.tables[-1]._tbl.tblPr.find(BUILD_WORD.qn("w:tblInd"))
        self.assertIsNotNone(form_table_indent)
        self.assertEqual(form_table_indent.get(BUILD_WORD.qn("w:w")), "249")
        self.assertEqual(
            [
                "".join(row.cells[0].paragraphs[0].text.split())
                for table in document.tables[:7]
                for row in table.rows
            ],
            [
                "题目",
                "学院（部）",
                "专业",
                "学生姓名",
                "学号",
                "指导教师（校内）",
                "指导教师（校外）",
            ],
        )
        self.assertTrue(
            all(
                row.cells[0].paragraphs[0].style.name
                == "CoverField"
                for table in document.tables[:7]
                for row in table.rows
            )
        )
        cover_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("本科生毕业论文开题报告", cover_text)
        self.assertIn("教务处制表", cover_text)
        self.assertNotIn("2026年7月", cover_text)
        self.assertEqual(len(document.inline_shapes), 2)
        self.assertEqual(
            document.paragraphs[0].paragraph_format.line_spacing_rule,
            WD_LINE_SPACING.AT_LEAST,
        )
        self.assertEqual(
            document.paragraphs[2].paragraph_format.line_spacing_rule,
            WD_LINE_SPACING.AT_LEAST,
        )
        self.assertAlmostEqual(document.sections[0].left_margin.pt, 85.05, places=2)
        self.assertAlmostEqual(document.sections[0].right_margin.pt, 62.35, places=2)
        self.assertAlmostEqual(document.sections[1].top_margin.pt, 70.9, places=2)
        self.assertAlmostEqual(document.sections[1].bottom_margin.pt, 51.05, places=2)
        self.assertAlmostEqual(document.sections[1].left_margin.pt, 85.05, places=2)
        self.assertAlmostEqual(document.sections[1].right_margin.pt, 62.35, places=2)
        page_breaks = [
            node
            for node in document.element.body.iter(BUILD_WORD.qn("w:br"))
            if node.get(BUILD_WORD.qn("w:type")) == "page"
        ]
        self.assertEqual(len(page_breaks), 0)
        self.assertIn(
            " PAGE ",
            self._field_instructions(document.sections[1].footer._element),
        )

    def test_literature_review_has_four_sections_and_native_word_fields(self) -> None:
        source = Document(self.reference_docx)
        source.paragraphs[0].text = BUILD_WORD.SECTION_REVIEW_COVER_END
        source.add_paragraph("摘要", style="FrontTitle")
        source.add_paragraph("中文摘要正文。", style="ChineseAbstract")
        source.add_paragraph("关键词：资料；标准；研究", style="Keywords")
        source.add_paragraph("Abstract", style="FrontTitle")
        source.add_paragraph("An English abstract.", style="EnglishAbstract")
        source.add_paragraph("Keywords: evidence; standards; research", style="Keywords")
        source.add_paragraph(BUILD_WORD.SECTION_REVIEW_ABSTRACT_END)
        source.add_paragraph("无标题引言正文。", style="Body Text")
        source.add_paragraph("研究背景", style=BUILD_WORD.STYLE_HEADING_1)
        source.add_paragraph("章节正文。", style="Body Text")
        source.add_paragraph("图1 技术路线", style="Caption")

        labels = BUILD_WORD.LabelRegistry()
        labels.register("fig:review", "1", "图")
        reference = labels.placeholder("fig:review", parenthesized=False)
        citations = BUILD_WORD.CitationRegistry()
        citation = citations.placeholder("source1")
        source.add_paragraph(f"见图{reference}，参考文献{citation}。")
        source.add_paragraph("结论与展望", style=BUILD_WORD.STYLE_HEADING_1)
        source.add_paragraph("结论正文。", style="Body Text")
        source.add_paragraph(BUILD_WORD.SECTION_REVIEW_REFERENCES)
        source.add_paragraph("参考文献", style="Heading 1")
        source.add_paragraph("[1] 示例文献", style="Bibliography")

        input_docx = self.root / "review-input.docx"
        output_docx = self.root / "review-output.docx"
        source.save(input_docx)
        BUILD_WORD._postprocess_review_docx(
            input_docx,
            output_docx,
            metadata=self._metadata(),
            project_root=ROOT,
            main_dir=self.root,
            metadata_dir=self.root,
            discipline="humanities",
            citation_mode="linked",
            labels=labels,
            citations=citations,
            bibliography_items=[],
            table_layouts=[],
            word_figure_sequence="图",
            word_table_sequence="表",
        )

        document = Document(output_docx)
        self.assertEqual(len(document.sections), 4)
        self.assertIsNone(self._page_number_attributes(document.sections[0]))
        self.assertEqual(
            self._page_number_attributes(document.sections[1]),
            {"format": "decimal", "start": "1"},
        )
        self.assertEqual(self._page_number_attributes(document.sections[2]), {"format": "decimal", "start": ""})
        self.assertEqual(
            next(p for p in document.paragraphs if p.text == "研究背景").style.name,
            BUILD_WORD.STYLE_REVIEW_HEADING_1,
        )
        self.assertEqual(next(p for p in document.paragraphs if p.text == "结论与展望").style.name, BUILD_WORD.STYLE_REVIEW_HEADING_1)
        self.assertEqual(next(p for p in document.paragraphs if p.text == "无标题引言正文。").style.name, BUILD_WORD.STYLE_REVIEW_BODY)
        self.assertFalse(document.tables)
        cover_rows = [p for p in document.paragraphs if p.style.name == BUILD_WORD.STYLE_REVIEW_COVER_FIELD]
        self.assertEqual(len(cover_rows), 7)
        self.assertTrue(cover_rows[0].text.startswith("题    目\t"))
        self.assertIn(self._metadata()["title"], cover_rows[0].text)
        self.assertAlmostEqual(cover_rows[0].paragraph_format.left_indent.cm, 3.7, places=2)
        self.assertAlmostEqual(cover_rows[0].paragraph_format.first_line_indent.cm, -3.7, places=2)
        self.assertTrue(cover_rows[1].text.startswith("学部（院）\t"))
        self.assertTrue(cover_rows[-2].text.startswith("指导教师1\t"))
        self.assertTrue(cover_rows[-1].text.startswith("指导教师2\t"))
        for row in cover_rows:
            self.assertIsNone(row._p.find(BUILD_WORD.qn("w:pPr") + "/" + BUILD_WORD.qn("w:pBdr")))
            for run in row.runs:
                if run.style.name == BUILD_WORD.STYLE_REVIEW_COVER_LABEL:
                    self.assertFalse(run.style.font.underline)
                elif run.style.name == BUILD_WORD.STYLE_REVIEW_COVER_VALUE:
                    self.assertTrue(run.style.font.underline)
        self.assertIn("职称", cover_rows[-1].text)
        self.assertEqual(next(p for p in document.paragraphs if p.style.name == BUILD_WORD.STYLE_REVIEW_THESIS_TITLE).text, self._metadata()["title"])
        self.assertEqual(next(p for p in document.paragraphs if p.style.name == BUILD_WORD.STYLE_REVIEW_AFFILIATION).text, "（示例学部 示例城市 000000）")
        self.assertEqual(next(p for p in document.paragraphs if p.text == "Abstract").style.name, BUILD_WORD.STYLE_REVIEW_ENGLISH_ABSTRACT_TITLE)
        self.assertFalse(any("TOC" in value for value in self._field_instructions(document.element)))

        caption = next(p for p in document.paragraphs if p.text == "图1 技术路线")
        self.assertEqual(caption.style.name, BUILD_WORD.STYLE_FIGURE_CAPTION)
        self.assertIn(
            " SEQ 图 \\r 1 \\* ARABIC ",
            self._field_instructions(caption._p),
        )
        bookmark_names = {
            node.get(BUILD_WORD.qn("w:name"))
            for node in caption._p.iter(BUILD_WORD.qn("w:bookmarkStart"))
        }
        self.assertIn(BUILD_WORD._label_bookmark_name("fig:review"), bookmark_names)
        self.assertIn(
            BUILD_WORD._label_number_bookmark_name("fig:review"),
            bookmark_names,
        )

        reference_paragraph = next(
            p for p in document.paragraphs if p.text == "见图1，参考文献[1]。"
        )
        instructions = self._field_instructions(reference_paragraph._p)
        self.assertIn(
            f" REF {BUILD_WORD._label_bookmark_name('fig:review')} \\h ",
            instructions,
        )
        citation_bookmark = BUILD_WORD._citation_bookmark_name("source1")
        self.assertIn(
            f" REF {citation_bookmark} \\h \\* CHARFORMAT ",
            instructions,
        )
        all_bookmarks = {
            node.get(BUILD_WORD.qn("w:name"))
            for node in document.element.iter(BUILD_WORD.qn("w:bookmarkStart"))
        }
        ref_targets = {
            match.group(1)
            for instruction in instructions
            if (match := re.search(r"\bREF\s+(\S+)", instruction))
        }
        self.assertTrue(ref_targets.issubset(all_bookmarks))

        citation_style = document.styles[BUILD_WORD.STYLE_CITATION]
        vertical = citation_style.element.find(
            BUILD_WORD.qn("w:rPr") + "/" + BUILD_WORD.qn("w:vertAlign")
        )
        self.assertIsNotNone(vertical)
        self.assertEqual(vertical.get(BUILD_WORD.qn("w:val")), "superscript")
        styled_runs = []
        for run in reference_paragraph._p.iter(BUILD_WORD.qn("w:r")):
            style = run.find(
                BUILD_WORD.qn("w:rPr") + "/" + BUILD_WORD.qn("w:rStyle")
            )
            if style is not None and style.get(BUILD_WORD.qn("w:val")) == citation_style.style_id:
                styled_runs.append(run)
        self.assertGreaterEqual(len(styled_runs), 7)
        self.assertTrue(
            all(
                run.find(
                    BUILD_WORD.qn("w:rPr") + "/" + BUILD_WORD.qn("w:vertAlign")
                )
                is None
                for run in styled_runs
            ),
            "引用上标应由字符样式驱动，而不是逐个 run 直接编码。",
        )

    def test_review_second_advisor_prefers_explicit_slot_and_accepts_legacy_keys(self) -> None:
        cases = (
            ({"external-advisor": "旧字段导师", "external-advisor-title": "讲师"}, "旧字段导师", "讲师"),
            ({"external-advisor": "旧字段导师", "second-advisor": "第二导师", "second-advisor-title": "副教授"}, "第二导师", "副教授"),
            ({"external-advisor": "旧字段导师", "second-advisor": "", "second-advisor-title": ""}, "", ""),
        )
        for extra, expected_name, expected_title in cases:
            with self.subTest(metadata=extra):
                document = Document(self.reference_docx)
                BUILD_WORD._profiles_module()._insert_review_cover(
                    document, document.paragraphs[0], {**self._metadata(), **extra},
                    project_root=ROOT, main_dir=self.root, metadata_dir=self.root,
                    discipline="humanities",
                )
                row = next(p for p in document.paragraphs if p.text.startswith("指导教师2\t"))
                values = [run.text.strip("\t") for run in row.runs if run.style.name == BUILD_WORD.STYLE_REVIEW_COVER_VALUE]
                self.assertEqual(values, [expected_name, expected_title])
                self.assertIn("职称", row.text)
                self.assertFalse(document.tables)

    def test_review_build_preserves_abstracts_math_and_standard_citations_for_both_disciplines(self) -> None:
        try:
            pandoc = BUILD_WORD.find_pandoc(None, ROOT)
        except BUILD_WORD.BuildError as error:
            self.skipTest(str(error))
        for discipline in ("humanities", "science"):
            with self.subTest(discipline=discipline):
                project = self.root / ("review-build-" + discipline)
                chapters = project / "chapters"
                chapters.mkdir(parents=True)
                (project / "metadata.tex").write_text(
                    r"\yibinsetup{template-year=2024,title={结构回归综述},author={测试作者},college={测试学部},affiliation-city={测试市},affiliation-postcode={000000}}",
                    encoding="utf-8",
                )
                main_text = (
                    "\\documentclass[literature-review," + discipline + "]{yibinthesis}\n"
                    "\\input{metadata}\n\\addbibresource{references.bib}\n"
                    "\\begin{document}\n\\makeyibinliteraturereviewcover\n"
                    "\\input{chapters/00-abstract-cn}\n\\input{chapters/01-abstract-en}\n"
                    "\\input{chapters/body}\n\\printyibinbibliography\n\\end{document}\n"
                )
                main = project / "main.tex"
                main.write_text(main_text, encoding="utf-8")
                (chapters / "00-abstract-cn.tex").write_text(
                    r"\begin{cnabstract}这是用于验证摘要内容保留的文字。\cnkeywords{资料；标准；研究}\end{cnabstract}",
                    encoding="utf-8",
                )
                (chapters / "01-abstract-en.tex").write_text(
                    r"\begin{enabstract}This abstract verifies the document conversion.\enkeywords{evidence; standards; research}\end{enabstract}",
                    encoding="utf-8",
                )
                (chapters / "body.tex").write_text(
                    "这是无标题引言。\n\n\\chapter{相关研究}\n"
                    "正文数学 $k$ 保持可编辑。\\yibincite{fixture_standard,fixture_math}\n"
                    "\\section{研究方法}\n方法正文。\n"
                    "\\chapter{结论与展望}\n结论正文。\n",
                    encoding="utf-8",
                )
                (project / "references.bib").write_text(
                    "@standard{fixture_standard, author={{Test Organization}}, title={Test standard}, number={TEST 1}, date={2024}, location={Test City}, publisher={Test Publisher}}\n"
                    "@inproceedings{fixture_math, author={Author, Test}, title={{$k$-means++}: CO\\textsubscript{2}, x\\textsuperscript{2}, \\textit{in vivo} and \\textbf{bold}}, booktitle={Test Proceedings}, date={2024}, pages={1--2}}\n"
                    "@article{fixture_uncited, author={Other, Test}, title={Uncited fixture}, journaltitle={Test Journal}, date={2024}}\n",
                    encoding="utf-8",
                )
                output = project / "review.docx"
                args = BUILD_WORD.make_parser().parse_args([
                    "--main", str(main), "--output", str(output),
                    "--reference-doc", str(self.reference_docx), "--pandoc", str(pandoc),
                    "--citation-mode", "native", "--keep-intermediate",
                ])
                expected_title = "结构回归综述"
                if discipline == "humanities":
                    override = project / "metadata-override.tex"
                    expected_title = "覆盖元数据综述"
                    override.write_text(
                        (project / "metadata.tex").read_text(encoding="utf-8").replace("结构回归综述", expected_title),
                        encoding="utf-8",
                    )
                    args.metadata = str(override)
                BUILD_WORD.build(args)
                document = Document(output)
                texts = [paragraph.text for paragraph in document.paragraphs]
                self.assertEqual(len(document.sections), 4)
                self.assertLess(texts.index(expected_title), texts.index("摘要"))
                self.assertLess(texts.index("摘要"), texts.index("Abstract"))
                self.assertLess(texts.index("Abstract"), texts.index("这是无标题引言。"))
                self.assertIn("关键词：资料；标准；研究", texts)
                self.assertIn("Keywords: evidence; standards; research", texts)
                self.assertIn("（测试学部 测试市 000000）", texts)
                conclusion = next(p for p in document.paragraphs if p.text == "结论与展望")
                self.assertEqual(conclusion.style.name, BUILD_WORD.STYLE_REVIEW_HEADING_1)
                self.assertNotIn("引言", texts)
                self.assertNotIn("目录", texts)
                self.assertFalse(any(text.startswith("YIBIN_INTERNAL_") for text in texts))
                self.assertNotIn("$k$", "\n".join(texts))
                math_values = [
                    "".join(part.text or "" for part in math.iter(BUILD_WORD.qn("m:t")))
                    for math in document.element.iter(BUILD_WORD.qn("m:oMath"))
                ]
                self.assertGreaterEqual(math_values.count("k"), 2)
                bibliography = "\n".join(p.text for p in document.paragraphs if p.style.name == BUILD_WORD.STYLE_REVIEW_BIBLIOGRAPHY)
                standard_entry = next(line for line in bibliography.splitlines() if "Test standard" in line)
                self.assertIn("[S]", standard_entry)
                self.assertNotIn("[Z]", standard_entry)
                self.assertNotIn("Uncited fixture", bibliography)
                math_entry = next(p for p in document.paragraphs if "means++" in p.text)
                vertical_alignments = {
                    node.get(BUILD_WORD.qn("w:val"))
                    for node in math_entry._p.iter(BUILD_WORD.qn("w:vertAlign"))
                }
                self.assertTrue({"subscript", "superscript"}.issubset(vertical_alignments))
                for tag, word in (("i", "vivo"), ("b", "bold")):
                    self.assertTrue(any(
                        word in "".join(node.text or "" for node in run.iter(BUILD_WORD.qn("w:t"))).lower()
                        and run.find(BUILD_WORD.qn("w:rPr") + "/" + BUILD_WORD.qn("w:" + tag)) is not None
                        for run in math_entry._p.iter(BUILD_WORD.qn("w:r"))
                    ))
                heading = document.styles[BUILD_WORD.STYLE_REVIEW_HEADING_1]
                self.assertFalse(heading.paragraph_format.page_break_before)
                self.assertEqual(heading.paragraph_format.alignment, BUILD_WORD.WD_ALIGN_PARAGRAPH.LEFT)
                indent = heading.element.get_or_add_pPr().find(BUILD_WORD.qn("w:ind"))
                self.assertEqual(indent.get(BUILD_WORD.qn("w:firstLineChars")), "0" if discipline == "science" else "200")
                with zipfile.ZipFile(output) as package:
                    self.assertTrue(any(
                        b"Test standard" in package.read(name)
                        for name in package.namelist()
                        if name.startswith("customXml/") and name.endswith(".xml")
                    ))
                abstract_path = chapters / "00-abstract-cn.tex"
                abstract_source = abstract_path.read_text(encoding="utf-8")
                abstract_path.write_text(abstract_source + "\n不应静默丢弃的正文。", encoding="utf-8")
                with self.assertRaisesRegex(BUILD_WORD.BuildError, "摘要文件不能混入正文"):
                    BUILD_WORD.build(args)
                abstract_path.write_text(abstract_source, encoding="utf-8")
                main.write_text(main_text.replace("\\input{chapters/00-abstract-cn}\n", ""), encoding="utf-8")
                with self.assertRaisesRegex(BUILD_WORD.BuildError, "中文摘要"):
                    BUILD_WORD.build(args)


if __name__ == "__main__":
    unittest.main()
