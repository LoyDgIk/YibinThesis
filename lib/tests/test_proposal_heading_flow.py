"""Proposal heading fixtures and real Pandoc-to-DOCX integration checks."""

from __future__ import annotations

import re
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "lib"))

import word_core as WORD  # noqa: E402
from build_reference_docx import build_reference_docx  # noqa: E402


FIELD_CASES = (
    ("significance", "意义", "10-significance.tex"),
    ("research-status", "现状", "20-research-status.tex"),
)
HEADING_COMMANDS = ("section", "subsection", "subsubsection", "paragraph")
LEVEL_NAMES = ("一级", "二级", "三级", "四级")
EXPECTED_LEVELS = (0, 1, 2, 3, 0, 1, 2, 3)
EXPECTED_FORMATS = {
    "humanities": (
        ("chineseCounting", "%1、"),
        ("chineseCounting", "（%2）"),
        ("decimal", "%3."),
        ("decimal", "（%4）"),
    ),
    "science": (
        ("decimal", "%1"),
        ("decimal", "%1.%2"),
        ("decimal", "%1.%2.%3"),
        ("decimal", "%1.%2.%3.%4"),
    ),
}


def _heading_titles(prefix: str) -> list[str]:
    return [f"{prefix}{branch}{level}" for branch in ("A", "B") for level in LEVEL_NAMES]


def _field_source(key: str, prefix: str) -> str:
    parts = []
    for branch in ("A", "B"):
        for index, (command, level) in enumerate(zip(HEADING_COMMANDS, LEVEL_NAMES)):
            title = f"{prefix}{branch}{level}"
            parts.append(rf"\{command}{{{title}}}")
            sentence = f"这是{title}的正文。"
            if branch == "A" and index == 0:
                sentence += r"共同参考资料见\yibincite{fixturebook}。"
            parts.append(sentence + "\n")
    parts.append(
        rf"""\begin{{table}}[htbp]
\centering
\caption{{{prefix}对照表}}\label{{tab:{key}}}
\begin{{tabular}}{{ll}}
\toprule
项目 & 内容 \\
\midrule
测试项 & 测试值 \\
\bottomrule
\end{{tabular}}
\end{{table}}

表\ref{{tab:{key}}}用于核对字段间的全局表序号。

\section*{{{prefix}不编号说明}}
这是{prefix}不编号说明的正文。
"""
    )
    return "\n".join(parts)


def write_fixture(project: Path, discipline: str) -> Path:
    """Write synthetic proposal sources; return an absolute main.tex path."""
    if discipline not in EXPECTED_FORMATS:
        raise ValueError(f"Unsupported discipline: {discipline}")
    project = project.resolve()
    fields = project / "fields"
    fields.mkdir(parents=True, exist_ok=True)
    (project / "metadata.tex").write_text(
        r"""\yibinsetup{
  template-year = 2022,
  title = {开题标题层级测试},
  author = {测试作者},
  student-id = {20260001},
  college = {测试学部},
  major = {测试专业},
  grade = {2026},
  class-name = {1班},
  advisor = {测试导师},
  advisor-title = {讲师},
  date = {2026年10月},
  logo = {builtin}
}
""",
        encoding="utf-8",
    )
    (project / "references.bib").write_text(
        """% Synthetic bibliography entry for this integration fixture only.
@book{fixturebook,
  author = {Author, Test},
  title = {Fixture Reference Book},
  location = {Test City},
  publisher = {Fixture Publisher},
  date = {2024}
}
""",
        encoding="utf-8",
    )
    for key, prefix, filename in FIELD_CASES:
        (fields / filename).write_text(_field_source(key, prefix), encoding="utf-8")
    main = project / "main.tex"
    main.write_text(
        rf"""\documentclass[proposal,{discipline}]{{yibinthesis}}
\input{{metadata}}
\addbibresource{{references.bib}}
\begin{{document}}
\makeyibinproposal
\yibinproposalfield{{significance}}{{\input{{fields/10-significance}}}}
\yibinproposalfield{{research-status}}{{\input{{fields/20-research-status}}}}
\yibinproposalfield{{research-content}}{{主要研究内容测试占位。}}
\yibinproposalfield{{research-approach}}{{研究思路测试占位。}}
\yibinproposalfield{{schedule}}{{研究进度测试占位。}}
\yibinproposalfield{{references}}{{\printyibinproposalbibliography}}
\yibinproposalfield{{advisor-opinion}}{{导师意见测试占位。}}
\end{{document}}
""",
        encoding="utf-8",
    )
    return main


def _field_instructions(element) -> list[str]:
    instructions = [node.text or "" for node in element.iter(qn("w:instrText"))]
    instructions.extend(
        node.get(qn("w:instr"), "") for node in element.iter(qn("w:fldSimple"))
    )
    return instructions


class ProposalHeadingFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        try:
            cls.pandoc = WORD.find_pandoc(None, ROOT)
        except WORD.BuildError as error:
            raise unittest.SkipTest(str(error)) from error
        cls.temporary = tempfile.TemporaryDirectory(prefix="proposal-heading-flow-")
        cls.addClassCleanup(cls.temporary.cleanup)
        cls.root = Path(cls.temporary.name)
        cls.reference_docx = cls.root / "reference.docx"
        build_reference_docx(cls.reference_docx)

    def _check_heading_style(self, paragraph, level: int, discipline: str) -> None:
        style = paragraph.style
        self.assertEqual(style.font.size.pt, (16, 15, 14, 14)[level])
        self.assertEqual(style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertIn(paragraph.paragraph_format.alignment, (None, WD_ALIGN_PARAGRAPH.LEFT))
        self.assertFalse(style.paragraph_format.page_break_before)
        self.assertFalse(paragraph.paragraph_format.page_break_before)
        self.assertTrue(style.paragraph_format.keep_with_next)
        self.assertIsNot(paragraph.paragraph_format.keep_with_next, False)
        indent = style.element.find(qn("w:pPr") + "/" + qn("w:ind"))
        self.assertIsNotNone(indent)
        self.assertEqual(
            indent.get(qn("w:firstLineChars")),
            "0" if discipline == "science" else "200",
        )
        properties = [node.tag for node in style.element.pPr]
        numbering_position = properties.index(qn("w:numPr"))
        for tag in ("keepNext", "keepLines", "pageBreakBefore"):
            self.assertLess(properties.index(qn(f"w:{tag}")), numbering_position)
        for tag in ("spacing", "ind", "jc"):
            self.assertGreater(properties.index(qn(f"w:{tag}")), numbering_position)

    def _check_numbering_definition(
        self, document, num_id: str, discipline: str, root_number: int
    ) -> None:
        numbering = document.part.numbering_part.element
        numbers = {node.get(qn("w:numId")): node for node in numbering.findall(qn("w:num"))}
        self.assertIn(num_id, numbers)
        instance = numbers[num_id]
        abstract_id = instance.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstracts = {
            node.get(qn("w:abstractNumId")): node
            for node in numbering.findall(qn("w:abstractNum"))
        }
        self.assertIn(abstract_id, abstracts)
        levels = sorted(
            abstracts[abstract_id].findall(qn("w:lvl")),
            key=lambda node: int(node.get(qn("w:ilvl"))),
        )
        self.assertEqual([node.get(qn("w:ilvl")) for node in levels], ["0", "1", "2", "3"])
        self.assertEqual(
            [
                (node.find(qn("w:numFmt")).get(qn("w:val")),
                 node.find(qn("w:lvlText")).get(qn("w:val")))
                for node in levels
            ],
            list(EXPECTED_FORMATS[discipline]),
        )
        self.assertTrue(all(node.find(qn("w:start")).get(qn("w:val")) == "1" for node in levels))

        # Explicit root numbers survive Word's pagination of the proposal table.
        first_level_overrides = [
            node for node in instance.findall(qn("w:lvlOverride"))
            if node.get(qn("w:ilvl")) == "0"
        ]
        self.assertEqual(len(first_level_overrides), 1, "Each root needs an explicit level-0 start.")
        restart = first_level_overrides[0].find(qn("w:startOverride"))
        self.assertIsNotNone(restart)
        self.assertEqual(restart.get(qn("w:val")), str(root_number))

    def _check_document(self, document, discipline: str) -> None:
        forms = [
            table for table in document.tables
            if table.style.name == WORD.STYLE_PROPOSAL_FORM_TABLE
        ]
        self.assertEqual(len(forms), 1)
        form = forms[0]
        self.assertEqual(len(form.rows), 7)
        self.assertEqual(
            ["".join(row.cells[0].text.split()) for row in form.rows],
            ["选题意义", "国内外研究现状概述", "主要研究内容", "拟采用的研究思路",
             "研究工作安排及进度", "参考文献目录", "指导教师意见"],
        )
        heading_styles = (
            WORD.STYLE_HEADING_1, WORD.STYLE_HEADING_2,
            WORD.STYLE_HEADING_3, WORD.STYLE_HEADING_4,
        )
        branch_num_ids = []
        caption_numbers = []
        citation_target = WORD._citation_bookmark_name("fixturebook")
        for row_index, (key, prefix, _) in enumerate(FIELD_CASES):
            with self.subTest(discipline=discipline, field=key):
                cell = form.rows[row_index].cells[1]
                headings = [p for p in cell.paragraphs if p.style.name in heading_styles]
                self.assertEqual(len(headings), 8)
                self.assertEqual([p.text.strip() for p in headings], _heading_titles(prefix))
                self.assertEqual(
                    [p.style.name for p in headings],
                    [heading_styles[level] for level in EXPECTED_LEVELS],
                )
                bindings = []
                for paragraph, level in zip(headings, EXPECTED_LEVELS):
                    num_pr = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
                    self.assertIsNotNone(num_pr, paragraph.text)
                    actual_level = int(num_pr.find(qn("w:ilvl")).get(qn("w:val")))
                    num_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
                    bindings.append((num_id, actual_level))
                    self._check_heading_style(paragraph, level, discipline)
                    self.assertTrue(any(
                        p.text.startswith(f"这是{paragraph.text.strip()}的正文。")
                        for p in cell.paragraphs
                    ))
                self.assertEqual([level for _, level in bindings], list(EXPECTED_LEVELS))
                for root_number, start in enumerate((0, 4), start=1):
                    branch = bindings[start:start + 4]
                    self.assertEqual(len({num_id for num_id, _ in branch}), 1)
                    branch_num_id = branch[0][0]
                    branch_num_ids.append(branch_num_id)
                    self._check_numbering_definition(
                        document, branch_num_id, discipline, root_number
                    )

                star = [p for p in cell.paragraphs if p.text == f"{prefix}不编号说明"]
                self.assertEqual(len(star), 1)
                self.assertEqual(star[0].style.name, WORD.STYLE_PROPOSAL_UNNUMBERED_HEADING)
                self.assertIsNone(star[0]._p.find(qn("w:pPr") + "/" + qn("w:numPr")))
                self.assertIsNone(star[0].style.element.find(qn("w:pPr") + "/" + qn("w:numPr")))
                self.assertIn(f"这是{prefix}不编号说明的正文。", [p.text for p in cell.paragraphs])

                self.assertEqual(len(cell.tables), 1)
                captions = [
                    p for p in cell.paragraphs
                    if p.style.name == WORD.STYLE_TABLE_CAPTION
                    and f"{prefix}对照表" in p.text
                ]
                self.assertEqual(len(captions), 1)
                caption_match = re.fullmatch(
                    rf"表\s*(\d+)\s*{re.escape(prefix)}对照表", captions[0].text.strip()
                )
                self.assertIsNotNone(caption_match, captions[0].text)
                caption_numbers.append(int(caption_match.group(1)))
                sequences = [
                    code for code in _field_instructions(captions[0]._p)
                    if re.match(r'^\s*SEQ\s+(?:"表"|表)(?:\s|$)', code, re.IGNORECASE)
                ]
                self.assertEqual(len(sequences), 1)
                for code in sequences:
                    reset = re.search(r"\\r\s+(\d+)", code)
                    if reset is not None:
                        self.assertEqual(int(reset.group(1)), row_index + 1)
                citation_paragraphs = [
                    p for p in cell.paragraphs
                    if any(citation_target in code for code in _field_instructions(p._p))
                ]
                self.assertEqual(len(citation_paragraphs), 1)
                self.assertIn("[1]", citation_paragraphs[0].text)

        self.assertEqual(len(branch_num_ids), 4)
        self.assertEqual(
            len(set(branch_num_ids)), 4,
            "Each root branch must have its own numbering instance across fields.",
        )
        self.assertEqual(caption_numbers, [1, 2], "Table numbering must remain document-wide.")
        starts = list(document.element.iter(qn("w:bookmarkStart")))
        ends = list(document.element.iter(qn("w:bookmarkEnd")))
        bookmark_names = {node.get(qn("w:name")) for node in starts}
        self.assertEqual(
            Counter(node.get(qn("w:id")) for node in starts),
            Counter(node.get(qn("w:id")) for node in ends),
        )
        references = [
            match.group(1)
            for code in _field_instructions(document.element)
            for match in re.finditer(r"\bREF\s+([A-Za-z0-9_]+)", code, re.IGNORECASE)
        ]
        self.assertGreaterEqual(len(references), 4)
        self.assertTrue(set(references).issubset(bookmark_names))
        self.assertEqual(references.count(citation_target), 2)
        for key, _, _ in FIELD_CASES:
            self.assertIn(WORD._label_bookmark_name(f"tab:{key}"), references)
        bibliography = [
            p for p in form.rows[5].cells[1].paragraphs
            if "fixture reference book" in p.text.casefold()
        ]
        self.assertEqual(len(bibliography), 1)
        self.assertRegex(bibliography[0].text, r"^\s*\[1\]")

    def test_heading_flow_for_both_disciplines(self) -> None:
        for discipline in ("humanities", "science"):
            with self.subTest(discipline=discipline):
                project = self.root / discipline
                main = write_fixture(project, discipline)
                output = project / "proposal.docx"
                args = WORD.make_parser().parse_args([
                    "--main", str(main),
                    "--output", str(output),
                    "--reference-doc", str(self.reference_docx),
                    "--pandoc", str(self.pandoc),
                    "--citation-mode", "linked",
                    "--keep-intermediate",
                ])
                WORD.build(args)
                self._check_document(Document(output), discipline)


if __name__ == "__main__":
    unittest.main()
