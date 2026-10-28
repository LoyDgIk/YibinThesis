"""Proposal and literature-review Word profile assembly."""

from __future__ import annotations

import sys
import unicodedata

from docx.enum.text import WD_LINE_SPACING


def bind_core(core):
    """Expose core helpers to profile functions without duplicating them."""
    module = sys.modules[__name__]
    profile_functions = {
        "_build_non_thesis",
        "_postprocess_proposal_docx",
        "_postprocess_review_docx",
    }
    for name in dir(core):
        if (
            not name.startswith("__")
            and name not in {"bind_core", "sys"}
            and name not in profile_functions
        ):
            setattr(module, name, getattr(core, name))
    return module


def _set_form_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def _set_form_cell_margins(cell, *, top: int = 90, left: int = 110, bottom: int = 90, right: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _wrap_proposal_title(text: str, width_cm: float, font_pt: float = 14) -> list[str]:
    """Split a long cover title into underlined, full-width Word paragraphs."""
    max_width = Cm(width_cm).pt - 4.5
    lines: list[str] = []
    line: list[str] = []
    line_width = 0.0
    for char in text.strip():
        char_width = font_pt if unicodedata.east_asian_width(char) in {"F", "W"} else font_pt / 2
        if line and line_width + char_width > max_width:
            lines.append("".join(line).rstrip())
            line = []
            line_width = 0.0
        line.append(char)
        line_width += char_width
    if line:
        lines.append("".join(line).rstrip())
    return lines or [""]


def _set_paragraph_between_border(paragraph) -> None:
    """Keep a rule between adjacent wrapped proposal-title paragraphs."""
    borders = paragraph._p.get_or_add_pPr().find(qn("w:pBdr"))
    between = borders.find(qn("w:between"))
    if between is None:
        between = OxmlElement("w:between")
        borders.append(between)
    between.set(qn("w:val"), "single")
    between.set(qn("w:sz"), "8")
    between.set(qn("w:space"), "0")
    between.set(qn("w:color"), "000000")


def _move_nodes_to_cell(cell, nodes: list[object], document: Document) -> None:
    anchor = cell.paragraphs[0]
    for node in nodes:
        anchor._p.addprevious(node)
    anchor.style = document.styles[STYLE_PROPOSAL_BODY]
    anchor.paragraph_format.space_before = Pt(0)
    anchor.paragraph_format.space_after = Pt(0)
    anchor.paragraph_format.line_spacing = Pt(1)
    for paragraph in cell.paragraphs:
        if paragraph is anchor:
            continue
        if paragraph.style.name in {STYLE_BODY, STYLE_FIRST_PARAGRAPH, "Normal", "Body Text"}:
            paragraph.style = document.styles[STYLE_PROPOSAL_BODY]


def _extract_marker_ranges(document: Document) -> dict[str, list[object]]:
    body = document._element.body
    children = list(body)
    markers: list[tuple[str, int]] = []
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
        if text.startswith(PROPOSAL_FIELD_MARKER_PREFIX):
            markers.append((text.removeprefix(PROPOSAL_FIELD_MARKER_PREFIX).lower(), index))
    expected = list(PROPOSAL_FIELD_ORDER)
    if [key for key, _ in markers] != expected:
        raise BuildError("Word 开题报告字段标记缺失或顺序错误。")
    result: dict[str, list[object]] = {}
    section_index = children.index(body.sectPr) if body.sectPr is not None else len(children)
    for marker_index, (key, start) in enumerate(markers):
        end = markers[marker_index + 1][1] if marker_index + 1 < len(markers) else section_index
        nodes = children[start + 1 : end]
        if key == "references":
            nodes = [
                node
                for node in nodes
                if not (
                    node.tag == qn("w:p")
                    and "".join(part.text or "" for part in node.iter(qn("w:t"))).strip()
                    == "参考文献"
                )
            ]
        result[key] = nodes
    for child in list(body):
        if child is not body.sectPr:
            body.remove(child)
    return result


def _add_profile_page_field(section, document: Document) -> None:
    section.footer.is_linked_to_previous = False
    _clear_story(section.footer)
    paragraph = section.footer.paragraphs[0]
    paragraph.style = document.styles["Footer"]
    _add_page_field(paragraph)
    for run in paragraph.runs:
        _set_run_fonts(run, "SimSun", "Times New Roman", 9)


def _insert_proposal_cover(
    document: Document,
    metadata: dict[str, str],
    *,
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
) -> None:
    """Insert the 2024 opening-report cover and its information rows."""

    assets = project_root / "assets"
    wordmark = assets / "yibin-university-wordmark.png"
    seal = assets / "yibin-university-seal.png"
    if not wordmark.is_file():
        try:
            wordmark = _resolve_cover_logo(
                metadata,
                project_root,
                main_dir,
                metadata_dir,
                allow_project_fallback=True,
            )
        except BuildError:
            # Unit fixtures and external profile callers may omit optional art.
            wordmark = None

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.space_before = Pt(7)
    logo_paragraph.paragraph_format.space_after = Pt(55.5)
    logo_paragraph.paragraph_format.line_spacing = Pt(66)
    logo_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    if wordmark is not None:
        logo_paragraph.add_run().add_picture(
            str(wordmark), width=Cm(6.24), height=Cm(2.33)
        )
    else:
        run = logo_paragraph.add_run("宜宾学院")
        _set_run_fonts(run, "SimHei", "SimHei", 22, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = Pt(38)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(26.5)
    _set_run_fonts(title.add_run("本科生毕业论文开题报告"), "SimHei", "SimHei", 32, bold=True)

    seal_paragraph = document.add_paragraph()
    seal_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    seal_paragraph.paragraph_format.space_before = Pt(0)
    seal_paragraph.paragraph_format.space_after = Pt(60)
    seal_paragraph.paragraph_format.line_spacing = Pt(115.5)
    seal_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    if seal.is_file():
        seal_paragraph.add_run().add_picture(
            str(seal), width=Cm(4.07), height=Cm(4.07)
        )

    fields_gap = document.add_paragraph()
    fields_gap.paragraph_format.space_before = Pt(0)
    fields_gap.paragraph_format.space_after = Pt(0)
    fields_gap.paragraph_format.line_spacing = Pt(1)

    cover_footer = document.add_paragraph()
    cover_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_footer.paragraph_format.space_before = Pt(34)
    cover_footer.paragraph_format.space_after = Pt(0)
    cover_footer.paragraph_format.line_spacing = Pt(18)
    _set_run_fonts(cover_footer.add_run("教务处制表"), "SimSun", "SimSun", 14)

    cover_rows = (
        (["label", "题  目", 2.61], ["value", metadata.get("title", ""), 10.75]),
        (["label", "学院（部）", 2.75], ["value", metadata.get("college", ""), 8.15]),
        (["label", "专  业", 2.75], ["value", metadata.get("major", ""), 8.15]),
        (["label", "学生姓名", 2.75], ["value", metadata.get("author", ""), 8.15]),
        (
            ["label", "学  号", 2.75], ["value", metadata.get("student-id", ""), 4.55],
            ["label", "年级", 1.2], ["value", format_grade_class(metadata), 2.4],
        ),
        (
            ["label", "指导教师（校内）", 4.7], ["value", metadata.get("advisor", ""), 2.45],
            ["label", "职称", 1.15], ["value", metadata.get("advisor-title", ""), 2.6],
        ),
        (
            ["label", "指导教师（校外）", 4.7],
            ["value", metadata.get("external-advisor", ""), 2.45],
            ["label", "职称", 1.15],
            ["value", metadata.get("external-advisor-title", ""), 2.6],
        ),
    )
    for row_index, cells in enumerate(cover_rows):
        title_lines = (
            _wrap_proposal_title(cells[1][1], cells[1][2]) if row_index == 0 else []
        )
        table = _insert_cover_table(
            document,
            cover_footer,
            [(kind, value, width) for kind, value, width in cells],
            row_height_pt=max(48, 24 * len(title_lines)) if row_index == 0 else 36,
        )
        table.alignment = (
            WD_TABLE_ALIGNMENT.RIGHT if row_index == 0 else WD_TABLE_ALIGNMENT.CENTER
        )
        if row_index == 0:
            table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, (kind, _, _) in zip(table.rows[0].cells, cells):
            if row_index == 0 and kind == "value":
                paragraph = cell.paragraphs[0]
                paragraph.text = title_lines[0] or "\u200b"
                lines = [paragraph]
                for line in title_lines[1:]:
                    continuation = cell.add_paragraph(style="CoverValueLine")
                    continuation.add_run(line)
                    lines.append(continuation)
                for line_index, title_paragraph in enumerate(lines):
                    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    title_paragraph.paragraph_format.space_before = Pt(0)
                    title_paragraph.paragraph_format.space_after = Pt(0)
                    title_paragraph.paragraph_format.line_spacing = Pt(24)
                    title_paragraph.paragraph_format.keep_together = True
                    _set_paragraph_bottom_border(title_paragraph)
                    _set_paragraph_between_border(title_paragraph)
                    run = title_paragraph.runs[0]
                    if line_index == 0:
                        run.style = "CoverValue"
                    _set_run_fonts(
                        run, "SimSun", "Times New Roman", 14, bold=False, underline=False
                    )
                continue
            for run in cell.paragraphs[0].runs:
                if kind == "label":
                    _set_run_fonts(run, "SimSun", "SimSun", 14, bold=True, underline=False)
                else:
                    _set_run_fonts(run, "SimSun", "Times New Roman", 14, bold=False, underline=False)

    # Keep every cover block free of inherited first-line indents.
    for paragraph in (logo_paragraph, title, seal_paragraph, fields_gap, cover_footer):
        indent = paragraph._p.get_or_add_pPr().get_or_add_ind()
        for attribute in ("left", "leftChars", "firstLine", "firstLineChars"):
            indent.set(qn(f"w:{attribute}"), "0")


def _append_proposal_table(
    document: Document,
    fields: tuple[str, ...],
    field_nodes: dict[str, list[object]],
    *,
    heading_abstract_id: int,
    advisor_signature: Path | None = None,
    signature_background: str = "preserve",
    signature_date: str = "",
) -> None:
    labels = {
        "significance": "选题意义",
        "research-status": "国内外研究现状概述",
        "research-content": "主要研究内容",
        "research-approach": "拟采用的研究思路",
        "schedule": "研究工作安排及进度",
        "references": "参考文献目录",
        "advisor-opinion": "指导教师意见",
    }
    minimum_heights = {
        "significance": 165.15,
        "research-status": 297.15,
        "research-content": 183.15,
        "schedule": 140.1,
        "references": 383.6,
        "advisor-opinion": 99.9,
    }
    # The long research-approach field uses its content height to avoid a blank page tail.
    table = document.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_fixed_table_geometry(
        table,
        (1.52, 14.23),
        style_name=STYLE_PROPOSAL_FORM_TABLE,
    )
    table_indent = table._tbl.tblPr.find(qn("w:tblInd"))
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "249")
    _set_form_table_borders(table)
    for row_index, key in enumerate(fields):
        row = table.rows[row_index]
        if key in minimum_heights:
            row.height = Pt(minimum_heights[key])
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row_properties = row._tr.get_or_add_trPr()
        cant_split = row_properties.find(qn("w:cantSplit"))
        if cant_split is not None:
            row_properties.remove(cant_split)
        label_cell, body_cell = row.cells
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        body_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_form_cell_margins(label_cell, left=40, right=40)
        _set_form_cell_margins(body_cell)
        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.style = document.styles[STYLE_PROPOSAL_LABEL]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_paragraph.add_run("\n".join(labels[key]))
        _move_nodes_to_cell(body_cell, field_nodes[key], document)
        _fit_nested_tables_to_cell(body_cell)
        _restart_cell_heading_numbering(document, body_cell, heading_abstract_id)
        if key == "advisor-opinion":
            signature = body_cell.add_paragraph(style=STYLE_PROPOSAL_SIGNATURE)
            signature.add_run("签名： ")
            if advisor_signature is None:
                signature.add_run("        ")
            else:
                picture_source: str | io.BytesIO = str(advisor_signature)
                if signature_background == "whiten":
                    picture_source = whiten_signature_image(advisor_signature)
                picture_run = signature.add_run()
                picture_run.add_picture(picture_source, width=Cm(2.6))
                for doc_property in picture_run._r.iter(qn("wp:docPr")):
                    doc_property.attrib.pop("descr", None)
            signature.add_run(
                f"    {signature_date}" if signature_date else "    年    月    日"
            )


def _configure_proposal_heading_styles(document: Document, discipline: str) -> None:
    """Apply the field outline without thesis chapter pagination."""
    for name in (STYLE_HEADING_1, STYLE_HEADING_2, STYLE_HEADING_3, STYLE_HEADING_4):
        style = document.styles[name]
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.page_break_before = False
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        indent = style.element.get_or_add_pPr().get_or_add_ind()
        indent.set(qn("w:firstLineChars"), "0" if discipline == "science" else "200")
        indent.set(qn("w:firstLine"), "0" if discipline == "science" else str(round(style.font.size.pt * 40)))


def _postprocess_proposal_docx(
    input_docx: Path,
    output_docx: Path,
    *,
    metadata: dict[str, str],
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    discipline: str,
    citation_mode: str,
    labels: LabelRegistry,
    citations: CitationRegistry,
    bibliography_items: list[dict[str, object]],
    table_layouts: list[LatexTableLayout],
    word_figure_sequence: str,
    word_table_sequence: str,
) -> None:
    document = Document(input_docx)
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith("YIBIN_INTERNAL_CITATION_SEED"):
            paragraph._p.getparent().remove(paragraph._p)
    heading_abstract_id = _apply_common_body_formatting(
        document,
        discipline=discipline,
        labels=labels,
        citations=citations,
        citation_mode=citation_mode,
        table_layouts=table_layouts,
        word_figure_sequence=word_figure_sequence,
        word_table_sequence=word_table_sequence,
    )
    _configure_proposal_heading_styles(document, discipline)
    field_nodes = _extract_marker_ranges(document)
    signature_background = metadata.get("signature-background", "preserve").strip().casefold()
    if signature_background not in {"preserve", "whiten"}:
        raise BuildError(
            "元数据 signature-background 只能是 preserve 或 whiten："
            + signature_background
        )
    advisor_signature = _resolve_signature_asset(
        metadata,
        "advisor-signature",
        main_dir,
        metadata_dir,
    )

    _insert_proposal_cover(
        document,
        metadata,
        project_root=project_root,
        main_dir=main_dir,
        metadata_dir=metadata_dir,
    )
    # The standalone opening report keeps the cover in its own unnumbered
    # section, so the form itself starts at page 1. Keep the form's existing
    # marker content after the break so long fields can still flow naturally.
    document.add_section(WD_SECTION_START.NEW_PAGE)

    title = document.add_paragraph(style=STYLE_PROPOSAL_TITLE)
    title.text = "宜宾学院本科毕业论文（设计）开题报告"
    subtitle = document.add_paragraph(style=STYLE_PROPOSAL_SUBTITLE)
    subtitle.text = "（学生填写）"
    _append_proposal_table(
        document,
        PROPOSAL_FIELD_ORDER,
        field_nodes,
        heading_abstract_id=heading_abstract_id,
        advisor_signature=advisor_signature,
        signature_background=signature_background,
        signature_date=metadata.get("advisor-signature-date", "").strip(),
    )

    if len(document.sections) != 2:
        raise BuildError(f"开题报告应生成 2 个 Word 分节，实际为 {len(document.sections)}。")
    cover, section = document.sections
    cover.start_type = WD_SECTION_START.NEW_PAGE
    cover.page_width = Cm(21)
    cover.page_height = Cm(29.7)
    cover.top_margin = Pt(70.9)
    cover.bottom_margin = Pt(51.05)
    cover.left_margin = Pt(85.05)
    cover.right_margin = Pt(62.35)
    cover.header_distance = Pt(36.85)
    cover.footer_distance = Pt(34)
    cover.header.is_linked_to_previous = False
    cover.footer.is_linked_to_previous = False
    _clear_story(cover.header)
    _clear_story(cover.footer)
    _set_page_number_format(cover, None)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Pt(70.9)
    section.bottom_margin = Pt(51.05)
    section.left_margin = Pt(85.05)
    section.right_margin = Pt(62.35)
    section.header_distance = Pt(36.85)
    section.footer_distance = Pt(34)
    section.header.is_linked_to_previous = False
    _clear_story(section.header)
    _set_page_number_format(section, "decimal", 1)
    _add_profile_page_field(section, document)

    document.core_properties.title = metadata.get("title", "")
    document.core_properties.author = metadata.get("author", "")
    document.core_properties.subject = "宜宾学院本科毕业论文（设计）开题报告"
    document.core_properties.keywords = "YibinThesis;document-type=proposal;template-year=2024"
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    _localize_bibliography_terms(document, citations, bibliography_items)
    document.save(output_docx)
    if citation_mode == "native":
        _inject_bibliography_sources(output_docx, citations, bibliography_items)
    _normalize_embedded_pngs(output_docx)
    Document(output_docx)


def _set_single_cell_paragraph(cell, document: Document, style_name: str, value: str):
    """Replace merge residue with one style-driven paragraph."""

    paragraph = cell.paragraphs[0]
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.style = document.styles[style_name]
    if value:
        paragraph.add_run(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return paragraph


def _insert_review_cover_field(
    document: Document,
    anchor,
    label: str,
    value: str,
    *,
    second_label: str = "",
    second_value: str = "",
    wrap_value: bool = False,
) -> None:
    """Use native paragraph tabs; underline only the editable value slots."""
    paragraph = _insert_paragraph_before(document, anchor, STYLE_REVIEW_COVER_FIELD)
    stops = ((3.7, "left"), (15.5, "right")) if wrap_value else (
        ((3.7, "left"), (6.35, "center"), (9.0, "right"),
         (9.3, "left"), (10.9, "left"), (13.2, "center"), (15.5, "right"))
        if second_label
        else ((3.7, "left"), (9.6, "center"), (15.5, "right"))
    )
    if wrap_value:
        paragraph.paragraph_format.left_indent = Cm(3.7)
        paragraph.paragraph_format.first_line_indent = -Cm(3.7)
    _set_tabs(paragraph, [(Cm(position).twips, alignment, None) for position, alignment in stops])
    label_run = paragraph.add_run(label + "\t")
    label_run.style = document.styles[STYLE_REVIEW_COVER_LABEL]
    value_run = paragraph.add_run(("" if wrap_value else "\t") + value.strip() + "\t")
    value_run.style = document.styles[STYLE_REVIEW_COVER_VALUE]
    if second_label:
        second_label_run = paragraph.add_run("\t" + second_label + "\t")
        second_label_run.style = document.styles[STYLE_REVIEW_COVER_LABEL]
        second_value_run = paragraph.add_run("\t" + second_value.strip() + "\t")
        second_value_run.style = document.styles[STYLE_REVIEW_COVER_VALUE]


def _insert_review_cover(
    document: Document,
    anchor,
    metadata: dict[str, str],
    *,
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    discipline: str,
) -> None:
    try:
        logo = _resolve_cover_logo(
            metadata,
            project_root,
            main_dir,
            metadata_dir,
            allow_project_fallback=True,
        )
    except BuildError:
        logo = None
    logo_paragraph = _insert_paragraph_before(document, anchor, "CoverLogo")
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.space_before = Pt(0)
    logo_paragraph.paragraph_format.space_after = Pt(12)
    if logo is not None:
        logo_paragraph.add_run().add_picture(str(logo), width=Cm(13.15), height=Cm(3.65))
    heading = _insert_paragraph_before(document, anchor, STYLE_REVIEW_DOCUMENT_TITLE)
    heading.text = (
        "本科生毕业论文（设计）文献综述"
        if discipline == "science"
        else "本科生毕业论文文献综述"
    )
    rows = (
        ("题    目", metadata.get("title", "")),
        ("学部（院）", metadata.get("college", "")),
        ("专    业", metadata.get("major", "")),
        ("学生姓名", metadata.get("author", "")),
    )
    for label, value in rows:
        _insert_review_cover_field(document, anchor, label, value, wrap_value=label == "题    目")
    _insert_review_cover_field(
        document, anchor, "学    号", metadata.get("student-id", ""),
        second_label="年级", second_value=format_grade_class(metadata),
    )
    _insert_review_cover_field(
        document, anchor, "指导教师1", metadata.get("advisor", ""),
        second_label="职称", second_value=metadata.get("advisor-title", ""),
    )
    _insert_review_cover_field(
        document, anchor, "指导教师2",
        metadata.get("second-advisor", metadata.get("external-advisor", "")),
        second_label="职称",
        second_value=metadata.get("second-advisor-title", metadata.get("external-advisor-title", "")),
    )
    date = _insert_paragraph_before(document, anchor, STYLE_REVIEW_DATE)
    date.text = metadata.get("date", "")


def _insert_review_front_matter(document: Document, anchor, metadata: dict[str, str]) -> None:
    affiliation = " ".join(
        metadata.get(key, "").strip()
        for key in ("college", "affiliation-city", "affiliation-postcode")
        if metadata.get(key, "").strip()
    )
    for style, text in (
        (STYLE_REVIEW_THESIS_TITLE, metadata.get("title", "")),
        (STYLE_REVIEW_AUTHOR, metadata.get("author", "")),
        (STYLE_REVIEW_AFFILIATION, f"（{affiliation}）" if affiliation else ""),
    ):
        _insert_paragraph_before(document, anchor, style).text = text


def _apply_review_styles(document: Document, discipline: str) -> None:
    """Keep the review's abstract page and flowing outline separate from thesis styles."""
    style_map = {
        STYLE_BODY: STYLE_REVIEW_BODY,
        STYLE_FIRST_PARAGRAPH: STYLE_REVIEW_BODY,
        STYLE_HEADING_1: STYLE_REVIEW_HEADING_1,
        STYLE_HEADING_2: STYLE_REVIEW_HEADING_2,
        STYLE_HEADING_3: STYLE_REVIEW_HEADING_3,
        STYLE_HEADING_4: STYLE_REVIEW_HEADING_4,
        STYLE_UNNUMBERED_HEADING: STYLE_REVIEW_UNNUMBERED_HEADING,
        STYLE_FRONT_TITLE: STYLE_REVIEW_ABSTRACT_TITLE,
        STYLE_ENGLISH_ABSTRACT_TITLE: STYLE_REVIEW_ENGLISH_ABSTRACT_TITLE,
        STYLE_CHINESE_ABSTRACT: STYLE_REVIEW_CHINESE_ABSTRACT,
        STYLE_ENGLISH_ABSTRACT: STYLE_REVIEW_ENGLISH_ABSTRACT,
        STYLE_KEYWORDS: STYLE_REVIEW_KEYWORDS,
        STYLE_BIBLIOGRAPHY: STYLE_REVIEW_BIBLIOGRAPHY,
    }
    for paragraph in document.paragraphs:
        target = style_map.get(paragraph.style.name)
        if target == STYLE_REVIEW_KEYWORDS and paragraph.text.startswith("Keywords:"):
            target = STYLE_REVIEW_ENGLISH_KEYWORDS
        if target:
            paragraph.style = document.styles[target]
    heading_styles = (
        STYLE_REVIEW_HEADING_1, STYLE_REVIEW_HEADING_2,
        STYLE_REVIEW_HEADING_3, STYLE_REVIEW_HEADING_4,
    )
    for style_name in heading_styles:
        style = document.styles[style_name]
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.page_break_before = False
        indent = style.element.get_or_add_pPr().get_or_add_ind()
        indent.set(qn("w:firstLineChars"), "0" if discipline == "science" else "200")
        indent.set(qn("w:firstLine"), "0" if discipline == "science" else str(round(style.font.size.pt * 40)))
    _configure_heading_numbering(document, discipline, style_names=heading_styles)


def _postprocess_review_docx(
    input_docx: Path,
    output_docx: Path,
    *,
    metadata: dict[str, str],
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    discipline: str,
    citation_mode: str,
    labels: LabelRegistry,
    citations: CitationRegistry,
    bibliography_items: list[dict[str, object]],
    table_layouts: list[LatexTableLayout],
    word_figure_sequence: str,
    word_table_sequence: str,
) -> None:
    document = Document(input_docx)
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith("YIBIN_INTERNAL_CITATION_SEED"):
            paragraph._p.getparent().remove(paragraph._p)
    marker = next(
        (p for p in document.paragraphs if p.text.strip() == SECTION_REVIEW_COVER_END),
        None,
    )
    if marker is None:
        raise BuildError("Word 文献综述信息页分节标记缺失。")
    body_sect_pr = document._element.body.sectPr
    if body_sect_pr is None:
        raise BuildError("Pandoc 输出缺少正文分节属性。")
    references_marker = next(
        (p for p in document.paragraphs if p.text.strip() == SECTION_REVIEW_REFERENCES),
        None,
    )
    abstract_marker = next(
        (p for p in document.paragraphs if p.text.strip() == SECTION_REVIEW_ABSTRACT_END),
        None,
    )
    if references_marker is None or abstract_marker is None:
        raise BuildError("文献综述摘要或参考文献分节标记缺失。")
    paragraphs = document.paragraphs
    first_abstract = paragraphs[next(i for i, p in enumerate(paragraphs) if p._p is marker._p) + 1]
    _insert_review_cover(
        document,
        marker,
        metadata,
        project_root=project_root,
        main_dir=main_dir,
        metadata_dir=metadata_dir,
        discipline=discipline,
    )
    _insert_review_front_matter(document, first_abstract, metadata)
    _add_section_break(marker, body_sect_pr, None)
    _add_section_break(abstract_marker, body_sect_pr, None)
    _add_section_break(references_marker, body_sect_pr, None)
    _apply_common_body_formatting(
        document,
        discipline=discipline,
        labels=labels,
        citations=citations,
        citation_mode=citation_mode,
        table_layouts=table_layouts,
        word_figure_sequence=word_figure_sequence,
        word_table_sequence=word_table_sequence,
    )
    _apply_review_styles(document, discipline)
    stage = output_docx.with_suffix(".review-sections.docx")
    document.save(stage)
    document = Document(stage)
    if len(document.sections) != 4:
        stage.unlink(missing_ok=True)
        raise BuildError(f"文献综述应生成 4 个 Word 分节，实际为 {len(document.sections)}。")
    cover, abstracts, body, references = document.sections
    for section in (cover, abstracts, body, references):
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        _clear_story(section.header)
        _clear_story(section.footer)
    _set_page_number_format(cover, None)
    _set_page_number_format(abstracts, "decimal", 1)
    _set_page_number_format(body, "decimal")
    _set_page_number_format(references, "decimal")
    body_header = body.header.paragraphs[0]
    body_header.style = document.styles["Header"]
    body_header.text = metadata.get("title", "")
    body_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_header_border(body_header)
    for run in body_header.runs:
        _set_run_fonts(run, "SimSun", "Times New Roman", 9)
    _add_profile_page_field(abstracts, document)
    _add_profile_page_field(body, document)
    _add_profile_page_field(references, document)

    document.core_properties.title = metadata.get("title", "")
    document.core_properties.author = metadata.get("author", "")
    document.core_properties.subject = "宜宾学院本科毕业论文（设计）文献综述"
    document.core_properties.keywords = "YibinThesis;document-type=literature-review;template-year=2024"
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    _localize_bibliography_terms(document, citations, bibliography_items)
    document.save(output_docx)
    if citation_mode == "native":
        _inject_bibliography_sources(output_docx, citations, bibliography_items)
    _normalize_embedded_pngs(output_docx)
    Document(output_docx)
    stage.unlink(missing_ok=True)


def _build_non_thesis(
    args: argparse.Namespace,
    *,
    profile: DocumentProfile,
    project_root: Path,
    main_path: Path,
    main_text: str,
    main_dir: Path,
    metadata_path: Path,
    metadata: dict[str, str],
    pandoc: Path,
    reference_doc: Path,
    csl: Path,
    bibliography: Path,
    output: Path,
    allow_project_fallback: bool,
) -> Path:
    warnings: list[str] = []
    notes = NoteRegistry()
    counters = HeadingCounters()
    float_counters = FloatCounters()
    labels = LabelRegistry()
    citations = CitationRegistry()
    table_layouts: list[LatexTableLayout] = []
    markdown_parts: list[str] = []

    graphic_paths = collect_graphic_paths(
        main_text,
        main_dir=main_dir,
        project_root=project_root,
        allow_project_fallback=allow_project_fallback,
    )
    resource_candidates = [
        main_dir.resolve(),
        (main_dir / "assets").resolve(),
        (main_dir / "chapters").resolve(),
        *graphic_paths,
    ]
    if allow_project_fallback:
        resource_candidates.extend(
            [project_root.resolve(), (project_root / "assets").resolve()]
        )
    resource_candidates = list(dict.fromkeys(resource_candidates))

    def image_target_resolver(target: str) -> str:
        return resolve_word_image_target(target, resource_candidates)

    expanded_main = expand_nested_inputs(
        main_path,
        project_root,
        allow_project_fallback=allow_project_fallback,
    )
    temporary_root = output.parent / ".tmp"
    temporary_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=f"yibinthesis-{profile.document_type}-word-",
        dir=temporary_root,
    ) as temporary:
        temp_dir = Path(temporary)
        if profile.document_type == "proposal":
            if "\\makeyibinproposal" not in strip_tex_comments(main_text):
                raise BuildError("开题报告入口未调用 \\makeyibinproposal。")
            fields = extract_proposal_fields(expanded_main)
            main_nocites = [
                event.value
                for event in parse_main_events(main_text)
                if event.kind == "nocite" and event.value
            ]
            bibliography_requested = False
            for index, (key, field_text) in enumerate(fields):
                bibliography_requested = bibliography_requested or (
                    "\\printyibinproposalbibliography" in field_text
                )
                field_text = field_text.replace("\\printyibinproposalbibliography", "")
                field_text = normalize_proposal_latex(field_text)
                detect_unsupported(field_text, main_path, warnings)
                # Each official form field is an independent outline.  A
                # ``\section`` inside every cell is level one and starts at
                # 一、 or 1 instead of continuing from the previous field.
                field_counters = HeadingCounters()
                chapter_before = field_counters.chapter
                converted = latex_to_markdown(
                    normalize_latex(
                        field_text,
                        notes,
                        profile.discipline,
                        labels,
                        citations,
                        table_layouts,
                    ),
                    pandoc=pandoc,
                    cwd=main_dir,
                    temp_dir=temp_dir,
                    name=f"proposal-{index}-{key}",
                )
                converted = normalize_proposal_headings(converted)
                converted = apply_heading_numbering(
                    converted,
                    profile.discipline,
                    field_counters,
                    labels,
                )
                converted = apply_float_numbering(
                    converted,
                    # Figures and tables retain their document-wide sequence;
                    # only the field's heading outline restarts.
                    "humanities",
                    chapter_before,
                    float_counters,
                    labels,
                    image_target_resolver,
                )
                markdown_parts.append(
                    styled_marker(PROPOSAL_FIELD_MARKER_PREFIX + key.upper())
                )
                markdown_parts.append(converted)
                if key == "references" and bibliography_requested:
                    for keys in main_nocites:
                        citations.add_nocite(keys)
                    seed = citation_seed_markdown(citations)
                    if seed:
                        markdown_parts.append(seed)
                    markdown_parts.append(bibliography_markdown())
        else:
            if "\\makeyibinliteraturereviewcover" not in strip_tex_comments(main_text):
                raise BuildError(
                    "文献综述入口未调用 \\makeyibinliteraturereviewcover。"
                )
            markdown_parts.append(styled_marker(SECTION_REVIEW_COVER_END))
            events = parse_main_events(main_text)
            bibliography_inserted = False
            expected_abstracts = ["cnabstract", "enabstract"]
            abstracts_seen: list[str] = []
            body_started = False
            for index, event in enumerate(events):
                if event.kind in {
                    "frontmatter",
                    "mainmatter",
                    "backmatter",
                    "makeyibinliteraturereviewcover",
                    "makeyibinproposal",
                }:
                    continue
                if event.kind in {"input", "include"} and event.value:
                    source = resolve_source(
                        event.value,
                        main_dir,
                        project_root,
                        allow_project_fallback=allow_project_fallback,
                    )
                    if source.resolve() == metadata_path.resolve():
                        continue
                    text = expand_nested_inputs(
                        source,
                        project_root,
                        allow_project_fallback=allow_project_fallback,
                    )
                    # An explicit --metadata replaces the values, not the
                    # original setup input's role in the source entry point.
                    if find_command_argument(strip_tex_comments(text), "yibinsetup") is not None:
                        continue
                    environments = [
                        environment for environment in expected_abstracts
                        if extract_environment(text, environment) is not None
                    ]
                    if environments:
                        pending = abstracts_seen + environments
                        if body_started or pending != expected_abstracts[:len(pending)]:
                            raise BuildError("文献综述须按中文摘要、英文摘要、正文的顺序组织，摘要不能重复。")
                        remainder = strip_tex_comments(text)
                        for environment in environments:
                            pattern = r"\\begin\s*\{" + environment + r"\}.*?\\end\s*\{" + environment + r"\}"
                            remainder, count = re.subn(pattern, "", remainder, flags=re.DOTALL)
                            if count != 1:
                                raise BuildError("文献综述摘要文件不能重复同一摘要环境。")
                        if remainder.strip():
                            raise BuildError("文献综述摘要文件不能混入正文；请把摘要环境以外的内容放入独立正文文件。")
                        for environment in environments:
                            markdown_parts.append(abstract_markdown(
                                text, environment=environment, pandoc=pandoc,
                                cwd=main_dir, temp_dir=temp_dir, notes=notes,
                                discipline=profile.discipline, labels=labels,
                                citations=citations, table_layouts=table_layouts,
                            ))
                        abstracts_seen = pending
                        continue
                    if not body_started:
                        if abstracts_seen != expected_abstracts:
                            raise BuildError("文献综述正文前须提供 cnabstract 和 enabstract 中英文摘要。")
                        markdown_parts.append(styled_marker(SECTION_REVIEW_ABSTRACT_END))
                        body_started = True
                    detect_unsupported(text, source, warnings)
                    chapter_before = counters.chapter
                    converted = latex_to_markdown(
                        normalize_latex(
                            text,
                            notes,
                            profile.discipline,
                            labels,
                            citations,
                            table_layouts,
                        ),
                        pandoc=pandoc,
                        cwd=main_dir,
                        temp_dir=temp_dir,
                        name=f"review-{index}",
                    )
                    converted = apply_heading_numbering(
                        converted,
                        profile.discipline,
                        counters,
                        labels,
                        unnumbered_style=STYLE_REVIEW_UNNUMBERED_HEADING,
                    )
                    converted = apply_float_numbering(
                        converted,
                        profile.discipline,
                        chapter_before,
                        float_counters,
                        labels,
                        image_target_resolver,
                    )
                    markdown_parts.append(converted)
                elif event.kind == "printyibinnotes":
                    rendered_notes = notes_markdown(notes)
                    if rendered_notes:
                        markdown_parts.append(rendered_notes)
                elif event.kind == "nocite" and event.value:
                    citations.add_nocite(event.value)
                elif event.kind == "printyibinbibliography":
                    if not body_started:
                        raise BuildError("文献综述参考文献前须提供中英文摘要和正文。")
                    seed = citation_seed_markdown(citations)
                    if seed:
                        markdown_parts.append(seed)
                    markdown_parts.append(styled_marker(SECTION_REVIEW_REFERENCES))
                    markdown_parts.append(bibliography_markdown())
                    bibliography_inserted = True
            if not body_started:
                raise BuildError("文献综述须提供中英文摘要和正文文件。")
            if not bibliography_inserted:
                seed = citation_seed_markdown(citations)
                if seed:
                    markdown_parts.append(seed)
                markdown_parts.append(styled_marker(SECTION_REVIEW_REFERENCES))
                markdown_parts.append(bibliography_markdown())

        assembled = "\n\n".join(part for part in markdown_parts if part.strip()) + "\n"
        assembled = labels.resolve(assembled)
        markdown_file = temp_dir / "assembled.md"
        markdown_file.write_text(assembled, encoding="utf-8", newline="\n")
        if args.keep_intermediate:
            kept = output.with_suffix(".pandoc.md")
            kept.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(markdown_file, kept)

        class_options = parse_document_class_options(main_text)
        citeproc_bibliography = prepare_bibliography(
            bibliography,
            pandoc=pandoc,
            cwd=main_dir,
            temp_dir=temp_dir,
            include_url="no-url" not in class_options,
            include_doi="no-doi" not in class_options,
        )
        bibliography_items: list[dict[str, object]] = []
        bibliography_json = citeproc_bibliography.with_suffix(".json") if citeproc_bibliography != bibliography else citeproc_bibliography
        if bibliography_json.suffix.lower() == ".json":
            parsed_items = json.loads(bibliography_json.read_text(encoding="utf-8"))
            if isinstance(parsed_items, list):
                bibliography_items = [
                    item for item in parsed_items if isinstance(item, dict)
                ]
        source_items_path = bibliography_source_items_path(citeproc_bibliography)
        if source_items_path.is_file():
            parsed_items = json.loads(source_items_path.read_text(encoding="utf-8"))
            if isinstance(parsed_items, list):
                bibliography_items = [
                    item for item in parsed_items if isinstance(item, dict)
                ]
        pandoc_docx = temp_dir / "pandoc.docx"
        resource_paths = os.pathsep.join(str(path) for path in resource_candidates)
        command = [
            str(pandoc),
            "--from=markdown+smart+raw_attribute+fenced_divs+bracketed_spans+citations+tex_math_dollars",
            "--to=docx",
            "--standalone",
            "--citeproc",
            f"--reference-doc={reference_doc}",
            f"--bibliography={citeproc_bibliography}",
            f"--csl={csl}",
            f"--resource-path={resource_paths}",
            "--metadata=lang:zh-CN",
            f"--output={pandoc_docx}",
            str(markdown_file),
        ]
        run_checked(command, cwd=main_dir)
        _normalize_embedded_pngs(pandoc_docx)
        if profile.document_type == "proposal":
            _postprocess_proposal_docx(
                pandoc_docx,
                output,
                metadata=metadata,
                project_root=project_root,
                main_dir=main_dir,
                metadata_dir=metadata_path.parent,
                discipline=profile.discipline,
                citation_mode=args.citation_mode,
                labels=labels,
                citations=citations,
                bibliography_items=bibliography_items,
                table_layouts=table_layouts,
                word_figure_sequence=args.word_figure_sequence,
                word_table_sequence=args.word_table_sequence,
            )
        else:
            _postprocess_review_docx(
                pandoc_docx,
                output,
                metadata=metadata,
                project_root=project_root,
                main_dir=main_dir,
                metadata_dir=metadata_path.parent,
                discipline=profile.discipline,
                citation_mode=args.citation_mode,
                labels=labels,
                citations=citations,
                bibliography_items=bibliography_items,
                table_layouts=table_layouts,
                word_figure_sequence=args.word_figure_sequence,
                word_table_sequence=args.word_table_sequence,
            )

    print(f"Generated: {output}")
    print("Word 输出可编辑；字段将在 Microsoft Word 打开或刷新时更新。")
    for warning in dict.fromkeys(warnings):
        print(f"WARNING: {warning}", file=sys.stderr)
    return output
