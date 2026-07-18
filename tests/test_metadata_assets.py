from __future__ import annotations

import base64
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "yibinthesis_build_word_metadata_assets", ROOT / "tools" / "build_word.py"
)
assert SPEC is not None and SPEC.loader is not None
BUILD_WORD = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = BUILD_WORD
SPEC.loader.exec_module(BUILD_WORD)

ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class MetadataAssetResolutionTests(unittest.TestCase):
    def test_builtin_logo_is_available_to_external_manuscripts(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            external = Path(temporary)
            self.assertEqual(
                BUILD_WORD._resolve_cover_logo(
                    {"logo": "builtin"},
                    ROOT,
                    external,
                    external,
                    allow_project_fallback=False,
                ),
                (ROOT / "assets" / "yibin-university-logo.png").resolve(),
            )

    def test_signature_is_resolved_from_metadata_directory_only(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            metadata_dir = root / "paper"
            main_dir = root / "other"
            metadata_dir.mkdir()
            main_dir.mkdir()
            signature = metadata_dir / "signature.png"
            signature.write_bytes(ONE_PIXEL_PNG)
            self.assertEqual(
                BUILD_WORD._resolve_signature_asset(
                    {"author-signature": "signature.png"},
                    "author-signature",
                    main_dir,
                    metadata_dir,
                ),
                signature.resolve(),
            )

    def test_empty_advisor_signature_is_a_clean_empty_slot(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            self.assertIsNone(
                BUILD_WORD._resolve_signature_asset(
                    {"advisor-signature": ""},
                    "advisor-signature",
                    root,
                    root,
                )
            )

    def test_declared_missing_signature_fails_loudly(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with self.assertRaises(BUILD_WORD.BuildError):
                BUILD_WORD._resolve_signature_asset(
                    {"author-signature": "missing.png"},
                    "author-signature",
                    root,
                    root,
                )


class DeclarationSignatureTests(unittest.TestCase):
    def test_author_signature_is_embedded_twice_and_advisor_remains_empty(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            signature = root / "signature.png"
            signature.write_bytes(ONE_PIXEL_PNG)
            document = Document(ROOT / "word" / "reference.docx")
            anchor = document.paragraphs[0]
            metadata = {"title": "测试论文", "secrecy": "public"}

            BUILD_WORD._build_originality_page(
                document,
                anchor,
                metadata,
                signature,
                "preserve",
            )
            BUILD_WORD._build_authorization_page(
                document,
                anchor,
                metadata,
                signature,
                None,
                "preserve",
            )

            self.assertEqual(len(document.inline_shapes), 2)
            authorization_signature_table = document.tables[-2]
            self.assertTrue(
                list(
                    authorization_signature_table.cell(0, 1)._tc.iter(
                        BUILD_WORD.qn("w:drawing")
                    )
                )
            )
            self.assertFalse(
                list(
                    authorization_signature_table.cell(0, 4)._tc.iter(
                        BUILD_WORD.qn("w:drawing")
                    )
                )
            )
            self.assertFalse(
                list(
                    authorization_signature_table.cell(0, 4)._tc.iter(
                        BUILD_WORD.qn("w:u")
                    )
                )
            )
            for cell_index in (1, 4):
                signature_cell = authorization_signature_table.cell(0, cell_index)
                self.assertFalse(
                    list(signature_cell._tc.iter(BUILD_WORD.qn("w:tcBorders")))
                )
                self.assertEqual(
                    len(
                        list(
                            signature_cell.paragraphs[0]._p.iter(
                                BUILD_WORD.qn("w:pBdr")
                            )
                        )
                    ),
                    1,
                )

    def test_latex_metadata_contract_exposes_both_signature_slots(self) -> None:
        source = (ROOT / "yibinthesis.cls").read_text(encoding="utf-8")
        self.assertIn("author-signature", source)
        self.assertIn("advisor-signature", source)
        self.assertIn(r"\yibinauthorsignature", source)
        self.assertIn(r"\yibinadvisorsignature", source)
        self.assertIn("logo                    = {builtin}", source)

    def test_signature_whitening_preserves_source_and_writes_white_background(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "signature.png"
            output = root / "signature-whitened.png"
            source.write_bytes(ONE_PIXEL_PNG)
            original = source.read_bytes()

            stream = BUILD_WORD.whiten_signature_image(source, output)

            self.assertEqual(source.read_bytes(), original)
            self.assertTrue(output.is_file())
            self.assertGreater(len(stream.getvalue()), 0)


if __name__ == "__main__":
    unittest.main()
