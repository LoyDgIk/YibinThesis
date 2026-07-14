#!/usr/bin/env python3
"""Generate the semantic Pandoc reference document for YibinThesis.

The output is a derived, editable Word style package.  It does not copy the
content or package structure of the University's 2024 Word example.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SIMSUN = "SimSun"
SIMHEI = "SimHei"
KAITI = "KaiTi"
TIMES = "Times New Roman"


def _get_or_add_style(document: Document, name: str, style_type: WD_STYLE_TYPE):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def _set_fonts(
    style,
    *,
    east_asia: str,
    latin: str = TIMES,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    style.font.name = latin
    style.font.bold = bold
    style.font.italic = italic
    if size is not None:
        style.font.size = Pt(size)

    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:hint"), "eastAsia")

    color = r_pr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        r_pr.append(color)
    color.set(qn("w:val"), "000000")
    for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
        color.attrib.pop(qn(attribute), None)

    if bold is not None:
        for tag in ("w:b", "w:bCs"):
            node = r_pr.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                r_pr.append(node)
            node.set(qn("w:val"), "1" if bold else "0")


def _set_first_line_chars(style, chars: int = 200) -> None:
    p_pr = style.element.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))


def _set_outline_level(style, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def _set_keep(style, *, next_paragraph: bool = False, lines: bool = False) -> None:
    p_pr = style.element.get_or_add_pPr()
    for tag, enabled in (("w:keepNext", next_paragraph), ("w:keepLines", lines)):
        node = p_pr.find(qn(tag))
        if enabled and node is None:
            p_pr.append(OxmlElement(tag))
        elif not enabled and node is not None:
            p_pr.remove(node)


def _set_doc_defaults(document: Document) -> None:
    styles = document.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)

    r_pr_default = defaults.find(qn("w:rPrDefault"))
    if r_pr_default is None:
        r_pr_default = OxmlElement("w:rPrDefault")
        defaults.append(r_pr_default)
    r_pr = r_pr_default.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        r_pr_default.append(r_pr)
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr, value in (
        ("w:ascii", TIMES),
        ("w:hAnsi", TIMES),
        ("w:cs", TIMES),
        ("w:eastAsia", SIMSUN),
        ("w:hint", "eastAsia"),
    ):
        r_fonts.set(qn(attr), value)
    size = r_pr.find(qn("w:sz"))
    if size is None:
        size = OxmlElement("w:sz")
        r_pr.append(size)
    size.set(qn("w:val"), "24")
    size_cs = r_pr.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        r_pr.append(size_cs)
    size_cs.set(qn("w:val"), "24")

    p_pr_default = defaults.find(qn("w:pPrDefault"))
    if p_pr_default is None:
        p_pr_default = OxmlElement("w:pPrDefault")
        defaults.append(p_pr_default)
    p_pr = p_pr_default.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        p_pr_default.append(p_pr)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), "360")
    spacing.set(qn("w:lineRule"), "auto")


def _configure_paragraph_style(
    document: Document,
    name: str,
    *,
    east_asia: str = SIMSUN,
    latin: str = TIMES,
    size: float = 12,
    bold: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
    first_line: bool = False,
    line_spacing: float = 1.5,
    before: float = 0,
    after: float = 0,
    keep_next: bool = False,
    keep_lines: bool = False,
    page_break_before: bool = False,
    base: str | None = None,
    outline_level: int | None = None,
):
    style = _get_or_add_style(document, name, WD_STYLE_TYPE.PARAGRAPH)
    if base:
        style.base_style = document.styles[base]
    _set_fonts(
        style,
        east_asia=east_asia,
        latin=latin,
        size=size,
        bold=bold,
    )
    fmt = style.paragraph_format
    fmt.alignment = alignment
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.first_line_indent = Pt(24) if first_line else Pt(0)
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_lines
    fmt.page_break_before = page_break_before
    fmt.widow_control = True
    if first_line:
        _set_first_line_chars(style)
    else:
        # Block a two-character indent inherited from Normal/Heading styles.
        _set_first_line_chars(style, 0)
    if outline_level is not None:
        _set_outline_level(style, outline_level)
    _set_keep(style, next_paragraph=keep_next, lines=keep_lines)
    return style


def _configure_character_style(
    document: Document,
    name: str,
    *,
    east_asia: str,
    latin: str = TIMES,
    size: float,
    bold: bool = False,
):
    style = _get_or_add_style(document, name, WD_STYLE_TYPE.CHARACTER)
    _set_fonts(
        style,
        east_asia=east_asia,
        latin=latin,
        size=size,
        bold=bold,
    )
    return style


def build_reference_docx(output: Path) -> Path:
    document = Document()
    _set_doc_defaults(document)

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    _configure_paragraph_style(
        document,
        "Normal",
        first_line=True,
        line_spacing=1.5,
    )
    for name in ("Body Text", "First Paragraph"):
        _configure_paragraph_style(
            document,
            name,
            first_line=True,
            line_spacing=1.5,
            base="Normal",
        )
    _configure_paragraph_style(
        document,
        "Compact",
        first_line=True,
        line_spacing=1.5,
        base="Normal",
    )

    # The committed reference document defaults to the 2024 humanities
    # profile.  Science-specific chapter alignment/pagination is applied by
    # build_word.py during post-processing.
    heading_specs = (
        ("Heading 1", SIMHEI, 16, 12, 6, False, 0),
        ("Heading 2", KAITI, 15, 9, 6, False, 1),
        ("Heading 3", SIMSUN, 14, 6, 3, False, 2),
        ("Heading 4", SIMSUN, 14, 6, 3, False, 3),
    )
    for name, font, size, before, after, page_break, outline in heading_specs:
        _configure_paragraph_style(
            document,
            name,
            east_asia=font,
            size=size,
            bold=True,
            first_line=True,
            before=before,
            after=after,
            keep_next=True,
            keep_lines=True,
            page_break_before=page_break,
            outline_level=outline,
        )
        _configure_paragraph_style(
            document,
            name.replace("Heading", "Yibin Heading"),
            east_asia=font,
            size=size,
            bold=True,
            first_line=True,
            before=before,
            after=after,
            keep_next=True,
            keep_lines=True,
            page_break_before=page_break,
            base="Normal",
            outline_level=outline,
        )

    _configure_paragraph_style(
        document,
        "Title",
        east_asia=SIMHEI,
        size=18,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.5,
        after=12,
    )
    _configure_paragraph_style(
        document,
        "Subtitle",
        east_asia=SIMSUN,
        size=16,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.5,
    )
    _configure_paragraph_style(
        document,
        "Header",
        east_asia=SIMSUN,
        size=9,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
    )
    _configure_paragraph_style(
        document,
        "Footer",
        east_asia=SIMSUN,
        size=9,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
    )
    _configure_paragraph_style(
        document,
        "Caption",
        east_asia=SIMSUN,
        size=10.5,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        before=3,
        after=3,
        keep_next=False,
        keep_lines=True,
    )
    for name in ("Table Caption", "Image Caption", "Figure Caption"):
        _configure_paragraph_style(
            document,
            name,
            east_asia=SIMSUN,
            size=10.5,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            line_spacing=1.0,
            before=3,
            after=3,
            keep_next=name == "Table Caption",
            keep_lines=True,
            base="Caption",
        )
    _configure_paragraph_style(
        document,
        "Bibliography",
        east_asia=SIMSUN,
        size=10.5,
        line_spacing=1.5,
        before=0,
        after=0,
    )
    bibliography = document.styles["Bibliography"]
    bibliography.paragraph_format.left_indent = Cm(0.74)
    bibliography.paragraph_format.first_line_indent = Cm(-0.74)
    _configure_paragraph_style(
        document,
        "Footnote Text",
        east_asia=SIMSUN,
        size=9,
        line_spacing=1.0,
    )

    _configure_paragraph_style(
        document,
        "TOC Heading",
        east_asia=SIMHEI,
        size=16,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.5,
        after=12,
        keep_next=True,
    )
    for level in range(1, 4):
        toc = _configure_paragraph_style(
            document,
            f"TOC {level}",
            east_asia=SIMSUN,
            size=12,
            line_spacing=1.5,
            base="Normal",
        )
        toc.paragraph_format.first_line_indent = Pt(0)
        toc.paragraph_format.left_indent = Pt((level - 1) * 24)
        toc_ppr = toc.element.get_or_add_pPr()
        toc_indent = toc_ppr.find(qn("w:ind"))
        if toc_indent is None:
            toc_indent = OxmlElement("w:ind")
            toc_ppr.append(toc_indent)
        toc_indent.set(qn("w:leftChars"), str((level - 1) * 200))

    custom_paragraphs = (
        ("CoverLogo", SIMHEI, 24, True, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 12, False, False, None),
        ("CoverSchool", SIMHEI, 26, True, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 12, False, False, None),
        ("CoverThesisType", SIMHEI, 22, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 18, 18, False, False, None),
        ("CoverTitle", SIMHEI, 18, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 18, 24, False, False, None),
        ("CoverField", SIMSUN, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 3, 3, False, False, None),
        ("CoverDate", SIMSUN, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 24, 0, False, False, None),
        ("DeclarationTitle", SIMHEI, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0, 18, True, False, None),
        ("DeclarationBody", SIMSUN, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 0, 0, False, False, None),
        ("FrontTitle", SIMHEI, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0, 12, True, False, 0),
        ("Yibin TOC Heading", SIMHEI, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0, 12, True, False, None),
        ("ChineseAbstract", SIMSUN, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 0, 0, False, False, None),
        ("EnglishAbstract", TIMES, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 0, 0, False, False, None),
        ("Keywords", SIMSUN, 12, False, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 6, 0, False, False, None),
        ("Unnumbered Heading 1", SIMHEI, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0, 12, True, True, 0),
        ("Notes", SIMSUN, 9, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 0, 0, False, False, None),
        ("YibinSectionMarker", SIMSUN, 1, False, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0, False, False, None),
    )
    for (
        name,
        font,
        size,
        bold,
        alignment,
        line_spacing,
        before,
        after,
        keep_next,
        page_break,
        outline,
    ) in custom_paragraphs:
        _configure_paragraph_style(
            document,
            name,
            east_asia=font,
            latin=TIMES,
            size=size,
            bold=bold,
            alignment=alignment,
            line_spacing=line_spacing,
            before=before,
            after=after,
            keep_next=keep_next,
            keep_lines=keep_next,
            page_break_before=page_break,
            outline_level=outline,
        )
    _set_first_line_chars(document.styles["DeclarationBody"])
    _set_first_line_chars(document.styles["ChineseAbstract"])

    _configure_character_style(
        document,
        "CoverLabel",
        east_asia=SIMHEI,
        size=18,
        bold=True,
    )
    _configure_character_style(
        document,
        "CoverValue",
        east_asia=SIMSUN,
        size=16,
        bold=True,
    )
    _configure_character_style(
        document,
        "KeywordLabel",
        east_asia=SIMHEI,
        size=12,
        bold=True,
    )

    marker = document.styles["YibinSectionMarker"]
    marker.font.hidden = True

    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.core_properties.title = "YibinThesis Pandoc reference document"
    document.core_properties.subject = "宜宾学院本科毕业论文非官方 LaTeX 模板的 Word 样式"
    document.core_properties.author = "YibinThesis contributors"
    document.core_properties.comments = (
        "Semantic reference.docx derived from public school formatting evidence; "
        "not an official University template."
    )

    # Pandoc imports styles and section settings, not the reference body content.
    if document.paragraphs:
        document.paragraphs[0].text = ""
    else:
        document.add_paragraph("")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main() -> int:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=project_root / "word" / "reference.docx",
        help="Output DOCX path (default: word/reference.docx).",
    )
    args = parser.parse_args()
    output = args.output.expanduser().resolve()
    build_reference_docx(output)
    print(f"Generated: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
