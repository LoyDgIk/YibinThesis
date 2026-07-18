#!/usr/bin/env python3
"""Build an editable YibinThesis DOCX from the template's LaTeX sources.

This converter intentionally supports the controlled source subset used by this
project: metadata in ``metadata.tex``, the ordered ``main.tex`` input list,
standard sectioning, ordinary paragraphs, Pandoc-readable tables/figures/math,
citations, and the YibinThesis abstract/note/unnumbered-chapter commands.

PDF remains the layout-primary output.  Complex TikZ, arbitrary custom macros,
manual page geometry, and unusually nested floats cannot be guaranteed to
round-trip to Word and are reported when detected.
"""

from __future__ import annotations

import argparse
import copy
import html as html_lib
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SECTION_COVER_END = "YIBIN_INTERNAL_SECTION_COVER_END_5B8D"
SECTION_ABSTRACT_END = "YIBIN_INTERNAL_SECTION_ABSTRACT_END_6458"
SECTION_FRONT_END = "YIBIN_INTERNAL_SECTION_FRONT_END_8F91"
TOC_MARKER = "YIBIN_INTERNAL_TOC_76EE"
STYLE_BODY = "宜宾论文-正文"
STYLE_FIRST_PARAGRAPH = "宜宾论文-首段"
STYLE_HEADING_1 = "宜宾论文-一级标题"
STYLE_HEADING_2 = "宜宾论文-二级标题"
STYLE_HEADING_3 = "宜宾论文-三级标题"
STYLE_HEADING_4 = "宜宾论文-四级标题"
STYLE_UNNUMBERED_HEADING = "宜宾论文-无编号标题"
STYLE_FRONT_TITLE = "宜宾论文-中文页标题"
STYLE_ENGLISH_ABSTRACT_TITLE = "宜宾论文-英文摘要标题"
STYLE_TOC_TITLE = "宜宾论文-目录标题"
STYLE_APPENDIX_HEADING = "宜宾论文-附录标题"
STYLE_CHINESE_ABSTRACT = "宜宾论文-中文摘要正文"
STYLE_ENGLISH_ABSTRACT = "宜宾论文-英文摘要正文"
STYLE_KEYWORDS = "宜宾论文-关键词"
STYLE_FIGURE = "宜宾论文-插图"
STYLE_EQUATION = "宜宾论文-公式"
STYLE_FIGURE_CAPTION = "宜宾论文-图题"
STYLE_TABLE_CAPTION = "宜宾论文-表题"
STYLE_TABLE_CONTINUATION = "宜宾论文-续表题"
STYLE_TABLE_TEXT = "宜宾论文-表格正文"
STYLE_TABLE_CENTER = "宜宾论文-表格居中"
STYLE_TABLE_RIGHT = "宜宾论文-表格右对齐"
STYLE_TABLE_HEADER = "宜宾论文-表头"
STYLE_TABLE_HEADER_LEFT = "宜宾论文-表头左对齐"
STYLE_TABLE_HEADER_RIGHT = "宜宾论文-表头右对齐"
STYLE_BIBLIOGRAPHY = "宜宾论文-参考文献"
STYLE_NOTES = "宜宾论文-注释"
STYLE_CITATION = "宜宾论文-文献上标"
STYLE_THREE_LINE_TABLE = "宜宾论文-三线表"
FIGURE_FILTER = (
    Path(__file__).resolve().parents[1]
    / "word"
    / "filters"
    / "latex-figure-to-image.lua"
)

PAGE_BREAK = """```{=openxml}
<w:p><w:r><w:br w:type="page"/></w:r></w:p>
```"""

UNSUPPORTED_PATTERNS = {
    r"\\begin\s*\{tikzpicture\}": "TikZ 图形",
    r"\\begin\s*\{pspicture\}": "PSTricks 图形",
    r"\\begin\s*\{minted\}": "minted 代码环境",
    r"\\begin\s*\{algorithm2e\}": "algorithm2e 环境",
    r"\\write18\b": "shell-escape 命令",
}


class BuildError(RuntimeError):
    """A user-actionable build failure."""


@dataclass
class MainEvent:
    phase: str
    kind: str
    value: str | None = None


@dataclass
class NoteRegistry:
    entries: list[str] = field(default_factory=list)

    def add(self, value: str) -> str:
        self.entries.append(value.strip())
        number = len(self.entries)
        marker = chr(0x245F + number) if number <= 20 else f"[{number}]"
        return rf"\textsuperscript{{{marker}}}"


@dataclass
class HeadingCounters:
    chapter: int = 0
    section: int = 0
    subsection: int = 0
    subsubsection: int = 0


@dataclass
class FloatCounters:
    chapter: int | None = None
    figure: int = 0
    table: int = 0
    equation: int = 0

    def synchronize(self, chapter: int, discipline: str) -> None:
        if discipline == "science" and self.chapter != chapter:
            self.figure = self.table = 0
        if self.chapter != chapter:
            self.equation = 0
        self.chapter = chapter


@dataclass(frozen=True)
class LatexTableColumn:
    """Column semantics distilled from a LaTeX table specification."""

    horizontal: str
    vertical: str = "center"
    width_fraction: float | None = None
    flexible: bool = False
    first_line_indent_pt: float | None = None


@dataclass(frozen=True)
class LatexTableLayout:
    environment: str
    columns: tuple[LatexTableColumn, ...]
    label: str | None = None
    table_width_fraction: float | None = None
    tabcolsep_pt: float | None = None


@dataclass(frozen=True)
class _LatexTablePreamble:
    environment: str
    begin_start: int
    specification_start: int
    specification_end: int
    specification: str
    table_width_fraction: float | None
    tabcolsep_pt: float | None


@dataclass(frozen=True)
class LabelTarget:
    number: str
    kind: str


@dataclass
class LabelRegistry:
    """Track LaTeX labels and deferred references across all source files."""

    targets: dict[str, LabelTarget] = field(default_factory=dict)
    references: dict[str, tuple[str, bool]] = field(default_factory=dict)
    next_reference: int = 0

    def placeholder(self, label: str, *, parenthesized: bool) -> str:
        label = label.strip()
        if not label:
            raise BuildError("检测到空的 LaTeX 交叉引用标签。")
        self.next_reference += 1
        token = f"YIBINXREF{self.next_reference:08d}"
        self.references[token] = (label, parenthesized)
        return token

    def register(self, label: str, number: str, kind: str) -> None:
        label = html_lib.unescape(label.strip())
        if not label:
            return
        if label in self.targets:
            previous = self.targets[label]
            raise BuildError(
                f"LaTeX 标签重复：{label}（{previous.kind} {previous.number} / "
                f"{kind} {number}）。"
            )
        self.targets[label] = LabelTarget(number=number, kind=kind)

    def resolve(self, text: str) -> str:
        missing: set[str] = set()
        for token, (label, parenthesized) in self.references.items():
            target = self.targets.get(label)
            if target is None:
                missing.add(label)
                continue
            # Keep opaque markers until the OOXML stage.  Replacing them here
            # would permanently flatten Word cross-references into plain text.
        if missing:
            raise BuildError("未找到交叉引用标签：" + "、".join(sorted(missing)))
        return text


@dataclass
class CitationRegistry:
    clusters: dict[str, list[str]] = field(default_factory=dict)
    order: list[str] = field(default_factory=list)
    next_cluster: int = 0

    def placeholder(self, keys: str) -> str:
        parsed = [item.strip() for item in split_top_level(keys) if item.strip()]
        if not parsed:
            raise BuildError("检测到空的文献引用键。")
        self.next_cluster += 1
        token = f"YIBINCITE{self.next_cluster:08d}"
        self.clusters[token] = parsed
        for key in parsed:
            if key not in self.order:
                self.order.append(key)
        return token


def strip_tex_comments(text: str) -> str:
    """Remove unescaped LaTeX comments while retaining line structure."""
    output: list[str] = []
    for line in text.splitlines():
        cut = len(line)
        for index, char in enumerate(line):
            if char != "%":
                continue
            backslashes = 0
            cursor = index - 1
            while cursor >= 0 and line[cursor] == "\\":
                backslashes += 1
                cursor -= 1
            if backslashes % 2 == 0:
                cut = index
                break
        output.append(line[:cut])
    return "\n".join(output)


def extract_balanced(text: str, opening: int) -> tuple[str, int]:
    if opening >= len(text) or text[opening] != "{":
        raise BuildError("内部解析错误：预期花括号参数。")
    depth = 0
    escaped = False
    for index in range(opening, len(text)):
        char = text[index]
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return text[opening + 1 : index], index + 1
    raise BuildError("LaTeX 参数花括号未闭合。")


def find_command_argument(text: str, command: str) -> str | None:
    match = re.search(rf"\\{re.escape(command)}\s*\{{", text)
    if not match:
        return None
    opening = text.find("{", match.start())
    value, _ = extract_balanced(text, opening)
    return value


def replace_braced_command(
    text: str,
    command: str,
    replacement: Callable[[str], str],
) -> str:
    pattern = re.compile(rf"\\{re.escape(command)}\s*\{{")
    cursor = 0
    parts: list[str] = []
    while match := pattern.search(text, cursor):
        opening = text.find("{", match.start())
        value, end = extract_balanced(text, opening)
        parts.append(text[cursor : match.start()])
        parts.append(replacement(value))
        cursor = end
    parts.append(text[cursor:])
    return "".join(parts)


def split_top_level(value: str, delimiter: str = ",") -> list[str]:
    depth = 0
    escaped = False
    start = 0
    items: list[str] = []
    for index, char in enumerate(value):
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
        elif char == delimiter and depth == 0:
            items.append(value[start:index])
            start = index + 1
    items.append(value[start:])
    return items


def clean_tex_scalar(value: str) -> str:
    value = value.strip()
    while value.startswith("{") and value.endswith("}"):
        try:
            inner, end = extract_balanced(value, 0)
        except BuildError:
            break
        if end != len(value):
            break
        value = inner.strip()
    replacements = {
        r"\&": "&",
        r"\%": "%",
        r"\#": "#",
        r"\_": "_",
        r"\{": "{",
        r"\}": "}",
        "~": " ",
        r"\\": " ",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    value = re.sub(r"\\(?:textbf|textit|emph|textrm|textsf|texttt)\s*\{([^{}]*)\}", r"\1", value)
    return re.sub(r"\s+", " ", value).strip()


def parse_metadata(path: Path) -> dict[str, str]:
    text = strip_tex_comments(path.read_text(encoding="utf-8"))
    payload = find_command_argument(text, "yibinsetup")
    if payload is None:
        raise BuildError(f"未在 {path} 中找到 \\yibinsetup{{...}}。")
    metadata: dict[str, str] = {}
    for item in split_top_level(payload):
        if not item.strip():
            continue
        key_value = split_top_level(item, delimiter="=")
        if len(key_value) < 2:
            raise BuildError(f"无法解析元数据项：{item.strip()}")
        key = key_value[0].strip()
        raw_value = "=".join(key_value[1:]).strip()
        metadata[key] = clean_tex_scalar(raw_value)
    return metadata


def format_grade_class(metadata: dict[str, str]) -> str:
    """Return the cover value while preserving legacy numeric grade metadata."""
    grade = metadata.get("grade", "").strip()
    class_name = metadata.get("class-name", "").strip()
    if grade and "级" not in grade:
        grade += "级"
    return grade + class_name


def parse_main_events(main_text: str) -> list[MainEvent]:
    text = strip_tex_comments(main_text)
    pattern = re.compile(
        r"\\(?P<kind>frontmatter|mainmatter|backmatter|"
        r"makeyibincover|makeyibindeclarations|tableofcontents|"
        r"printyibinnotes|printyibinbibliography|input|include)"
        r"(?:\s*\{(?P<value>[^{}]+)\})?"
    )
    phase = "pre"
    events: list[MainEvent] = []
    for match in pattern.finditer(text):
        kind = match.group("kind")
        value = match.group("value")
        if kind == "frontmatter":
            phase = "front"
            events.append(MainEvent(phase, kind))
        elif kind == "mainmatter":
            phase = "main"
            events.append(MainEvent(phase, kind))
        elif kind == "backmatter":
            phase = "back"
            events.append(MainEvent(phase, kind))
        else:
            events.append(MainEvent(phase, kind, value.strip() if value else None))
    return events


def parse_discipline(main_text: str) -> str:
    match = re.search(
        r"\\documentclass(?:\[(?P<options>[^]]*)\])?\s*\{yibinthesis\}",
        strip_tex_comments(main_text),
    )
    if not match:
        return "humanities"
    options = {item.strip().lower() for item in (match.group("options") or "").split(",")}
    return "science" if "science" in options else "humanities"


def path_is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except ValueError:
        return False
    return True


def resolve_source(
    name: str,
    main_dir: Path,
    project_root: Path,
    *,
    allow_project_fallback: bool = True,
) -> Path:
    raw = Path(name)
    names = [raw] if raw.suffix else [raw.with_suffix(".tex"), raw]
    bases = [main_dir]
    if allow_project_fallback and main_dir.resolve() != project_root.resolve():
        bases.append(project_root)
    for base in bases:
        for candidate_name in names:
            candidate = (base / candidate_name).resolve()
            if candidate.is_file():
                return candidate
    scope = "入口文件目录" if not allow_project_fallback else "入口文件或模板目录"
    raise BuildError(f"找不到主文件引用的源文件：{name}（已检查{scope}）")


def expand_nested_inputs(
    path: Path,
    project_root: Path,
    seen: set[Path] | None = None,
    *,
    allow_project_fallback: bool = True,
) -> str:
    seen = seen or set()
    resolved = path.resolve()
    if resolved in seen:
        raise BuildError(f"检测到循环 \\input：{resolved}")
    seen.add(resolved)
    text = path.read_text(encoding="utf-8")
    pattern = re.compile(r"\\(?:input|include)\s*\{([^{}]+)\}")

    def include(match: re.Match[str]) -> str:
        child = resolve_source(
            match.group(1).strip(),
            path.parent,
            project_root,
            allow_project_fallback=allow_project_fallback,
        )
        return expand_nested_inputs(
            child,
            project_root,
            seen.copy(),
            allow_project_fallback=allow_project_fallback,
        )

    return pattern.sub(include, text)


def extract_environment(text: str, environment: str) -> str | None:
    match = re.search(
        rf"\\begin\s*\{{{re.escape(environment)}\}}(?P<body>.*?)"
        rf"\\end\s*\{{{re.escape(environment)}\}}",
        text,
        flags=re.DOTALL,
    )
    return match.group("body").strip() if match else None


def detect_unsupported(text: str, source: Path, warnings: list[str]) -> None:
    for pattern, label in UNSUPPORTED_PATTERNS.items():
        if re.search(pattern, text, flags=re.IGNORECASE):
            warnings.append(f"{source}: 检测到{label}；Word 转换不保证版式或内容完整。")


def _latex_dimension_fraction(value: str) -> float | None:
    """Convert a controlled LaTeX width to a fraction of the 15.5 cm text area."""

    compact = re.sub(r"\s+", "", value.strip())
    relative = re.fullmatch(
        r"(?P<factor>(?:\d+(?:\.\d*)?|\.\d+)?)"
        r"\\(?:textwidth|linewidth|columnwidth)",
        compact,
    )
    if relative:
        return float(relative.group("factor") or "1")

    absolute = re.fullmatch(
        r"(?P<value>\d+(?:\.\d*)?|\.\d+)"
        r"(?P<unit>cm|mm|in|pt|bp|pc)",
        compact,
        flags=re.IGNORECASE,
    )
    if not absolute:
        return None
    amount = float(absolute.group("value"))
    unit = absolute.group("unit").casefold()
    centimetres = {
        "cm": amount,
        "mm": amount / 10,
        "in": amount * 2.54,
        "pt": amount * 2.54 / 72.27,
        "bp": amount * 2.54 / 72,
        "pc": amount * 12 * 2.54 / 72.27,
    }[unit]
    return centimetres / 15.5


def _latex_length_points(value: str) -> float | None:
    compact = re.sub(r"\s+", "", value.strip())
    match = re.fullmatch(
        r"(?P<value>-?(?:\d+(?:\.\d*)?|\.\d+))"
        r"(?P<unit>pt|bp|pc|cm|mm|in|em)",
        compact,
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    amount = float(match.group("value"))
    return {
        "pt": amount,
        "bp": amount * 72.27 / 72,
        "pc": amount * 12,
        "cm": amount * 72.27 / 2.54,
        "mm": amount * 72.27 / 25.4,
        "in": amount * 72.27,
        # Table text is five-size (10.5 pt), so source em indents remain
        # proportional to the style rather than to Word's Normal style.
        "em": amount * 10.5,
    }[match.group("unit").casefold()]


def _declaration_alignment(value: str) -> str | None:
    if re.search(r"\\centering\b", value):
        return "center"
    if re.search(r"\\raggedleft\b", value):
        return "right"
    if re.search(r"\\raggedright\b", value):
        return "left"
    return None


def _declaration_first_line_indent(value: str) -> float | None:
    match = re.search(
        r"\\setlength\s*\{\s*\\parindent\s*\}\s*\{([^{}]+)\}",
        value,
    )
    return _latex_length_points(match.group(1)) if match else None


def parse_latex_table_columns(specification: str) -> list[LatexTableColumn]:
    """Parse widths and alignment from the table-column subset used by the class."""

    columns: list[LatexTableColumn] = []
    pending_alignment: str | None = None
    pending_first_line_indent: float | None = None
    index = 0
    length = len(specification)

    def skip_space(cursor: int) -> int:
        while cursor < length and specification[cursor].isspace():
            cursor += 1
        return cursor

    while index < length:
        token = specification[index]
        if token in "><@!":
            cursor = skip_space(index + 1)
            if cursor < length and specification[cursor] == "{":
                declaration, index = extract_balanced(specification, cursor)
                if token == ">":
                    pending_alignment = (
                        _declaration_alignment(declaration) or pending_alignment
                    )
                    declared_indent = _declaration_first_line_indent(declaration)
                    if declared_indent is not None:
                        pending_first_line_indent = declared_indent
                continue

        if token == "*":
            cursor = skip_space(index + 1)
            if cursor < length and specification[cursor] == "{":
                repeats_text, cursor = extract_balanced(specification, cursor)
                cursor = skip_space(cursor)
                if cursor < length and specification[cursor] == "{":
                    repeated_spec, index = extract_balanced(specification, cursor)
                    try:
                        repeats = max(0, int(repeats_text.strip()))
                    except ValueError as error:
                        raise BuildError(
                            f"Word 转换不支持非整数表格重复列次数：{repeats_text.strip()}"
                        ) from error
                    repeated_columns = parse_latex_table_columns(repeated_spec)
                    for _ in range(repeats):
                        columns.extend(repeated_columns)
                    pending_alignment = None
                    pending_first_line_indent = None
                    continue

        if token in "LCRpmb":
            cursor = skip_space(index + 1)
            if cursor < length and specification[cursor] == "{":
                width, index = extract_balanced(specification, cursor)
                if token in "LCR":
                    horizontal = {
                        "L": "left",
                        "C": "center",
                        "R": "right",
                    }[token]
                    vertical = "center"
                else:
                    horizontal = pending_alignment or "left"
                    vertical = {"p": "top", "m": "center", "b": "bottom"}[token]
                columns.append(
                    LatexTableColumn(
                        horizontal=pending_alignment or horizontal,
                        vertical=vertical,
                        width_fraction=_latex_dimension_fraction(width),
                        first_line_indent_pt=pending_first_line_indent,
                    )
                )
                pending_alignment = None
                pending_first_line_indent = None
                continue

        if token in "lcrXS":
            if token in "lcr":
                horizontal = {"l": "left", "c": "center", "r": "right"}[token]
            elif token == "S":
                horizontal = "right"
            else:
                horizontal = "left"
            columns.append(
                LatexTableColumn(
                    horizontal=pending_alignment or horizontal,
                    vertical="center",
                    flexible=token == "X",
                    first_line_indent_pt=pending_first_line_indent,
                )
            )
            pending_alignment = None
            pending_first_line_indent = None
            index += 1
            if token == "S":
                index = skip_space(index)
                if index < length and specification[index] == "[":
                    closing = specification.find("]", index + 1)
                    if closing < 0:
                        raise BuildError("S 表格列选项缺少右方括号。")
                    index = closing + 1
            continue

        if token == "\\":
            command = re.match(r"\\[A-Za-z@]+|\\.", specification[index:])
            index += len(command.group(0)) if command else 1
            index = skip_space(index)
            while index < length and specification[index] == "{":
                _, index = extract_balanced(specification, index)
                index = skip_space(index)
            continue
        if token.isalpha():
            raise BuildError(f"Word 转换不支持表格列类型：{token}")
        index += 1
    return columns


def count_latex_table_columns(specification: str) -> int:
    """Count columns in the controlled LaTeX table specification subset."""

    return len(parse_latex_table_columns(specification))


def _table_local_length_points(
    text: str,
    begin_start: int,
    command: str,
) -> float | None:
    r"""Return the nearest table-local ``\setlength`` value before a preamble.

    Table spacing declarations in the supplied sources live either inside a
    ``table`` float or a ``\begingroup`` block immediately surrounding a
    ``longtable``.  Stop at the nearest opening/closing table boundary so a
    value from an earlier table cannot leak into a later one.
    """

    boundary_markers = (
        r"\begingroup",
        r"\endgroup",
        r"\begin{table}",
        r"\begin{table*}",
        r"\end{table}",
        r"\end{table*}",
        r"\end{longtable}",
    )
    scope_start = max(
        (text.rfind(marker, 0, begin_start) for marker in boundary_markers),
        default=-1,
    )
    segment = text[scope_start if scope_start >= 0 else 0 : begin_start]
    matches = list(
        re.finditer(
            rf"\\setlength\s*\{{\s*\\{re.escape(command)}\s*\}}"
            r"\s*\{([^{}]+)\}",
            segment,
        )
    )
    return _latex_length_points(matches[-1].group(1)) if matches else None


def _find_latex_table_preambles(text: str) -> list[_LatexTablePreamble]:
    preambles: list[_LatexTablePreamble] = []
    pattern = re.compile(
        r"\\begin\s*\{(?P<environment>longtable|tabularx|tabular\*|tabular)\}"
    )

    def skip_space(cursor: int) -> int:
        while cursor < len(text) and text[cursor].isspace():
            cursor += 1
        return cursor

    def skip_optional_position(cursor: int) -> int:
        cursor = skip_space(cursor)
        if cursor >= len(text) or text[cursor] != "[":
            return cursor
        closing = text.find("]", cursor + 1)
        return skip_space(closing + 1) if closing >= 0 else len(text)

    for match in pattern.finditer(text):
        environment = match.group("environment")
        cursor = skip_space(match.end())
        table_width_fraction: float | None = None
        try:
            if environment in {"tabularx", "tabular*"}:
                if cursor >= len(text) or text[cursor] != "{":
                    continue
                width, cursor = extract_balanced(text, cursor)
                table_width_fraction = _latex_dimension_fraction(width)
                cursor = skip_optional_position(cursor)
            else:
                cursor = skip_optional_position(cursor)
            cursor = skip_space(cursor)
            if cursor >= len(text) or text[cursor] != "{":
                continue
            specification_start = cursor + 1
            specification, cursor = extract_balanced(text, cursor)
            specification_end = cursor - 1
        except BuildError:
            continue
        preambles.append(
            _LatexTablePreamble(
                environment=environment,
                begin_start=match.start(),
                specification_start=specification_start,
                specification_end=specification_end,
                specification=specification,
                table_width_fraction=(
                    1.0 if environment == "longtable" else table_width_fraction
                ),
                tabcolsep_pt=_table_local_length_points(
                    text,
                    match.start(),
                    "tabcolsep",
                ),
            )
        )
    return preambles


def _table_layout_label(text: str, preamble: _LatexTablePreamble) -> str | None:
    if preamble.environment == "longtable":
        end = text.find(r"\end{longtable}", preamble.specification_end)
        segment = text[
            preamble.begin_start : end if end >= 0 else len(text)
        ]
    else:
        prefix = text[: preamble.begin_start]
        table_start = prefix.rfind(r"\begin{table}")
        if table_start < 0:
            table_start = prefix.rfind(r"\begin{table*}")
        last_table_end = max(prefix.rfind(r"\end{table}"), prefix.rfind(r"\end{table*}"))
        if table_start > last_table_end:
            candidates = [
                position
                for marker in (r"\end{table}", r"\end{table*}")
                if (position := text.find(marker, preamble.specification_end)) >= 0
            ]
            end = min(candidates) if candidates else len(text)
            segment = text[table_start:end]
        else:
            end_marker = rf"\end{{{preamble.environment}}}"
            end = text.find(end_marker, preamble.specification_end)
            segment = text[
                preamble.begin_start : end if end >= 0 else len(text)
            ]
    labels = re.findall(r"\\label\s*\{([^{}]+)\}", segment)
    return next((label for label in labels if label.strip().startswith("tab:")), labels[0] if labels else None)


def extract_latex_table_layouts(text: str) -> list[LatexTableLayout]:
    """Collect semantic table layouts in document order before Pandoc runs."""

    layouts: list[LatexTableLayout] = []
    for preamble in _find_latex_table_preambles(text):
        columns = parse_latex_table_columns(preamble.specification)
        if columns:
            layouts.append(
                LatexTableLayout(
                    environment=preamble.environment,
                    columns=tuple(columns),
                    label=_table_layout_label(text, preamble),
                    table_width_fraction=preamble.table_width_fraction,
                    tabcolsep_pt=preamble.tabcolsep_pt,
                )
            )
    return layouts


def normalize_latex_table_preambles(text: str) -> str:
    """Make only table preambles Pandoc-readable; leave body prose untouched."""

    replacements: list[tuple[int, int, str]] = []
    for preamble in _find_latex_table_preambles(text):
        specification = re.sub(
            r"(?<![A-Za-z@\\])(?P<column>[LCR])\s*\{[^{}]+\}",
            lambda match: {"L": "l", "C": "c", "R": "r"}[
                match.group("column")
            ],
            preamble.specification,
        )
        if specification != preamble.specification:
            replacements.append(
                (
                    preamble.specification_start,
                    preamble.specification_end,
                    specification,
                )
            )
    for start, end, replacement in reversed(replacements):
        text = text[:start] + replacement + text[end:]
    return text


def normalize_longtable_headers(text: str) -> str:
    """Remove repeated-page header definitions before Pandoc reads longtable.

    Pandoc already turns the first longtable header into a Markdown table
    header.  Keeping the block between ``\\endfirsthead`` and ``\\endhead``
    adds a second header plus a blank row, which makes the later Markdown to
    DOCX pass treat the whole table as plain text.
    """

    replacements: list[tuple[int, int, str]] = []
    for preamble in _find_latex_table_preambles(text):
        if preamble.environment != "longtable":
            continue
        end_match = re.search(
            r"\\end\s*\{longtable\}",
            text[preamble.specification_end + 1 :],
        )
        if end_match is None:
            continue
        end_start = preamble.specification_end + 1 + end_match.start()
        end = preamble.specification_end + 1 + end_match.end()
        body = text[preamble.specification_end + 1 : end_start]
        body = re.sub(
            r"\\endfirsthead\b.*?\\endhead\b",
            "\n",
            body,
            flags=re.DOTALL,
        )
        body = re.sub(r"\\end(?:firsthead|head)\b", "", body)
        columns = parse_latex_table_columns(preamble.specification)
        simplified_spec = "".join(
            {"left": "l", "center": "c", "right": "r"}[column.horizontal]
            for column in columns
        ) or "l"
        replacements.append(
            (
                preamble.begin_start,
                end,
                text[preamble.begin_start : preamble.specification_start]
                + simplified_spec
                + text[preamble.specification_end : preamble.specification_end + 1]
                + body
                + text[end_start:end],
            )
        )
    for start, end, replacement in reversed(replacements):
        text = text[:start] + replacement + text[end:]
    return text


def resolve_word_image_target(target: str, resource_roots: Iterable[Path]) -> str:
    """Prefer a Word-compatible raster sibling for PDF/SVG figure targets."""

    value = html_lib.unescape(target).strip()
    if not value or re.match(r"^(?:https?|data):", value, flags=re.IGNORECASE):
        return value

    relative = Path(value.replace("/", os.sep))
    if relative.suffix.casefold() not in {".pdf", ".svg", ".eps"}:
        return value

    candidates = [relative] if relative.is_absolute() else [root / relative for root in resource_roots]
    for candidate in candidates:
        for suffix in (".png", ".jpg", ".jpeg"):
            raster = candidate.with_suffix(suffix)
            if not raster.is_file():
                continue
            if relative.is_absolute():
                return raster.resolve().as_posix()
            return relative.with_suffix(suffix).as_posix()
    return value


def normalize_latex(
    text: str,
    notes: NoteRegistry,
    discipline: str = "humanities",
    labels: LabelRegistry | None = None,
    citations: CitationRegistry | None = None,
    table_layouts: list[LatexTableLayout] | None = None,
) -> str:
    text = strip_tex_comments(text)
    if table_layouts is not None:
        table_layouts.extend(extract_latex_table_layouts(text))
    text = normalize_longtable_headers(text)
    # Pandoc understands native l/c/r columns but discards the alignment on
    # array's paragraph columns.  Width/vertical semantics are retained in the
    # table layout registry and reapplied to Word after conversion.
    text = normalize_latex_table_preambles(text)
    labels = labels if labels is not None else LabelRegistry()
    citations = citations if citations is not None else CitationRegistry()
    chapter_command = "chapter" if discipline == "science" else "chapter*"

    def thematic_chapter(match: re.Match[str], default: str) -> str:
        title = match.group("title") if match.group("title") is not None else default
        return rf"\{chapter_command}{{{title}}}"

    def unnumbered_thematic_chapter(match: re.Match[str], default: str) -> str:
        title = match.group("title") if match.group("title") is not None else default
        return rf"\chapter*{{{title}}}"

    text = re.sub(
        r"\\yibinopeningchapter(?![A-Za-z@])(?:\s*\[(?P<title>[^]]*)\])?",
        lambda match: thematic_chapter(match, "绪论"),
        text,
    )
    text = re.sub(
        r"\\yibinclosingchapter(?![A-Za-z@])(?:\s*\[(?P<title>[^]]*)\])?",
        lambda match: unnumbered_thematic_chapter(match, "结论"),
        text,
    )
    text = replace_braced_command(
        text,
        "yibinunnumberedchapter",
        lambda title: rf"\chapter*{{{title}}}",
    )
    text = re.sub(
        r"\\begin\s*\{yibinappendices\}",
        "",
        text,
    )
    text = re.sub(r"\\end\s*\{yibinappendices\}", "", text)
    appendix_index = 0

    def appendix_heading(title: str) -> str:
        nonlocal appendix_index
        appendix_index += 1
        marker = chr(ord("A") + appendix_index - 1) if appendix_index <= 26 else str(appendix_index)
        return rf"\chapter*{{附录{marker} {title}}}"

    text = replace_braced_command(text, "yibinappendix", appendix_heading)
    text = replace_braced_command(
        text,
        "yibincite",
        citations.placeholder,
    )
    text = replace_braced_command(text, "cite", citations.placeholder)
    text = replace_braced_command(text, "yibinnote", notes.add)
    # Pandoc resolves each standalone LaTeX fragment with a fresh counter, so
    # its native \ref text is stale for cross-file targets.  Preserve the
    # semantic reference as an opaque token and resolve it after all labels
    # have received their template-specific global number.
    text = replace_braced_command(
        text,
        "eqref",
        lambda label: labels.placeholder(label, parenthesized=True),
    )
    text = replace_braced_command(
        text,
        "ref",
        lambda label: labels.placeholder(label, parenthesized=False),
    )
    text = re.sub(r"\\FloatBarrier\b", "", text)
    # The example's framed rule is a LaTeX-only visual placeholder.  Preserve
    # its semantic presence instead of silently dropping the whole figure.
    text = re.sub(
        r"\\fbox\s*\{\s*\\rule\s*\{[^{}]*\}\s*\{[^{}]*\}\s*"
        r"\\rule\s*\{[^{}]*\}\s*\{[^{}]*\}\s*\}",
        r"\\textit{[插图占位框]}",
        text,
    )
    return text


def run_checked(
    command: list[str],
    *,
    cwd: Path,
    capture_output: bool = False,
) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        command,
        cwd=cwd,
        text=True,
        encoding="utf-8",
        errors="replace",
        stdout=subprocess.PIPE if capture_output else None,
        stderr=subprocess.PIPE,
        check=False,
    )
    if result.returncode != 0:
        rendered = subprocess.list2cmdline(command)
        detail = (result.stderr or "").strip()
        raise BuildError(f"命令执行失败（{result.returncode}）：{rendered}\n{detail}")
    return result


def collect_graphic_paths(
    main_text: str,
    *,
    main_dir: Path,
    project_root: Path,
    allow_project_fallback: bool = True,
) -> list[Path]:
    """Return existing directories declared through ``\\graphicspath``.

    LaTeX resolves graphics against its process/search path, while the
    Pandoc DOCX writer resolves image targets through ``--resource-path``.
    Entry-file-relative directories take priority. Project-root-relative
    fallback is retained only for entry files that live inside the template
    repository, where bundled examples intentionally use root-relative paths.
    """
    paths: list[Path] = []
    cleaned = strip_tex_comments(main_text)
    for outer in re.finditer(
        r"\\graphicspath\s*\{(?P<body>(?:\s*\{[^{}]*\}\s*)+)\}",
        cleaned,
        flags=re.DOTALL,
    ):
        for item in re.findall(r"\{([^{}]+)\}", outer.group("body")):
            raw = item.strip().replace("/", os.sep)
            if not raw:
                continue
            candidate = Path(raw)
            if candidate.is_absolute():
                options = [candidate]
            else:
                options = [main_dir / candidate]
                if allow_project_fallback and main_dir.resolve() != project_root.resolve():
                    options.append(project_root / candidate)
            for option in options:
                resolved = option.resolve()
                if resolved.is_dir() and resolved not in paths:
                    paths.append(resolved)
    return paths


def latex_to_markdown(
    text: str,
    *,
    pandoc: Path,
    cwd: Path,
    temp_dir: Path,
    name: str,
) -> str:
    source = temp_dir / f"{name}.tex"
    source.write_text(text, encoding="utf-8", newline="\n")
    command = [
        str(pandoc),
        "--from=latex-smart",
        "--to=markdown+raw_attribute+fenced_divs+bracketed_spans+tex_math_dollars",
        "--wrap=none",
    ]
    if FIGURE_FILTER.is_file():
        command.append(f"--lua-filter={FIGURE_FILTER}")
    command.append(str(source))
    result = run_checked(
        command,
        cwd=cwd,
        capture_output=True,
    )
    markdown = (result.stdout or "").strip()

    def restore_chinese_quotes(match: re.Match[str]) -> str:
        content = match.group(1)
        if re.search(r"[\u3400-\u9fff]", content):
            return f"“{content}”"
        return match.group(0)

    return re.sub(r'"([^"\n]+)"', restore_chinese_quotes, markdown)


def chinese_number(value: int) -> str:
    digits = "零一二三四五六七八九"
    if value < 10:
        return digits[value]
    if value == 10:
        return "十"
    if value < 20:
        return "十" + digits[value % 10]
    if value < 100:
        suffix = "" if value % 10 == 0 else digits[value % 10]
        return digits[value // 10] + "十" + suffix
    return str(value)


def apply_heading_numbering(
    markdown: str,
    discipline: str,
    counters: HeadingCounters,
    labels: LabelRegistry | None = None,
) -> str:
    labels = labels or LabelRegistry()
    output: list[str] = []
    heading = re.compile(r"^(#{1,4})\s+(.+?)\s*$")
    attrs = re.compile(r"^(?P<title>.*?)(?P<attrs>\s+\{[^{}]*\})?$")
    for line in markdown.splitlines():
        match = heading.match(line)
        if not match:
            output.append(line)
            continue
        level = len(match.group(1))
        parsed = attrs.match(match.group(2))
        assert parsed is not None
        title = parsed.group("title").strip()
        attribute_text = parsed.group("attrs") or ""
        unnumbered = ".unnumbered" in attribute_text or re.search(r"(?:^|\s)-(?:\s|})", attribute_text)

        if unnumbered:
            if level == 1:
                counters.section = counters.subsection = counters.subsubsection = 0
            output.append(f"{'#' * level} {title}{attribute_text}")
            continue

        if level == 1:
            counters.chapter += 1
            counters.section = counters.subsection = counters.subsubsection = 0
        elif level == 2:
            counters.section += 1
            counters.subsection = counters.subsubsection = 0
        elif level == 3:
            counters.subsection += 1
            counters.subsubsection = 0
        else:
            counters.subsubsection += 1

        if discipline == "science":
            numbers = [
                counters.chapter,
                counters.section,
                counters.subsection,
                counters.subsubsection,
            ][:level]
            prefix = ".".join(str(number) for number in numbers)
            reference_number = prefix
        elif level == 1:
            reference_number = f"{chinese_number(counters.chapter)}、"
        elif level == 2:
            reference_number = f"（{chinese_number(counters.section)}）"
        elif level == 3:
            reference_number = f"{counters.subsection}."
        else:
            reference_number = f"（{counters.subsubsection}）"
        identifier = re.search(r"(?:^|[\s{])#([^\s}]+)", attribute_text)
        if identifier:
            labels.register(identifier.group(1), reference_number, "标题")
        if level == 1:
            output.append(f"<!-- YIBIN_INTERNAL_CHAPTER_{counters.chapter} -->")
        # Numbering is attached to Word's built-in Heading 1-4 styles during
        # OOXML post-processing.  Keep the paragraph text itself unnumbered so
        # Word users can renumber, reorder, and cross-reference headings.
        output.append(f"{'#' * level} {title}{attribute_text}")
    return "\n".join(output)


def apply_float_numbering(
    markdown: str,
    discipline: str,
    chapter: int,
    counters: FloatCounters,
    labels: LabelRegistry | None = None,
    image_target_resolver: Callable[[str], str] | None = None,
) -> str:
    """Add global caption/equation numbers and register their LaTeX labels.

    ``chapter`` is the chapter active before this Markdown fragment.  Internal
    markers emitted by :func:`apply_heading_numbering` let one source file
    contain multiple chapters without assigning every float to the final one.
    """

    labels = labels or LabelRegistry()

    def markdown_identifier(attributes: str | None) -> str | None:
        if not attributes:
            return None
        match = re.search(r"(?:^|\s)#([^\s}]+)", attributes)
        return match.group(1) if match else None

    def process_chunk(chunk: str, current_chapter: int) -> str:
        counters.synchronize(current_chapter, discipline)

        def next_number(kind: str) -> str:
            value = getattr(counters, kind) + 1
            setattr(counters, kind, value)
            if kind == "equation" or discipline == "science":
                return f"{current_chapter}.{value}"
            return str(value)

        # Pandoc wraps labelled tables in a fenced div and emits the caption as
        # an indented ``: caption`` line.  Track the surrounding div id while
        # numbering so the label and caption receive the same global number.
        table_lines: list[str] = []
        div_stack: list[str | None] = []
        for line in chunk.splitlines():
            opening = re.match(r"^\s*:::\s+\{(?P<attrs>[^{}]*)\}\s*$", line)
            closing = re.match(r"^\s*:::\s*$", line)
            if opening:
                div_stack.append(markdown_identifier(opening.group("attrs")))
                table_lines.append(line)
                continue
            if closing:
                if div_stack:
                    div_stack.pop()
                table_lines.append(line)
                continue
            caption = re.match(
                r"^(?P<indent>[ \t]*):\s+(?P<caption>[^\n]+)$",
                line,
            )
            if caption:
                number = next_number("table")
                identifier = next(
                    (item for item in reversed(div_stack) if item),
                    None,
                )
                if identifier:
                    labels.register(identifier, number, "表")
                line = (
                    f"{caption.group('indent')}: 表{number} "
                    f"{caption.group('caption').strip()}"
                )
            table_lines.append(line)
        chunk = "\n".join(table_lines)

        def markdown_figure(match: re.Match[str]) -> str:
            number = next_number("figure")
            identifier = markdown_identifier(match.group("attrs"))
            if identifier:
                labels.register(identifier, number, "图")
            return (
                f"![图{number} {match.group('caption').strip()}]"
                f"({match.group('target')}){match.group('attrs') or ''}"
            )

        chunk = re.sub(
            r"(?m)^!\[(?P<caption>[^\]\n]*)\]"
            r"\((?P<target><[^>]+>|[^)\n]+)\)"
            r"(?P<attrs>\{[^{}]*\})?\s*$",
            markdown_figure,
            chunk,
        )

        def image(match: re.Match[str]) -> str:
            attributes = match.group("attributes")
            src_match = re.search(
                r'\bsrc="([^"]+)"',
                attributes,
                flags=re.IGNORECASE,
            )
            if not src_match:
                return match.group(0)
            src = html_lib.unescape(src_match.group(1))
            if image_target_resolver is not None:
                src = image_target_resolver(src)
            alt_match = re.search(
                r'\balt="([^"]*)"',
                attributes,
                flags=re.IGNORECASE,
            )
            alt = html_lib.unescape(alt_match.group(1)) if alt_match else ""
            style_match = re.search(
                r'\bstyle="([^"]*)"',
                attributes,
                flags=re.IGNORECASE,
            )
            options: list[str] = []
            if style_match:
                style = style_match.group(1)
                for dimension in ("width", "height"):
                    dimension_value = re.search(
                        rf"(?:^|;)\s*{dimension}\s*:\s*([^;]+)",
                        style,
                        flags=re.IGNORECASE,
                    )
                    if dimension_value:
                        options.append(f"{dimension}={dimension_value.group(1).strip()}")
            target_value = f"<{src}>" if re.search(r"[\s()]", src) else src
            attributes_md = "{" + " ".join(options) + "}" if options else ""
            return f"![{alt}]({target_value}){attributes_md}"

        def figure(match: re.Match[str]) -> str:
            number = next_number("figure")
            attrs = match.group("attrs")
            identifier = re.search(r'\bid="([^"]+)"', attrs)
            if identifier:
                labels.register(identifier.group(1), number, "图")
            caption = re.sub(r"<[^>]+>", "", match.group("caption")).strip()
            body = match.group("body").strip()
            body = re.sub(r"^<p>(.*?)</p>$", r"\1", body, flags=re.DOTALL)

            body = re.sub(
                r"<(?:img|embed)\s+(?P<attributes>[^>]*?)/?>",
                image,
                body,
                flags=re.IGNORECASE,
            )
            opening = f"::: {{#{identifier.group(1)}}}" if identifier else "::: {}"
            return (
                f"{opening}\n{body}\n:::\n\n"
                + custom_block("Caption", f"图{number} {caption}")
            )

        chunk = re.sub(
            r"<figure(?P<attrs>[^>]*)>\s*(?P<body>.*?)\s*"
            r"<figcaption>(?P<caption>.*?)</figcaption>\s*</figure>",
            figure,
            chunk,
            flags=re.DOTALL | re.IGNORECASE,
        )

        chunk = re.sub(
            r"<(?:img|embed)\s+(?P<attributes>[^>]*?)/?>",
            image,
            chunk,
            flags=re.IGNORECASE,
        )

        def equation(match: re.Match[str]) -> str:
            number = next_number("equation")
            body = match.group("body")
            equation_labels = re.findall(r"\\label\s*\{([^{}]+)\}", body)
            for label in equation_labels:
                labels.register(label, number, "公式")
            body = re.sub(r"\\label\s*\{[^{}]+\}", "", body).strip()
            marker = custom_block(
                "YibinSectionMarker",
                f"YIBIN_INTERNAL_EQUATION_{number.replace('.', '_')}",
            )
            return f"$$\n{body}\n$$\n\n{marker}"

        return re.sub(
            r"\$\$\s*\\begin\s*\{equation\}(?P<body>.*?)"
            r"\\end\s*\{equation\}\s*\$\$",
            equation,
            chunk,
            flags=re.DOTALL,
        )

    marker = re.compile(r"(?m)^<!-- YIBIN_INTERNAL_CHAPTER_(\d+) -->\s*$")
    output: list[str] = []
    cursor = 0
    current_chapter = chapter
    for match in marker.finditer(markdown):
        output.append(process_chunk(markdown[cursor : match.start()], current_chapter))
        current_chapter = int(match.group(1))
        cursor = match.end()
    output.append(process_chunk(markdown[cursor:], current_chapter))
    return "".join(output)


def prepare_bibliography(
    bibliography: Path,
    *,
    pandoc: Path,
    cwd: Path,
    temp_dir: Path,
) -> Path:
    """Convert BibLaTeX to CSL JSON and preserve @standard as CSL standard.

    Pandoc maps BibLaTeX ``@standard`` to ``legislation`` by default, which the
    bundled GB/T style correctly renders as archival material [A].  The source
    semantics are unambiguous here, so restore only those entries to the CSL
    ``standard`` type before citeproc runs.
    """
    if bibliography.suffix.lower() not in {".bib", ".biblatex"}:
        return bibliography
    result = run_checked(
        [str(pandoc), "--from=biblatex", "--to=csljson", str(bibliography)],
        cwd=cwd,
        capture_output=True,
    )
    try:
        items = json.loads(result.stdout or "[]")
    except json.JSONDecodeError as error:
        raise BuildError(f"Pandoc 无法解析参考文献 JSON：{error}") from error
    source = strip_tex_comments(bibliography.read_text(encoding="utf-8"))
    entry_types = {
        key.strip(): entry_type.lower()
        for entry_type, key in re.findall(
            r"@([A-Za-z]+)\s*\{\s*([^,\s]+)",
            source,
        )
    }
    for item in items:
        if entry_types.get(str(item.get("id", ""))) == "standard":
            item["type"] = "standard"
    output = temp_dir / "references.csl.json"
    output.write_text(
        json.dumps(items, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    return output


def custom_block(style: str, content: str) -> str:
    return f'::: {{custom-style="{style}"}}\n{content.strip()}\n:::'


def styled_marker(value: str) -> str:
    return custom_block("YibinSectionMarker", value)


def escape_markdown(value: str) -> str:
    for char in ("\\", "[", "]", "*", "_"):
        value = value.replace(char, "\\" + char)
    return value


def abstract_markdown(
    source_text: str,
    *,
    environment: str,
    pandoc: Path,
    cwd: Path,
    temp_dir: Path,
    notes: NoteRegistry,
    discipline: str,
    labels: LabelRegistry,
    citations: CitationRegistry,
    table_layouts: list[LatexTableLayout],
) -> str:
    body = extract_environment(source_text, environment)
    if body is None:
        raise BuildError(f"摘要文件缺少 {environment} 环境。")
    keyword_command = "cnkeywords" if environment == "cnabstract" else "enkeywords"
    keywords = find_command_argument(body, keyword_command) or ""
    body = replace_braced_command(body, keyword_command, lambda _: "")
    body_md = latex_to_markdown(
        normalize_latex(
            body,
            notes,
            discipline,
            labels,
            citations,
            table_layouts,
        ),
        pandoc=pandoc,
        cwd=cwd,
        temp_dir=temp_dir,
        name=environment,
    )
    if environment == "cnabstract":
        heading = custom_block("FrontTitle", "摘要")
        body_block = custom_block("ChineseAbstract", body_md)
        label = "关键词："
        separator = ""
    else:
        heading = custom_block("FrontTitle", "Abstract")
        body_block = custom_block("EnglishAbstract", body_md)
        label = "Keywords:"
        separator = " "
    keyword_line = (
        f'[{label}]{{custom-style="KeywordLabel"}}{separator}'
        f'{escape_markdown(clean_tex_scalar(keywords))}'
    )
    return "\n\n".join([heading, body_block, custom_block("Keywords", keyword_line)])


def bibliography_markdown() -> str:
    return "# 参考文献 {.unnumbered}\n\n::: {#refs .references}\n:::"


def citation_seed_markdown(citations: CitationRegistry) -> str:
    """Feed citeproc every cited key without leaving a visible citation.

    Inline citations are replaced by opaque markers so they can become Word
    fields later.  Citeproc still needs semantic citation nodes to construct
    the bibliography; a removable seed paragraph supplies them in first-use
    order.
    """
    if not citations.order:
        return ""
    items = "; ".join(f"@{key}" for key in citations.order)
    return f"YIBIN_INTERNAL_CITATION_SEED [{items}]"


def notes_markdown(notes: NoteRegistry) -> str:
    if not notes.entries:
        return ""
    lines = ["# 注释 {.unnumbered}"]
    for index, entry in enumerate(notes.entries, start=1):
        marker = chr(0x245F + index) if index <= 20 else f"[{index}]"
        lines.append(custom_block("Notes", f"{marker} {entry}"))
    return "\n\n".join(lines)


def find_pandoc(explicit: str | None, project_root: Path) -> Path:
    candidates: list[Path] = []
    if explicit:
        candidates.append(Path(explicit).expanduser())
    if os.environ.get("PANDOC"):
        candidates.append(Path(os.environ["PANDOC"]).expanduser())
    candidates.extend(
        [
            project_root / ".tools" / "pandoc-3.9.0.2" / "pandoc.exe",
            project_root / ".tools" / "pandoc" / "pandoc.exe",
            project_root / "tools" / "pandoc.exe",
            project_root / "tools" / "pandoc" / "pandoc.exe",
            project_root / "vendor" / "pandoc" / "pandoc.exe",
        ]
    )
    on_path = shutil.which("pandoc")
    if on_path:
        candidates.append(Path(on_path))
    for candidate in candidates:
        candidate = candidate.resolve()
        if candidate.is_file():
            return candidate
    raise BuildError(
        "未找到 Pandoc。请安装 Pandoc 并加入 PATH，或使用 --pandoc 指定可执行文件。"
    )


def _set_page_number_format(section, format_name: str | None, start: int | None = None) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if format_name is None:
        if existing is not None:
            sect_pr.remove(existing)
        return
    if existing is None:
        existing = OxmlElement("w:pgNumType")
        sect_pr.append(existing)
    existing.set(qn("w:fmt"), format_name)
    if start is not None:
        existing.set(qn("w:start"), str(start))


def _strip_section_links(sect_pr) -> None:
    for child in list(sect_pr):
        if child.tag in {
            qn("w:headerReference"),
            qn("w:footerReference"),
            qn("w:pgNumType"),
            qn("w:titlePg"),
            qn("w:type"),
        }:
            sect_pr.remove(child)


def _add_section_break(paragraph, base_sect_pr, page_format: str | None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for existing in p_pr.findall(qn("w:sectPr")):
        p_pr.remove(existing)
    sect_pr = copy.deepcopy(base_sect_pr)
    _strip_section_links(sect_pr)
    section_type = OxmlElement("w:type")
    section_type.set(qn("w:val"), "nextPage")
    sect_pr.insert(0, section_type)
    if page_format:
        page_num = OxmlElement("w:pgNumType")
        page_num.set(qn("w:fmt"), page_format)
        page_num.set(qn("w:start"), "1")
        sect_pr.append(page_num)
    p_pr.append(sect_pr)
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)


def _insert_toc_after(paragraph) -> None:
    toc_paragraph = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    p_style = OxmlElement("w:pStyle")
    p_style.set(qn("w:val"), "Normal")
    p_pr.append(p_style)
    toc_paragraph.append(p_pr)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    field.set(qn("w:dirty"), "true")
    toc_paragraph.append(field)
    paragraph._p.addnext(toc_paragraph)


def _format_equation_number(
    paragraph,
    number: str,
    labels: LabelRegistry | None = None,
    *,
    discipline: str = "science",
    equation_sequence: str = "Equation",
) -> None:
    """Center an OMML equation and add a native Word ``SEQ 公式`` number.

    ``equation_sequence`` is detected from Word's built-in CaptionLabels when
    the PowerShell build entry is used.  The visible form remains ``(2.1)``.
    """
    if paragraph._p.find(qn("m:oMathPara")) is None and paragraph._p.find(qn("m:oMath")) is None:
        raise BuildError(f"公式编号 {number} 前未找到 Word 公式对象。")
    parsed_number = re.fullmatch(r"(?P<chapter>\d+)\.(?P<sequence>\d+)", number)
    if parsed_number is None:
        raise BuildError(f"Word 公式编号必须为章号.章内序号：{number}")
    chapter_number = parsed_number.group("chapter")
    sequence_number = parsed_number.group("sequence")
    p_pr = paragraph._p.get_or_add_pPr()
    # Tab positions, spacing, alignment, and fonts live in the reusable
    # 宜宾论文-公式 paragraph style.  Keep only the two tab characters and the
    # fields as paragraph content.
    for name in ("tabs", "jc", "ind", "spacing"):
        node = p_pr.find(qn(f"w:{name}"))
        if node is not None:
            p_pr.remove(node)

    leading_run = OxmlElement("w:r")
    leading_run.append(OxmlElement("w:tab"))
    paragraph._p.insert(1 if paragraph._p[0] is p_pr else 0, leading_run)
    trailing_tab = OxmlElement("w:r")
    trailing_tab.append(OxmlElement("w:tab"))
    paragraph._p.append(trailing_tab)

    matching_labels = [
        label
        for label, target in (labels.targets.items() if labels else [])
        if target.kind == "公式" and target.number == number
    ]
    document = paragraph.part.document
    bookmark_id = _next_bookmark_id(document)
    full_pairs = []
    number_pairs = []
    bookmark_seeds: list[str | None] = matching_labels or [None]
    for label in bookmark_seeds:
        seed = label or f"anonymous:equation:{number}:{bookmark_id}"
        full_start = OxmlElement("w:bookmarkStart")
        full_start.set(qn("w:id"), str(bookmark_id))
        full_start.set(
            qn("w:name"),
            _label_bookmark_name(label)
            if label
            else _hidden_ref_bookmark_name(f"full:{seed}"),
        )
        full_end = OxmlElement("w:bookmarkEnd")
        full_end.set(qn("w:id"), str(bookmark_id))
        bookmark_id += 1
        number_start = OxmlElement("w:bookmarkStart")
        number_start.set(qn("w:id"), str(bookmark_id))
        number_start.set(
            qn("w:name"),
            _label_number_bookmark_name(label)
            if label
            else _hidden_ref_bookmark_name(
                f"number:{seed}",
                prefix="_RefNum",
            ),
        )
        number_end = OxmlElement("w:bookmarkEnd")
        number_end.set(qn("w:id"), str(bookmark_id))
        bookmark_id += 1
        full_pairs.append((full_start, full_end))
        number_pairs.append((number_start, number_end))

    for start, _ in full_pairs:
        paragraph._p.append(start)
    open_run = OxmlElement("w:r")
    open_text = OxmlElement("w:t")
    open_text.text = "("
    open_run.append(open_text)
    paragraph._p.append(open_run)
    for start, _ in number_pairs:
        paragraph._p.append(start)

    if discipline == "science":
        # The science profile's thesis Heading 1 style owns an Arabic native
        # multilevel number, so STYLEREF keeps the chapter prefix synchronized.
        for node in _complex_field_runs(
            f' STYLEREF "{STYLE_HEADING_1}" \\n ',
            chapter_number,
        ):
            paragraph._p.append(node)
    else:
        # Humanities headings use Chinese counters (一、/（一）), while the
        # school equation form remains Arabic (2.1).  The Arabic chapter value
        # therefore remains source-derived; the equation sequence is native.
        chapter_run = _plain_run(chapter_number)
        paragraph._p.append(chapter_run)

    separator_run = _plain_run(".")
    paragraph._p.append(separator_run)
    equation_instruction = f" SEQ {_word_field_identifier(equation_sequence)} "
    if sequence_number == "1":
        equation_instruction += "\\r 1 "
    equation_instruction += "\\* ARABIC "
    for node in _complex_field_runs(equation_instruction, sequence_number):
        paragraph._p.append(node)
    for _, end in reversed(number_pairs):
        paragraph._p.append(end)
    close_run = OxmlElement("w:r")
    close_text = OxmlElement("w:t")
    close_text.text = ")"
    close_run.append(close_text)
    paragraph._p.append(close_run)
    for _, end in reversed(full_pairs):
        paragraph._p.append(end)


def _clear_story(story) -> None:
    paragraphs = list(story.paragraphs)
    if not paragraphs:
        story.add_paragraph()
        return
    first = paragraphs[0]
    first.clear()
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def _add_page_field(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph._p.append(_field_run(" PAGE ", "1"))


def _add_header_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")


def _set_run_fonts(
    run,
    east_asia: str,
    latin: str,
    size: float,
    *,
    bold: bool | None = None,
    underline: bool | None = None,
) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if underline is not None:
        run.font.underline = underline
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    color = r_pr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        r_pr.append(color)
    color.set(qn("w:val"), "000000")
    for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
        color.attrib.pop(qn(attribute), None)


def _png_has_transparency(data: bytes) -> bool:
    """Return whether a PNG uses an alpha channel or palette transparency."""

    if len(data) < 33 or data[:8] != b"\x89PNG\r\n\x1a\n":
        return False
    color_type = data[25]
    if color_type in {4, 6}:
        return True
    if color_type != 3:
        return False
    offset = 8
    while offset + 12 <= len(data):
        length = int.from_bytes(data[offset : offset + 4], "big")
        chunk_type = data[offset + 4 : offset + 8]
        if chunk_type == b"tRNS":
            return True
        offset += 12 + length
        if chunk_type == b"IEND":
            break
    return False


def _flatten_transparent_png(data: bytes, media_name: str) -> bytes:
    try:
        from PIL import Image
    except ImportError as error:
        raise BuildError(
            "检测到带透明通道的 PNG（"
            + media_name
            + "），Word 兼容转换需要 Pillow；请运行 python -m pip install Pillow。"
        ) from error

    with Image.open(io.BytesIO(data)) as source:
        source.load()
        rgba = source.convert("RGBA")
        white = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        white.alpha_composite(rgba)
        rgb = white.convert("RGB")
        output = io.BytesIO()
        rgb.save(output, format="PNG", compress_level=6)
        return output.getvalue()


def whiten_signature_image(source_path: Path, output_path: Path | None = None) -> io.BytesIO:
    """Map a photographed signature's paper background to white without touching the source."""

    try:
        from PIL import Image
    except ImportError as error:
        raise BuildError("签名背景漂白需要 Pillow；请安装 requirements-word.txt。") from error

    with Image.open(source_path) as source:
        source.load()
        grayscale = source.convert("L")
        histogram = grayscale.histogram()
        pixel_count = sum(histogram)

        def percentile(fraction: float) -> int:
            threshold = pixel_count * fraction
            cumulative = 0
            for value, count in enumerate(histogram):
                cumulative += count
                if cumulative >= threshold:
                    return value
            return 255

        black_point = percentile(0.02)
        white_point = percentile(0.20)
        if white_point - black_point < 20:
            whitened = Image.new("RGB", grayscale.size, (255, 255, 255))
        else:
            scale = 255.0 / (white_point - black_point)
            lookup = [
                0
                if value <= black_point
                else 255
                if value >= white_point
                else round((value - black_point) * scale)
                for value in range(256)
            ]
            whitened = grayscale.point(lookup).convert("RGB")
        stream = io.BytesIO()
        whitened.save(stream, format="PNG", compress_level=6, dpi=(300, 300))
        stream.seek(0)
        if output_path is not None:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(stream.getvalue())
            stream.seek(0)
        return stream


def _normalize_embedded_pngs(docx_path: Path) -> int:
    """Flatten transparent embedded PNG media without changing relationships."""

    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(docx_path, "r") as package:
        for item in package.infolist():
            media_name = item.filename
            if not media_name.startswith("word/media/") or not media_name.lower().endswith(".png"):
                continue
            data = package.read(item)
            if _png_has_transparency(data):
                replacements[media_name] = _flatten_transparent_png(data, media_name)
    if not replacements:
        return 0

    temporary = docx_path.with_name(docx_path.stem + ".rgb-normalized.docx")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for item in source.infolist():
            data = replacements.get(item.filename, source.read(item))
            target.writestr(item, data)
    os.replace(temporary, docx_path)
    return len(replacements)


def _insert_paragraph_before(document: Document, anchor, style_name: str):
    paragraph = document.add_paragraph(style=style_name)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def _set_tabs(paragraph, stops: Iterable[tuple[int, str, str | None]]) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:tabs"))
    if existing is not None:
        p_pr.remove(existing)
    tabs = OxmlElement("w:tabs")
    for position, alignment, leader in stops:
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), alignment)
        tab.set(qn("w:pos"), str(position))
        if leader:
            tab.set(qn("w:leader"), leader)
        tabs.append(tab)
    p_pr.append(tabs)


def _add_tab(paragraph) -> None:
    paragraph.add_run().add_tab()


def _add_cover_label(paragraph, value: str) -> None:
    run = paragraph.add_run(value)
    run.style = "CoverLabel"
    _set_run_fonts(run, "SimHei", "SimHei", 18, bold=True, underline=False)
    if value.startswith("指导教师"):
        # Eight 18 pt CJK glyphs occupy the official 144 pt label slot exactly.
        # A tiny style-neutral character condensation prevents Word's table end
        # mark from wrapping the final two glyphs without changing the font size.
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "-4")
        run._r.get_or_add_rPr().append(spacing)


def _add_cover_value(paragraph, value: str, *, suffix: str = "") -> None:
    # A zero-width character keeps an intentionally empty slot alive when Word
    # opens and saves the layout table.  It is invisible and, unlike NBSP or
    # ordinary spaces, cannot create a second underline.
    visible = value.strip()
    run = paragraph.add_run((visible if visible else "\u200b") + suffix)
    run.style = "CoverValue"
    _set_run_fonts(run, "SimSun", "SimSun", 16, bold=True, underline=False)


def _prepare_cover_field(paragraph, stops: Iterable[tuple[int, str, str | None]]) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    _set_tabs(paragraph, stops)


def _set_paragraph_bottom_border(target) -> None:
    """Give a paragraph or paragraph style one Word-native bottom border."""

    p_pr = target._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")


def _hide_run(run) -> None:
    r_pr = run._r.get_or_add_rPr(); vanish = OxmlElement("w:vanish"); r_pr.append(vanish)


def _set_fixed_table_geometry(table, widths_cm: Iterable[float], *, style_name: str) -> None:
    """Give Word a complete fixed-width grid instead of width hints.

    Setting only ``cell.width`` leaves the original equal-column ``tblGrid``
    created by python-docx in place.  Word then trusts that grid on save and
    can collapse or expand the cover columns.  Keep tblW, tblGrid and every
    tcW in exact agreement.
    """
    widths = [int(Cm(width).twips) for width in widths_cm]
    tbl_pr = table._tbl.tblPr
    style = tbl_pr.find(qn("w:tblStyle"))
    if style is None:
        style = OxmlElement("w:tblStyle")
        tbl_pr.insert(0, style)
    style.set(qn("w:val"), style_name)

    table_width = tbl_pr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        tbl_pr.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(widths)))

    table_indent = tbl_pr.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        tbl_pr.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "0")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    cell_margins = tbl_pr.find(qn("w:tblCellMar"))
    if cell_margins is None:
        cell_margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_margins)
    for edge in ("top", "left", "bottom", "right"):
        margin = cell_margins.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            cell_margins.append(margin)
        margin.set(qn("w:type"), "dxa")
        margin.set(qn("w:w"), "0")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        cant_split = row._tr.get_or_add_trPr().find(qn("w:cantSplit"))
        if cant_split is None:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths, strict=True):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def _set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    else:
        for child in list(borders):
            borders.remove(child)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def _insert_layout_separator(anchor) -> None:
    """Prevent Word from merging adjacent layout tables on open/save."""
    paragraph = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "1")
    spacing.set(qn("w:lineRule"), "exact")
    p_pr.append(spacing)
    mark_properties = OxmlElement("w:rPr")
    mark_properties.append(OxmlElement("w:vanish"))
    mark_size = OxmlElement("w:sz")
    mark_size.set(qn("w:val"), "2")
    mark_properties.append(mark_size)
    mark_size_cs = OxmlElement("w:szCs")
    mark_size_cs.set(qn("w:val"), "2")
    mark_properties.append(mark_size_cs)
    p_pr.append(mark_properties)
    paragraph.append(p_pr)
    anchor._p.addprevious(paragraph)


def _insert_cover_table(
    document: Document,
    anchor,
    cells: list[tuple[str, str, float]],
    *,
    row_height_pt: float,
    hidden_tokens: str = "",
):
    """Insert one official cover row as a borderless layout table."""
    table = document.add_table(rows=1, cols=len(cells))
    anchor._p.addprevious(table._tbl)
    # The official cover starts every field row at the 3 cm text margin, while
    # the underline length differs from row to row.  Left alignment preserves
    # that common origin; centring narrower rows shifts every underline.
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_fixed_table_geometry(table, (cell[2] for cell in cells), style_name="YibinCoverLayout")
    _set_table_borders_none(table)
    row = table.rows[0]
    row.height = Pt(row_height_pt)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    for index, (kind, value, width_cm) in enumerate(cells):
        cell = table.cell(0, index)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_together = True
        p_pr = p._p.get_or_add_pPr()
        indent = p_pr.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            p_pr.append(indent)
        for attribute in ("left", "right", "firstLine", "leftChars", "rightChars", "firstLineChars"):
            indent.set(qn(f"w:{attribute}"), "0")
        if kind == "label":
            if index == 0:
                p.style = document.styles["CoverField"]
                if value.startswith("指导教师"):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
            _add_cover_label(p, value)
        else:
            p.style = document.styles["CoverValueLine"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_cover_value(p, value)
    _insert_layout_separator(anchor)
    return table


def _calibrate_cover_styles(document: Document) -> None:
    """Apply the official first-page rhythm through reusable Word styles."""

    logo = document.styles["CoverLogo"].paragraph_format
    logo.space_before = Pt(2.0)
    logo.space_after = Pt(0)

    thesis_type = document.styles["CoverThesisType"].paragraph_format
    thesis_type.space_before = Pt(24.35)
    thesis_type.space_after = Pt(52.35)

    title = document.styles["CoverTitle"].paragraph_format
    title.space_before = Pt(0)
    title.space_after = Pt(47.3)
    title.line_spacing = 2.0

    field = document.styles["CoverField"].paragraph_format
    field.space_before = Pt(0)
    field.space_after = Pt(0)
    field.line_spacing = 1.0

    try:
        value_line_style = document.styles["CoverValueLine"]
    except KeyError:
        value_line_style = document.styles.add_style(
            "CoverValueLine", WD_STYLE_TYPE.PARAGRAPH
        )
    value_line_style.base_style = document.styles["Normal"]
    value_line = value_line_style.paragraph_format
    value_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    value_line.first_line_indent = Pt(0)
    value_line.left_indent = Pt(0)
    value_line.right_indent = Pt(0)
    value_line.space_before = Pt(0)
    value_line.space_after = Pt(0)
    # An exact 20 pt line box keeps the rule at the official underline height
    # even when the slot is empty and contains only the zero-width keeper.
    value_line.line_spacing = Pt(20.0)
    value_line.keep_together = True
    _set_paragraph_bottom_border(value_line_style)


def _insert_declaration_table(
    document: Document,
    anchor,
    cells: list[tuple[str, str | Path | None, float]],
    *,
    space_before: float = 0,
    signature_background: str = "preserve",
):
    """Insert fixed signature/date slots that survive a Word save."""
    table = document.add_table(rows=1, cols=len(cells))
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_fixed_table_geometry(table, (cell[2] for cell in cells), style_name="YibinFrontLayout")
    _set_table_borders_none(table)
    for index, (kind, value, width_cm) in enumerate(cells):
        cell = table.cell(0, index)
        paragraph = cell.paragraphs[0]
        paragraph.style = document.styles[
            "DeclarationDate" if kind == "date" else "DeclarationSignature"
        ]
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(space_before)
        paragraph.paragraph_format.space_after = Pt(0)
        p_pr = paragraph._p.get_or_add_pPr()
        indent = p_pr.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            p_pr.append(indent)
        for attribute in ("left", "right", "firstLine", "leftChars", "rightChars", "firstLineChars"):
            indent.set(qn(f"w:{attribute}"), "0")
        if kind == "label":
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if kind == "signature":
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Keep the rule attached to the signature paragraph.  A cell
            # border sits at the bottom of the tallest image row and therefore
            # drifts visibly downward compared with Word's official line.
            _set_paragraph_bottom_border(paragraph)
            if isinstance(value, Path):
                picture_run = paragraph.add_run()
                picture_source: str | io.BytesIO = str(value)
                if signature_background == "whiten":
                    picture_source = whiten_signature_image(value)
                picture_run.add_picture(
                    picture_source,
                    width=Cm(max(0.5, min(width_cm - 0.2, 2.6))),
                )
                for doc_property in picture_run._r.iter(qn("wp:docPr")):
                    doc_property.attrib.pop("descr", None)
            else:
                paragraph.add_run("\u200b")
        else:
            run = paragraph.add_run(str(value) if value else "\u200b")
            _set_run_fonts(run, "SimSun", "SimSun", 12, bold=False, underline=False)
    _insert_layout_separator(anchor)
    return table


def _resolve_cover_logo(
    metadata: dict[str, str],
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    *,
    allow_project_fallback: bool,
) -> Path | None:
    configured = metadata.get("logo", "builtin").strip()
    normalized = configured.casefold()
    if normalized in {"builtin", "built-in", "default"}:
        logo = (project_root / "assets" / "yibin-university-logo.png").resolve()
        if not logo.is_file():
            raise BuildError(f"模板内置封面校徽不存在：{logo}")
        return logo
    if not configured or normalized == "none":
        return None
    configured_path = Path(configured).expanduser()
    if configured_path.is_absolute():
        candidates = [configured_path]
    else:
        candidates = [
            main_dir / configured_path,
            metadata_dir / configured_path,
        ]
    if allow_project_fallback:
        candidates.extend(
            [
                project_root / "assets" / "yibin-university-logo.png",
                project_root / "assets" / "yibin-logo.png",
            ]
        )
    logo = next(
        (
            candidate.resolve()
            for candidate in dict.fromkeys(candidates)
            if candidate.is_file()
        ),
        None,
    )
    if logo is None and configured and not allow_project_fallback:
        raise BuildError(
            f"外部入口声明的封面校徽不存在：{configured}；"
            "不会回退到模板目录中的同名资源。"
        )
    return logo


def _resolve_signature_asset(
    metadata: dict[str, str],
    key: str,
    main_dir: Path,
    metadata_dir: Path,
) -> Path | None:
    configured = metadata.get(key, "").strip()
    if not configured or configured.casefold() in {"none", "null"}:
        return None
    configured_path = Path(configured).expanduser()
    candidates = (
        [configured_path]
        if configured_path.is_absolute()
        else [metadata_dir / configured_path, main_dir / configured_path]
    )
    asset = next(
        (
            candidate.resolve()
            for candidate in dict.fromkeys(candidates)
            if candidate.is_file()
        ),
        None,
    )
    if asset is None:
        raise BuildError(f"元数据 {key} 声明的签名图片不存在：{configured}")
    return asset


def _build_cover_page(
    document: Document,
    anchor,
    metadata: dict[str, str],
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    *,
    allow_project_fallback: bool,
    page_break_after: bool,
) -> None:
    _calibrate_cover_styles(document)
    logo = _resolve_cover_logo(
        metadata,
        project_root,
        main_dir,
        metadata_dir,
        allow_project_fallback=allow_project_fallback,
    )
    logo_paragraph = _insert_paragraph_before(document, anchor, "CoverLogo")
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.first_line_indent = Pt(0)
    if logo is not None:
        picture_run = logo_paragraph.add_run()
        picture_run.add_picture(str(logo), width=Cm(13.15), height=Cm(3.65))
        for doc_property in picture_run._r.iter(qn("wp:docPr")):
            doc_property.attrib.pop("descr", None)
    else:
        run = logo_paragraph.add_run("宜宾学院")
        _set_run_fonts(run, "SimHei", "Times New Roman", 26, bold=True)

    thesis_type = _insert_paragraph_before(document, anchor, "CoverThesisType")
    thesis_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    thesis_type.paragraph_format.first_line_indent = Pt(0)
    type_run = thesis_type.add_run("本科生毕业论文（设计）")
    _set_run_fonts(type_run, "SimHei", "SimHei", 28, bold=True, underline=False)

    title_paragraph = _insert_paragraph_before(document, anchor, "CoverTitle")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.first_line_indent = Pt(0)
    title_run = title_paragraph.add_run(metadata.get("title", "").strip() or "（填写中文题目）")
    _set_run_fonts(title_run, "SimHei", "SimHei", 18, bold=True, underline=True)

    # Widths are distilled from the official Word cover at page coordinates
    # (points): college 84.85|183.85|513.00, major 84.85|183.85|522.00,
    # student 84.85|174.85|513.00.  Each value cell owns the only visible rule.
    _insert_cover_table(
        document,
        anchor,
        [("label", "学院（部）", 3.4925), ("value", metadata.get("college", ""), 11.6121)],
        row_height_pt=46.30,
    )
    _insert_cover_table(
        document,
        anchor,
        [("label", "专    业", 3.4925), ("value", metadata.get("major", ""), 11.9296)],
        row_height_pt=46.25,
    )
    _insert_cover_table(
        document,
        anchor,
        [("label", "学生姓名", 3.1750), ("value", metadata.get("author", ""), 11.9296)],
        row_height_pt=47.60,
    )

    _insert_cover_table(
        document, anchor,
        [("label", "学    号", 3.1750), ("value", metadata.get("student-id", ""), 5.6686),
         ("label", "年级", 1.5875), ("value", format_grade_class(metadata), 4.5367)],
        row_height_pt=46.25,
    )

    advisor_rows = (
        (
            "指导教师（校内）",
            metadata.get("advisor", ""),
            metadata.get("advisor-title", ""),
        ),
        (
            "指导教师（校外）",
            metadata.get("external-advisor", ""),
            metadata.get("external-advisor-title", ""),
        ),
    )
    last_table = None
    for label, person, professional_title in advisor_rows:
        last_table = _insert_cover_table(
            document, anchor,
            [("label", label, 5.0800), ("value", person, 5.1259),
             ("label", "职称", 1.5416), ("value", professional_title, 3.1289)],
            row_height_pt=46.30,
        )


def _estimate_declaration_lines(text: str) -> int:
    units = sum(0.5 if ord(character) < 128 else 1.0 for character in text)
    if units <= 28:
        return 1
    return 1 + int((units - 28 + 30.999) // 31)


def _add_declaration_text(document: Document, anchor, style_name: str, text: str):
    paragraph = _insert_paragraph_before(document, anchor, style_name)
    run = paragraph.add_run(text)
    _set_run_fonts(run, "SimSun", "SimSun", 12, bold=False, underline=False)
    return paragraph


def _build_originality_page(
    document: Document,
    anchor,
    metadata: dict[str, str],
    author_signature: Path | None,
    signature_background: str,
) -> None:
    heading = _insert_paragraph_before(document, anchor, "DeclarationTitle")
    heading.paragraph_format.page_break_before = True
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run("原创性声明")
    _set_run_fonts(heading_run, "SimHei", "SimHei", 16, bold=True, underline=False)

    title = metadata.get("title", "").strip() or "（填写中文题目）"
    statement = (
        f"本人呈交的学位论文（设计）《{title}》，是在导师的指导下，独立进行研究取得的成果。"
        "除文中已经注明引用的内容外，本论文（设计）不包括其他个人或集体已经发表或撰写过的作品成果。"
        "对本文（设计）做出贡献的个人和集体，均已在文中以明确方式标明。"
        "本人完全意识到本声明的法律后果，因本声明而产生的法律后果由本人承担。"
    )
    body = _add_declaration_text(document, anchor, "DeclarationBody", statement)
    body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body.paragraph_format.first_line_indent = Pt(24)

    extra_lines = max(0, _estimate_declaration_lines(statement) - 5)
    _insert_declaration_table(
        document,
        anchor,
        [
            ("spacer", "", 0.85),
            ("label", "学位论文作者：", 3.6),
            ("signature", author_signature, 3.0),
            ("spacer", "", 8.05),
        ],
        space_before=max(0, 36 - extra_lines * 31.2),
        signature_background=signature_background,
    )

    date = _add_declaration_text(
        document,
        anchor,
        "DeclarationDate",
        "日期：    年   月   日",
    )
    date.paragraph_format.first_line_indent = Pt(24)

    lead = _insert_paragraph_before(document, anchor, "DeclarationRegulationLead")
    lead_run = lead.add_run("附：")
    lead_run.style = "DeclarationRegulationLeadLabel"
    _set_run_fonts(lead_run, "SimSun", "SimSun", 16, bold=True, underline=False)
    regulation_title = lead.add_run(
        "《普通高等学校学生管理规定》（中华人民共和国教育部令第41号）"
    )
    regulation_title.style = "DeclarationRegulationText"
    _set_run_fonts(regulation_title, "SimSun", "SimSun", 14, bold=True, underline=False)

    clause = _insert_paragraph_before(document, anchor, "DeclarationRegulationClause")
    clause_run = clause.add_run(
        "第五十二条\u00a0学生有下列情形之一，学校可以给予开除学籍处分："
    )
    clause_run.style = "DeclarationRegulationText"
    _set_run_fonts(clause_run, "SimSun", "SimSun", 14, bold=True, underline=False)

    item = _insert_paragraph_before(document, anchor, "DeclarationRegulationItem")
    item_run = item.add_run(
        "（五）学位论文、公开发表的研究成果存在抄袭、篡改、伪造等学术不端行为，"
        "情节严重的，或者代写论文、买卖论文的；"
    )
    item_run.style = "DeclarationRegulationText"
    _set_run_fonts(item_run, "SimSun", "SimSun", 14, bold=True, underline=False)
    item.add_run().add_break(WD_BREAK.PAGE)


def _build_authorization_page(
    document: Document,
    anchor,
    metadata: dict[str, str],
    author_signature: Path | None,
    advisor_signature: Path | None,
    signature_background: str,
) -> None:
    heading = _insert_paragraph_before(document, anchor, "DeclarationTitle")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run("学位论文（设计）版权使用授权书")
    _set_run_fonts(heading_run, "SimHei", "SimHei", 16, bold=True, underline=False)

    authorization = (
        "本学位论文（设计）作者完全了解学校有关保留、使用学位论文（设计）的规定，"
        "同意学校保留并向国家有关部门或机构送交论文（设计）的复印件和电子版，允许论文（设计）"
        "被查阅和借阅。本人授权宜宾学院将本学位论文（设计）的全部或部分内容编入有关数据库进行检索，"
        "可以采用影印、缩印或扫描等复制手段保存和汇编。"
    )
    body = _add_declaration_text(document, anchor, "DeclarationBody", authorization)
    body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body.paragraph_format.first_line_indent = Pt(24)

    choice_intro = _insert_paragraph_before(document, anchor, "DeclarationBody")
    choice_intro.paragraph_format.first_line_indent = Pt(24)
    choice_intro.paragraph_format.line_spacing = 1.5
    choice_intro.paragraph_format.space_before = Pt(36)
    intro_run = choice_intro.add_run("本学位论文（设计）属于")
    _set_run_fonts(intro_run, "SimSun", "SimSun", 12, bold=False, underline=False)
    instruction = choice_intro.add_run("（请在以下相应方框内打“√”）作品")
    _set_run_fonts(instruction, "SimSun", "SimSun", 12, bold=False, underline=True)

    confidential = metadata.get("secrecy", "public").strip().lower() == "confidential"
    declassify = metadata.get("declassify-year", "").strip() or "____"
    for text in (
        f"保  密{'■' if confidential else '□'}，在 {declassify} 年解密后适用本授权书。",
        f"不保密{'□' if confidential else '■'}。",
    ):
        paragraph = _add_declaration_text(document, anchor, "DeclarationBody", text)
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.5

    _insert_declaration_table(
        document,
        anchor,
        [
            ("label", "作者（签名）：", 3.5),
            ("signature", author_signature, 2.8),
            ("spacer", "", 0.2),
            ("label", "指导教师（签名）：", 4.2),
            ("signature", advisor_signature, 2.8),
        ],
        space_before=72,
        signature_background=signature_background,
    )
    _insert_declaration_table(
        document,
        anchor,
        [
            ("date", "日期：    年   月   日", 6.2),
            ("spacer", "", 1.3),
            ("date", "日期：    年   月   日", 6.0),
        ],
    )


def _build_front_matter(
    document: Document,
    anchor,
    metadata: dict[str, str],
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    *,
    allow_project_fallback: bool,
    include_cover: bool,
    include_declarations: bool,
) -> None:
    if include_cover:
        _build_cover_page(
            document,
            anchor,
            metadata,
            project_root,
            main_dir,
            metadata_dir,
            allow_project_fallback=allow_project_fallback,
            page_break_after=include_declarations,
        )
    if include_declarations:
        signature_background = metadata.get("signature-background", "preserve").strip().casefold()
        if signature_background not in {"preserve", "whiten"}:
            raise BuildError(
                "元数据 signature-background 只能是 preserve 或 whiten："
                + signature_background
            )
        author_signature = _resolve_signature_asset(
            metadata,
            "author-signature",
            main_dir,
            metadata_dir,
        )
        advisor_signature = _resolve_signature_asset(
            metadata,
            "advisor-signature",
            main_dir,
            metadata_dir,
        )
        _build_originality_page(
            document,
            anchor,
            metadata,
            author_signature,
            signature_background,
        )
        _build_authorization_page(
            document,
            anchor,
            metadata,
            author_signature,
            advisor_signature,
            signature_background,
        )


def _clear_direct_paragraph_format(paragraph, *names: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for name in names:
        node = p_pr.find(qn(f"w:{name}"))
        if node is not None:
            p_pr.remove(node)


def _clear_direct_run_typography(run) -> None:
    r_pr = run._r.find(qn("w:rPr"))
    if r_pr is None:
        return
    for name in ("rFonts", "sz", "szCs"):
        node = r_pr.find(qn(f"w:{name}"))
        if node is not None:
            r_pr.remove(node)
    if len(r_pr) == 0:
        run._r.remove(r_pr)


def _clear_title_run_typography(run) -> None:
    r_pr = run._r.find(qn("w:rPr"))
    if r_pr is None:
        return
    for name in (
        "rFonts",
        "b",
        "bCs",
        "i",
        "iCs",
        "color",
        "u",
        "sz",
        "szCs",
        "highlight",
        "caps",
        "smallCaps",
        "strike",
        "dstrike",
        "vertAlign",
        "spacing",
        "position",
        "kern",
    ):
        node = r_pr.find(qn(f"w:{name}"))
        if node is not None:
            r_pr.remove(node)
    if len(r_pr) == 0:
        run._r.remove(r_pr)


def _next_numbering_id(numbering, tag: str, attribute: str) -> int:
    values = [
        int(value)
        for element in numbering.findall(qn(f"w:{tag}"))
        if (value := element.get(qn(f"w:{attribute}"))) is not None
        and value.isdigit()
    ]
    return max(values, default=0) + 1


def _set_style_numbering(style, num_id: int, level: int) -> None:
    p_pr = style._element.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    insertion = 0
    for index, child_element in enumerate(p_pr):
        if child_element.tag == qn("w:pStyle"):
            insertion = index + 1
    p_pr.insert(insertion, num_pr)


def _configure_heading_numbering(document: Document, discipline: str) -> None:
    """Attach a real four-level Word list to the reusable heading styles."""

    numbering = document.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "abstractNum", "abstractNumId")
    num_id = _next_numbering_id(numbering, "num", "numId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), hashlib.sha1(f"YibinHeading:{discipline}".encode()).hexdigest()[:8].upper())
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.extend([nsid, multi])

    if discipline == "science":
        formats = ["decimal"] * 4
        level_texts = ["%1", "%1.%2", "%1.%2.%3", "%1.%2.%3.%4"]
        suffix = "space"
    else:
        formats = [
            "chineseCounting",
            "chineseCounting",
            "decimal",
            "decimal",
        ]
        level_texts = ["%1、", "（%2）", "%3.", "（%4）"]
        suffix = "nothing"

    style_names = [
        STYLE_HEADING_1,
        STYLE_HEADING_2,
        STYLE_HEADING_3,
        STYLE_HEADING_4,
    ]
    for level, (style_name, number_format, level_text) in enumerate(
        zip(style_names, formats, level_texts)
    ):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), number_format)
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), document.styles[style_name].style_id)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), suffix)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), level_text)
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        # Word's default for a multilevel list is to restart each lower level
        # when its immediately higher level advances.  Explicit lvlRestart
        # values generated by hand are easy to get off by one, so retain the
        # native default verified through Word COM.
        lvl.extend([start, num_fmt, p_style, suff, text, justification])
        abstract.append(lvl)
        _set_style_numbering(document.styles[style_name], num_id, level)

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def _resolve_table_column_widths(
    layout: LatexTableLayout | None,
    raw_widths: list[int],
    target_width: int,
) -> list[int]:
    if layout is not None and len(layout.columns) == len(raw_widths):
        fractions = [column.width_fraction for column in layout.columns]
        if any(value is not None and value > 0 for value in fractions):
            known_sum = sum(value or 0 for value in fractions)
            unknown = [index for index, value in enumerate(fractions) if value is None]
            if unknown:
                outer_fraction = layout.table_width_fraction or 1.0
                if 0 < known_sum < outer_fraction:
                    fallback = (outer_fraction - known_sum) / len(unknown)
                else:
                    known = [value for value in fractions if value is not None and value > 0]
                    fallback = (sum(known) / len(known)) if known else 1
                weights = [fallback if value is None else value for value in fractions]
            else:
                weights = [value or 0 for value in fractions]
            if sum(weights) > 0:
                scale = target_width / sum(weights)
                widths = [max(1, round(weight * scale)) for weight in weights]
                widths[-1] += target_width - sum(widths)
                return widths

    usable = raw_widths if len(raw_widths) > 0 and sum(raw_widths) > 0 else [1]
    scale = target_width / sum(usable)
    widths = [max(1, round(width * scale)) for width in usable]
    widths[-1] += target_width - sum(widths)
    return widths


def _table_style_for_alignment(row_index: int, horizontal: str) -> str:
    if row_index == 0:
        return {
            "left": STYLE_TABLE_HEADER_LEFT,
            "center": STYLE_TABLE_HEADER,
            "right": STYLE_TABLE_HEADER_RIGHT,
        }.get(horizontal, STYLE_TABLE_HEADER_LEFT)
    return {
        "left": STYLE_TABLE_TEXT,
        "center": STYLE_TABLE_CENTER,
        "right": STYLE_TABLE_RIGHT,
    }.get(horizontal, STYLE_TABLE_TEXT)


def _caption_number_by_table(document: Document) -> dict[object, str]:
    result: dict[object, str] = {}
    body = document._element.body
    children = list(body)
    for index, element in enumerate(children):
        if element.tag != qn("w:tbl"):
            continue
        for previous in reversed(children[max(0, index - 3) : index]):
            if previous.tag == qn("w:tbl"):
                break
            if previous.tag != qn("w:p"):
                continue
            value = "".join(node.text or "" for node in previous.iter(qn("w:t"))).strip()
            match = re.match(r"^表\s*([0-9]+(?:\.[0-9]+)?)\b", value)
            if match:
                result[element] = match.group(1)
                break
    return result


def _format_tables(
    document: Document,
    table_layouts: Iterable[LatexTableLayout] = (),
    labels: LabelRegistry | None = None,
) -> None:
    page_target_width = int(Cm(15.5).twips)
    layouts = list(table_layouts)
    caption_numbers = _caption_number_by_table(document)
    numbered_layouts: dict[str, LatexTableLayout] = {}
    if labels is not None:
        for layout in layouts:
            target = labels.targets.get(layout.label or "")
            if target is not None and target.kind == "表":
                numbered_layouts[target.number] = layout
    unused_layouts = list(layouts)
    used_layout_ids: set[int] = set()
    semantic_table_count = 0
    for table in document.tables:
        style = table._tbl.tblPr.find(qn("w:tblStyle"))
        if style is not None and style.get(qn("w:val")) in {"YibinCoverLayout", "YibinFrontLayout"}:
            continue
        semantic_table_count += 1
        caption_number = caption_numbers.get(table._tbl)
        source_layout = numbered_layouts.get(caption_number or "")
        if source_layout is None:
            source_layout = next(
                (layout for layout in unused_layouts if id(layout) not in used_layout_ids),
                None,
            )
        if source_layout is not None:
            used_layout_ids.add(id(source_layout))
        target_width = page_target_width
        if (
            source_layout is not None
            and source_layout.table_width_fraction is not None
            and source_layout.table_width_fraction > 0
        ):
            effective_fraction = source_layout.table_width_fraction
            known_widths = [
                column.width_fraction
                for column in source_layout.columns
                if column.width_fraction is not None
                and column.width_fraction > 0
            ]
            if len(known_widths) == len(source_layout.columns):
                effective_fraction = min(
                    effective_fraction,
                    sum(known_widths),
                )
            target_width = max(
                1,
                round(
                    page_target_width
                    * min(1.0, effective_fraction)
                ),
            )
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = document.styles[STYLE_THREE_LINE_TABLE]
        table.autofit = False
        tbl_pr = table._tbl.tblPr

        table_width = tbl_pr.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            tbl_pr.append(table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(target_width))

        table_indent = tbl_pr.find(qn("w:tblInd"))
        if table_indent is None:
            table_indent = OxmlElement("w:tblInd")
            tbl_pr.append(table_indent)
        table_indent.set(qn("w:type"), "dxa")
        table_indent.set(qn("w:w"), "0")

        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")

        cell_margins = tbl_pr.find(qn("w:tblCellMar"))
        if cell_margins is None:
            cell_margins = OxmlElement("w:tblCellMar")
            tbl_pr.append(cell_margins)
        horizontal_margin = max(
            0,
            round(
                20
                * (
                    source_layout.tabcolsep_pt
                    if source_layout is not None
                    and source_layout.tabcolsep_pt is not None
                    else 6.0
                )
            ),
        )
        for edge, width in (
            ("top", 80),
            ("left", horizontal_margin),
            ("bottom", 80),
            ("right", horizontal_margin),
        ):
            margin = cell_margins.find(qn(f"w:{edge}"))
            if margin is None:
                margin = OxmlElement(f"w:{edge}")
                cell_margins.append(margin)
            margin.set(qn("w:type"), "dxa")
            margin.set(qn("w:w"), str(width))

        grid = table._tbl.tblGrid
        grid_columns = list(grid.iterchildren(qn("w:gridCol")))
        column_count = max(1, len(table.columns))
        raw_widths = [int(column.get(qn("w:w"), "0")) for column in grid_columns]
        if len(raw_widths) != column_count or sum(raw_widths) <= 0:
            raw_widths = [1] * column_count
            while len(grid_columns) < column_count:
                column = OxmlElement("w:gridCol")
                grid.append(column)
                grid_columns.append(column)
        if source_layout is not None and len(source_layout.columns) != column_count:
            print(
                "WARNING: LaTeX 表格列数与 Word 表格不一致，已回退到 Pandoc 列宽/对齐："
                f"LaTeX={len(source_layout.columns)}, Word={column_count}",
                file=sys.stderr,
            )
            source_layout = None
        widths = _resolve_table_column_widths(source_layout, raw_widths, target_width)
        for column, width in zip(grid_columns, widths):
            column.set(qn("w:w"), str(width))

        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge, value, size in (
            ("top", "single", "12"),
            ("bottom", "single", "12"),
            ("left", "nil", "0"),
            ("right", "nil", "0"),
            ("insideH", "nil", "0"),
            ("insideV", "nil", "0"),
        ):
            node = borders.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                borders.append(node)
            node.set(qn("w:val"), value)
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), "000000")
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if row_index == 0:
                header = tr_pr.find(qn("w:tblHeader"))
                if header is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    tr_pr.append(header)
            else:
                # Short rows remain atomic; rows containing substantial text
                # are deliberately left splittable to avoid large white gaps.
                text_length = sum(len(p.text) for c in row.cells for p in c.paragraphs)
                if text_length < 160:
                    cant_split = tr_pr.find(qn("w:cantSplit"))
                    if cant_split is None:
                        tr_pr.append(OxmlElement("w:cantSplit"))
            grid_index = 0
            for cell_xml in row._tr.findall(qn("w:tc")):
                tc_pr = cell_xml.find(qn("w:tcPr"))
                if tc_pr is None:
                    tc_pr = OxmlElement("w:tcPr")
                    cell_xml.insert(0, tc_pr)
                grid_span = tc_pr.find(qn("w:gridSpan"))
                span = int(grid_span.get(qn("w:val"), "1")) if grid_span is not None else 1
                cell_width = sum(widths[grid_index : grid_index + span])
                grid_index += span
                tc_width = tc_pr.find(qn("w:tcW"))
                if tc_width is None:
                    tc_width = OxmlElement("w:tcW")
                    tc_pr.append(tc_width)
                tc_width.set(qn("w:type"), "dxa")
                tc_width.set(qn("w:w"), str(cell_width))

            seen_cells: set[int] = set()
            for column_index, cell in enumerate(row.cells):
                cell_key = id(cell._tc)
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                source_column = (
                    source_layout.columns[column_index]
                    if source_layout is not None
                    and column_index < len(source_layout.columns)
                    else None
                )
                vertical = source_column.vertical if source_column is not None else "center"
                cell.vertical_alignment = {
                    "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                    "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                    "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
                }.get(vertical, WD_CELL_VERTICAL_ALIGNMENT.CENTER)
                if row_index == 0:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    cell_borders = tc_pr.find(qn("w:tcBorders"))
                    if cell_borders is None:
                        cell_borders = OxmlElement("w:tcBorders")
                        tc_pr.append(cell_borders)
                    bottom = cell_borders.find(qn("w:bottom"))
                    if bottom is None:
                        bottom = OxmlElement("w:bottom")
                        cell_borders.append(bottom)
                    bottom.set(qn("w:val"), "single")
                    bottom.set(qn("w:sz"), "8")
                    bottom.set(qn("w:color"), "000000")
                for paragraph in cell.paragraphs:
                    source_alignment = paragraph.alignment
                    tc_pr = cell._tc.get_or_add_tcPr()
                    grid_span = tc_pr.find(qn("w:gridSpan"))
                    is_merged_cell = (
                        grid_span is not None
                        and int(grid_span.get(qn("w:val"), "1")) > 1
                    )
                    if source_column is not None and not is_merged_cell:
                        # For ordinary cells the LaTeX preamble is the source
                        # of truth.  Pandoc may leave direct paragraph
                        # alignment behind while converting the simplified
                        # l/c/r preamble; do not let that incidental OOXML
                        # override L/C/R/X or p/m/b column semantics.
                        horizontal = source_column.horizontal
                    elif source_alignment == WD_ALIGN_PARAGRAPH.CENTER:
                        horizontal = "center"
                    elif source_alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                        horizontal = "right"
                    elif source_alignment == WD_ALIGN_PARAGRAPH.LEFT:
                        horizontal = "left"
                    else:
                        horizontal = "left"
                    paragraph.style = document.styles[
                        _table_style_for_alignment(row_index, horizontal)
                    ]
                    _clear_direct_paragraph_format(
                        paragraph,
                        "ind",
                        "spacing",
                        "jc",
                    )
                    if (
                        source_column is not None
                        and source_column.first_line_indent_pt is not None
                        and source_column.first_line_indent_pt != 0
                    ):
                        paragraph.paragraph_format.first_line_indent = Pt(
                            source_column.first_line_indent_pt
                        )
                    for run in paragraph.runs:
                        _clear_direct_run_typography(run)

    if layouts and (
        semantic_table_count != len(layouts) or len(used_layout_ids) != len(layouts)
    ):
        print(
            "WARNING: LaTeX 表格布局数量与 Word 语义表格数量不一致，"
            f"LaTeX={len(layouts)}, Word={semantic_table_count}, "
            f"matched={len(used_layout_ids)}；未匹配表格已使用 Pandoc 回退。",
            file=sys.stderr,
        )


def _clamp_images(document: Document) -> None:
    maximum_width = int(15.5 * 360000)  # 15.5 cm usable width in EMU.
    for inline in document.element.iter(qn("wp:inline")):
        extent = inline.find(qn("wp:extent"))
        if extent is None:
            continue
        width = int(extent.get("cx", "0"))
        height = int(extent.get("cy", "0"))
        if width <= maximum_width or width <= 0:
            continue
        scale = maximum_width / width
        new_width = maximum_width
        new_height = max(1, int(height * scale))
        extent.set("cx", str(new_width))
        extent.set("cy", str(new_height))
        for drawing_extent in inline.iter(qn("a:ext")):
            drawing_extent.set("cx", str(new_width))
            drawing_extent.set("cy", str(new_height))


def _format_images_and_captions(
    document: Document,
    labels: LabelRegistry | None = None,
    discipline: str = "science",
    *,
    figure_sequence: str = "图",
    table_sequence: str = "表",
) -> None:
    """Apply Word-style-driven layout to embedded figures and captions."""

    for paragraph in document.paragraphs:
        has_drawing = bool(paragraph._p.xpath(".//w:drawing|.//w:pict"))
        if has_drawing and paragraph.style.name not in {"CoverLogo"}:
            paragraph.style = document.styles[STYLE_FIGURE]
            _clear_direct_paragraph_format(
                paragraph,
                "keepNext",
                "keepLines",
                "spacing",
                "ind",
                "jc",
            )

        style_name = paragraph.style.name
        text = paragraph.text.strip()
        if style_name not in {
            "Caption",
            "Image Caption",
            "Figure Caption",
            "Table Caption",
            STYLE_FIGURE_CAPTION,
            STYLE_TABLE_CAPTION,
        }:
            continue
        is_table_caption = text.startswith("表")
        # The reusable thesis styles inherit Word's built-in Caption style,
        # while keeping table-vs-figure pagination explicit and editable.
        paragraph.style = document.styles[
            STYLE_TABLE_CAPTION if is_table_caption else STYLE_FIGURE_CAPTION
        ]
        _clear_direct_paragraph_format(
            paragraph,
            "keepNext",
            "keepLines",
            "spacing",
            "ind",
            "jc",
        )
        for run in paragraph.runs:
            _clear_title_run_typography(run)

    _promote_caption_fields(
        document,
        labels,
        discipline,
        figure_sequence=figure_sequence,
        table_sequence=table_sequence,
    )

    for blip in document.element.iter(qn("a:blip")):
        if blip.get(qn("r:link")):
            raise BuildError("检测到 Word 图片外链；模板只允许内嵌图片关系。")


def _word_field_identifier(value: str) -> str:
    value = value.strip()
    if not value:
        raise BuildError("Word 题注序列名不能为空。")
    if re.fullmatch(r'[^\s"\\]+', value):
        return value
    return '"' + value.replace('"', '""') + '"'


def _field_run(
    instruction: str,
    result: str = "",
    *,
    locked: bool = False,
):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    field.set(qn("w:dirty"), "true")
    if locked:
        field.set(qn("w:fldLock"), "true")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = result
    run.append(text)
    field.append(run)
    return field


def _promote_caption_fields(
    document: Document,
    labels: LabelRegistry | None = None,
    discipline: str = "science",
    *,
    figure_sequence: str = "图",
    table_sequence: str = "表",
) -> None:
    """Turn Pandoc caption text into native Word caption sequences.

    The visible Chinese label and the registered Word CaptionLabel identifier
    intentionally match.  This keeps generated captions, Word's Insert Caption
    dialog and the Cross-reference dialog on the same ``图``/``表`` sequence.
    """

    bookmark_id = _next_bookmark_id(document)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        match = re.match(r"^(图|表)([0-9]+(?:\.[0-9]+)?)\s*(.*)$", text)
        if not match or paragraph.style.name not in {
            "Figure Caption",
            "Table Caption",
            "Caption",
            "Image Caption",
            STYLE_FIGURE_CAPTION,
            STYLE_TABLE_CAPTION,
        }:
            continue
        kind, number, rest = match.groups()
        # Rebuild paragraph while retaining its caption style and formatting.
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
        chapter_number = ""
        sequence_number = number
        sequence_name = figure_sequence if kind == "图" else table_sequence
        if discipline == "science" and "." in number:
            chapter, sequence_number = number.split(".", 1)
            chapter_number = chapter

        matched = next(
            (
                label
                for label, target in (labels.targets.items() if labels else [])
                if target.kind == kind and target.number == number
            ),
            None,
        )
        seed = matched or f"anonymous:{kind}:{number}:{bookmark_id}"
        full_name = (
            _label_bookmark_name(matched)
            if matched
            else _hidden_ref_bookmark_name(f"full:{seed}")
        )
        number_name = (
            _label_number_bookmark_name(matched)
            if matched
            else _hidden_ref_bookmark_name(f"number:{seed}", prefix="_RefNum")
        )
        full_id = bookmark_id
        number_id = bookmark_id + 1
        bookmark_id += 2

        full_start = OxmlElement("w:bookmarkStart")
        full_start.set(qn("w:id"), str(full_id))
        full_start.set(qn("w:name"), full_name)
        full_end = OxmlElement("w:bookmarkEnd")
        full_end.set(qn("w:id"), str(full_id))
        number_start = OxmlElement("w:bookmarkStart")
        number_start.set(qn("w:id"), str(number_id))
        number_start.set(qn("w:name"), number_name)
        number_end = OxmlElement("w:bookmarkEnd")
        number_end.set(qn("w:id"), str(number_id))

        paragraph._p.append(full_start)
        paragraph._p.append(_plain_run(kind))
        paragraph._p.append(number_start)
        if chapter_number:
            for node in _complex_field_runs(
                f' STYLEREF "{STYLE_HEADING_1}" \\n ',
                chapter_number,
            ):
                paragraph._p.append(node)
            paragraph._p.append(_plain_run("."))
        sequence_instruction = f" SEQ {_word_field_identifier(sequence_name)} "
        if sequence_number == "1":
            sequence_instruction += "\\r 1 "
        sequence_instruction += "\\* ARABIC "
        for node in _complex_field_runs(sequence_instruction, sequence_number):
            paragraph._p.append(node)
        paragraph._p.append(number_end)
        paragraph._p.append(full_end)
        if rest:
            paragraph._p.append(_plain_run(" " + rest))


def _label_bookmark_name(label: str) -> str:
    return _hidden_ref_bookmark_name(f"full:{label}")


def _label_number_bookmark_name(label: str) -> str:
    return _hidden_ref_bookmark_name(f"number:{label}", prefix="_RefNum")


def _hidden_ref_bookmark_name(seed: str, *, prefix: str = "_Ref") -> str:
    """Return a deterministic hidden bookmark that resembles Word's own.

    Leading underscores keep the target out of Word's normal bookmark list.
    A decimal digest is used because native cross-reference bookmarks are
    conventionally named ``_Ref#########``.
    """

    digest = int(hashlib.sha1(seed.encode("utf-8")).hexdigest()[:15], 16)
    return f"{prefix}{digest % 10**15:015d}"


def _next_bookmark_id(document: Document) -> int:
    values = [
        int(value)
        for node in document.element.iter(qn("w:bookmarkStart"))
        if (value := node.get(qn("w:id"))) is not None and value.isdigit()
    ]
    return max(values, default=0) + 1


def _citation_bookmark_name(key: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", key or "")[:17]
    digest = hashlib.sha1((key or "").encode("utf-8")).hexdigest()[:8]
    return f"YibinCitation_{safe}_{digest}"


def _run_style_properties(character_style_id: str | None):
    if not character_style_id:
        return None
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), character_style_id)
    r_pr.append(r_style)
    return r_pr


def _plain_run(text: str, *, character_style_id: str | None = None):
    run = OxmlElement("w:r")
    r_pr = _run_style_properties(character_style_id)
    if r_pr is not None:
        run.append(r_pr)
    node = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    return run


def _complex_field_runs(
    instruction: str,
    result: str,
    *,
    character_style_id: str | None = None,
    locked: bool = False,
) -> list:
    """Return a schema-valid complex field whose result keeps its style.

    Word recreates an unlocked REF result during F9.  CHARFORMAT reapplies the
    character style carried by the field-code runs, so the citation number
    remains a 9 pt superscript after refresh instead of dropping to baseline.
    """

    def styled_run(child):
        run = OxmlElement("w:r")
        r_pr = _run_style_properties(character_style_id)
        if r_pr is not None:
            run.append(r_pr)
        run.append(child)
        return run

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    if locked:
        begin.set(qn("w:fldLock"), "true")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = result
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    return [
        styled_run(begin),
        styled_run(code),
        styled_run(separate),
        styled_run(text),
        styled_run(end),
    ]


def _replace_reference_markers(document: Document, labels: LabelRegistry, citations: CitationRegistry, mode: str) -> None:
    """Replace opaque LaTeX reference/citation markers with Word fields."""
    label_names = {label: _label_bookmark_name(label) for label in labels.targets}
    label_number_names = {
        label: _label_number_bookmark_name(label) for label in labels.targets
    }
    citation_names = {key: _citation_bookmark_name(key) for key in citations.order}
    citation_numbers = {
        key: index for index, key in enumerate(citations.order, start=1)
    }
    generated_names = [
        *label_names.values(),
        *label_number_names.values(),
        *citation_names.values(),
    ]
    if any(len(name) > 40 for name in generated_names):
        raise BuildError("Word 书签名超过 40 字符限制。")
    folded_names = [name.casefold() for name in generated_names]
    if len(folded_names) != len(set(folded_names)):
        raise BuildError("Word 书签名发生大小写不敏感碰撞。")
    citation_style_id = document.styles[STYLE_CITATION].style_id
    for paragraph in document.paragraphs:
        for run in list(paragraph.runs):
            value = run.text or ""
            markers = list(re.finditer(r"YIBINXREF\d{8}|YIBINCITE\d{8}", value))
            if not markers:
                continue
            parent = run._r.getparent(); index = parent.index(run._r)
            parent.remove(run._r)
            cursor = 0
            for match in markers:
                token = match.group(0)
                if token.startswith("YIBINXREF"):
                    label, paren = labels.references.get(token, ("", False)); target = labels.targets.get(label)
                    if target is None: raise BuildError(f"未找到交叉引用标签：{label}")
                    prefix = value[cursor:match.start()]
                    use_full_target = bool(paren and target.kind == "公式")
                    if not paren and target.kind in {"图", "表"} and prefix.endswith(target.kind):
                        prefix = prefix[: -len(target.kind)]
                        use_full_target = True
                    if prefix:
                        parent.insert(index, _plain_run(prefix)); index += 1
                    bookmark = (
                        label_names[label]
                        if use_full_target
                        else label_number_names[label]
                    )
                    if use_full_target and target.kind in {"图", "表"}:
                        cached_result = f"{target.kind}{target.number}"
                    elif use_full_target and target.kind == "公式":
                        cached_result = f"({target.number})"
                    else:
                        cached_result = (
                            f"({target.number})" if paren else target.number
                        )
                    for field_run in _complex_field_runs(
                        f" REF {bookmark} \\h ",
                        cached_result,
                    ):
                        parent.insert(index, field_run)
                        index += 1
                else:
                    prefix = value[cursor:match.start()]
                    if prefix:
                        parent.insert(index, _plain_run(prefix)); index += 1
                    keys = sorted(
                        citations.clusters.get(token, []),
                        key=lambda key: citation_numbers[key],
                    )
                    open_run = _plain_run("[", character_style_id=citation_style_id)
                    parent.insert(index, open_run); index += 1
                    for key_index, key in enumerate(keys):
                        if key_index:
                            sep = copy.deepcopy(open_run); sep.find(qn("w:t")).text = ","; parent.insert(index, sep); index += 1
                        number = citation_numbers[key]
                        instruction = (
                            f" CITATION {key} \\l 2052 \\* CHARFORMAT "
                            if mode == "native"
                            else f" REF {citation_names[key]} \\h \\* CHARFORMAT "
                        )
                        for field_run in _complex_field_runs(
                            instruction,
                            str(number),
                            character_style_id=citation_style_id,
                            locked=mode == "native",
                        ):
                            parent.insert(index, field_run)
                            index += 1
                    close_run = copy.deepcopy(open_run); close_run.find(qn("w:t")).text = "]"; parent.insert(index, close_run); index += 1
                cursor = match.end()
            if cursor < len(value):
                parent.insert(index, _plain_run(value[cursor:]))
    # Bookmark bibliography entries by citation order where possible.
    for paragraph in document.paragraphs:
        if paragraph.style.name.casefold() not in {
            "bibliography",
            STYLE_BIBLIOGRAPHY.casefold(),
        }:
            continue
        m = re.match(r"^\s*\[?(\d+)\]?\s*", paragraph.text or "")
        if not m: continue
        number = int(m.group(1))
        if 1 <= number <= len(citations.order):
            key = citations.order[number - 1]; bid = 3000 + number
            # Isolate the displayed number so REF returns only the number,
            # never the complete bibliography entry.
            first = paragraph.runs[0] if paragraph.runs else None
            number_run = None
            first_match = (
                re.match(r"^(\s*)\[?\d+\]?", first.text)
                if first is not None
                else None
            )
            if first is not None and first_match is not None:
                leading = first_match.group(1)
                remainder = first.text[first_match.end():]
                parent = first._r.getparent()
                position = parent.index(first._r)
                r_pr = first._r.find(qn("w:rPr"))
                parent.remove(first._r)
                pieces = (leading + "[", str(number), "]" + remainder)
                created = []
                for piece in pieces:
                    node = OxmlElement("w:r")
                    if r_pr is not None:
                        node.append(copy.deepcopy(r_pr))
                    text_node = OxmlElement("w:t")
                    if piece.startswith(" ") or piece.endswith(" "):
                        text_node.set(qn("xml:space"), "preserve")
                    text_node.text = piece
                    node.append(text_node)
                    parent.insert(position, node)
                    position += 1
                    created.append(node)
                number_run = created[1]
            start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(bid)); start.set(qn("w:name"), citation_names[key])
            end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bid))
            if number_run is not None:
                pos = paragraph._p.index(number_run)
                paragraph._p.insert(pos, start)
                paragraph._p.insert(pos + 2, end)
            else:
                raise BuildError(f"无法隔离参考文献编号书签：{paragraph.text[:80]}")


def _word_source_type(csl_type: str) -> str:
    return {
        "article-journal": "JournalArticle",
        "article-magazine": "ArticleInAPeriodical",
        "article-newspaper": "ArticleInAPeriodical",
        "book": "Book",
        "chapter": "BookSection",
        "paper-conference": "ConferenceProceedings",
        "report": "Report",
        "thesis": "Report",
        "webpage": "DocumentFromInternetSite",
        "post-weblog": "InternetSite",
        "patent": "Patent",
    }.get(csl_type, "Misc")


def _append_word_source_text(source: ET.Element, name: str, value: object) -> None:
    text = str(value or "").strip()
    if not text:
        return
    ET.SubElement(source, f"{{{BIBLIOGRAPHY_NS}}}{name}").text = text


BIBLIOGRAPHY_NS = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
CUSTOM_XML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/customXml"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _bibliography_sources_xml(
    citations: CitationRegistry,
    bibliography_items: list[dict[str, object]],
) -> bytes:
    ET.register_namespace("b", BIBLIOGRAPHY_NS)
    by_key = {str(item.get("id", "")): item for item in bibliography_items}
    root = ET.Element(
        f"{{{BIBLIOGRAPHY_NS}}}Sources",
        {
            "SelectedStyle": "\\Yibin-GB-T-7714-Numeric.xsl",
            "StyleName": "GB/T 7714 数字引用（YibinThesis）",
            "Version": "6",
        },
    )
    for number, key in enumerate(citations.order, start=1):
        item = by_key.get(key, {})
        source = ET.SubElement(root, f"{{{BIBLIOGRAPHY_NS}}}Source")
        _append_word_source_text(source, "Tag", key)
        _append_word_source_text(source, "SourceType", _word_source_type(str(item.get("type", ""))))
        _append_word_source_text(
            source,
            "Guid",
            "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, f"yibinthesis:{key}")).upper() + "}",
        )
        _append_word_source_text(source, "LCID", "2052")
        _append_word_source_text(source, "RefOrder", number)

        authors = item.get("author")
        if isinstance(authors, list) and authors:
            author_node = ET.SubElement(source, f"{{{BIBLIOGRAPHY_NS}}}Author")
            author_role = ET.SubElement(author_node, f"{{{BIBLIOGRAPHY_NS}}}Author")
            people = [author for author in authors if isinstance(author, dict) and not author.get("literal")]
            corporate = [str(author.get("literal", "")).strip() for author in authors if isinstance(author, dict) and author.get("literal")]
            if people:
                name_list = ET.SubElement(author_role, f"{{{BIBLIOGRAPHY_NS}}}NameList")
                for author in people:
                    person = ET.SubElement(name_list, f"{{{BIBLIOGRAPHY_NS}}}Person")
                    _append_word_source_text(person, "Last", author.get("family"))
                    _append_word_source_text(person, "First", author.get("given"))
            elif corporate:
                _append_word_source_text(author_role, "Corporate", "；".join(corporate))

        title = item.get("title") or key
        _append_word_source_text(source, "Title", title)
        container = item.get("container-title")
        if isinstance(container, list):
            container = "; ".join(str(value) for value in container)
        _append_word_source_text(source, "JournalName", container)
        issued = item.get("issued")
        if isinstance(issued, dict):
            date_parts = issued.get("date-parts")
            if isinstance(date_parts, list) and date_parts and isinstance(date_parts[0], list):
                parts = date_parts[0]
                if parts:
                    _append_word_source_text(source, "Year", parts[0])
                if len(parts) > 1:
                    _append_word_source_text(source, "Month", parts[1])
                if len(parts) > 2:
                    _append_word_source_text(source, "Day", parts[2])
        for csl_name, word_name in (
            ("publisher", "Publisher"),
            ("publisher-place", "City"),
            ("volume", "Volume"),
            ("issue", "Issue"),
            ("page", "Pages"),
            ("DOI", "DOI"),
            ("URL", "URL"),
            ("number", "StandardNumber"),
        ):
            _append_word_source_text(source, word_name, item.get(csl_name))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _inject_bibliography_sources(
    docx_path: Path,
    citations: CitationRegistry,
    bibliography_items: list[dict[str, object]],
) -> None:
    """Add a fully related Word bibliography custom XML data store."""
    if not citations.order:
        return
    temp = docx_path.with_suffix(".sources.docx")
    with zipfile.ZipFile(docx_path, "r") as src:
        names = set(src.namelist())
        index = 1
        while f"customXml/item{index}.xml" in names:
            index += 1
        item_name = f"customXml/item{index}.xml"
        props_name = f"customXml/itemProps{index}.xml"
        item_rels_name = f"customXml/_rels/item{index}.xml.rels"

        document_rels = ET.fromstring(src.read("word/_rels/document.xml.rels"))
        relationship_ids = {
            relationship.get("Id", "") for relationship in document_rels
        }
        rel_index = 1
        while f"rIdYibinBibliography{rel_index}" in relationship_ids:
            rel_index += 1
        ET.SubElement(
            document_rels,
            f"{{{PACKAGE_REL_NS}}}Relationship",
            {
                "Id": f"rIdYibinBibliography{rel_index}",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml",
                "Target": f"../{item_name}",
            },
        )

        content_types = ET.fromstring(src.read("[Content_Types].xml"))
        ET.SubElement(
            content_types,
            f"{{{CONTENT_TYPES_NS}}}Override",
            {
                "PartName": "/" + props_name,
                "ContentType": "application/vnd.openxmlformats-officedocument.customXmlProperties+xml",
            },
        )

        store_id = "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, "yibinthesis:bibliography-store")).upper() + "}"
        ET.register_namespace("ds", CUSTOM_XML_NS)
        props = ET.Element(f"{{{CUSTOM_XML_NS}}}datastoreItem", {f"{{{CUSTOM_XML_NS}}}itemID": store_id})
        schema_refs = ET.SubElement(props, f"{{{CUSTOM_XML_NS}}}schemaRefs")
        ET.SubElement(schema_refs, f"{{{CUSTOM_XML_NS}}}schemaRef", {f"{{{CUSTOM_XML_NS}}}uri": BIBLIOGRAPHY_NS})

        ET.register_namespace("", PACKAGE_REL_NS)
        item_rels = ET.Element(f"{{{PACKAGE_REL_NS}}}Relationships")
        ET.SubElement(
            item_rels,
            f"{{{PACKAGE_REL_NS}}}Relationship",
            {
                "Id": "rId1",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps",
                "Target": f"itemProps{index}.xml",
            },
        )

        replacements = {
            "word/_rels/document.xml.rels": ET.tostring(document_rels, encoding="utf-8", xml_declaration=True),
            "[Content_Types].xml": ET.tostring(content_types, encoding="utf-8", xml_declaration=True),
        }
        with zipfile.ZipFile(temp, "w") as dst:
            for member in src.infolist():
                dst.writestr(member, replacements.get(member.filename, src.read(member)))
            dst.writestr(item_name, _bibliography_sources_xml(citations, bibliography_items))
            dst.writestr(props_name, ET.tostring(props, encoding="utf-8", xml_declaration=True))
            dst.writestr(item_rels_name, ET.tostring(item_rels, encoding="utf-8", xml_declaration=True))
    os.replace(temp, docx_path)


def postprocess_docx(
    input_docx: Path,
    output_docx: Path,
    *,
    metadata: dict[str, str],
    project_root: Path,
    main_dir: Path,
    metadata_dir: Path,
    discipline: str,
    allow_project_fallback: bool,
    include_cover: bool,
    include_declarations: bool,
    citation_mode: str = "linked",
    labels: LabelRegistry | None = None,
    citations: CitationRegistry | None = None,
    bibliography_items: list[dict[str, object]] | None = None,
    table_layouts: list[LatexTableLayout] | None = None,
    word_figure_sequence: str = "图",
    word_table_sequence: str = "表",
    word_equation_sequence: str = "公式",
) -> None:
    document = Document(input_docx)
    labels = labels or LabelRegistry()
    citations = citations or CitationRegistry()
    bibliography_items = bibliography_items or []
    table_layouts = table_layouts or []
    if "CoverValue" in [style.name for style in document.styles]:
        document.styles["CoverValue"].font.underline = False
    title = metadata.get("title", "")
    english_title = metadata.get("english-title", "")
    author = metadata.get("author", "")
    body_sect_pr = document._element.body.sectPr
    if body_sect_pr is None:
        raise BuildError("Pandoc 输出缺少正文分节属性。")

    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith("YIBIN_INTERNAL_CITATION_SEED"):
            paragraph._p.getparent().remove(paragraph._p)

    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        match = re.fullmatch(
            r"YIBIN_INTERNAL_EQUATION_(\d+)_(\d+)",
            paragraph.text.strip(),
        )
        if not match:
            continue
        equation_paragraph = next(
            (
                candidate
                for candidate in reversed(paragraphs[:index])
                if candidate._p.find(qn("m:oMathPara")) is not None
                or candidate._p.find(qn("m:oMath")) is not None
            ),
            None,
        )
        if equation_paragraph is None:
            raise BuildError("公式编号标记前未找到公式。")
        equation_paragraph.style = document.styles[STYLE_EQUATION]
        _format_equation_number(
            equation_paragraph,
            f"{match.group(1)}.{match.group(2)}",
            labels,
            discipline=discipline,
            equation_sequence=word_equation_sequence,
        )
        paragraph._p.getparent().remove(paragraph._p)

    cover_marker = None
    abstract_marker = None
    front_marker = None
    toc_paragraph = None
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value == SECTION_COVER_END:
            cover_marker = paragraph
        elif value == SECTION_ABSTRACT_END:
            abstract_marker = paragraph
        elif value == SECTION_FRONT_END:
            front_marker = paragraph
        elif value == TOC_MARKER:
            toc_paragraph = paragraph
    if (
        cover_marker is None
        or abstract_marker is None
        or front_marker is None
        or toc_paragraph is None
    ):
        missing = [
            name
            for value, name in (
                (cover_marker, "封面分节"),
                (abstract_marker, "摘要分节"),
                (front_marker, "前置分节"),
                (toc_paragraph, "目录"),
            )
            if value is None
        ]
        raise BuildError("Word 后处理标记缺失：" + "、".join(missing))

    _build_front_matter(
        document,
        cover_marker,
        metadata,
        project_root,
        main_dir,
        metadata_dir,
        allow_project_fallback=allow_project_fallback,
        include_cover=include_cover,
        include_declarations=include_declarations,
    )

    toc_paragraph.text = "目录"
    toc_paragraph.style = document.styles[STYLE_TOC_TITLE]
    _insert_toc_after(toc_paragraph)
    _add_section_break(cover_marker, body_sect_pr, None)
    _add_section_break(abstract_marker, body_sect_pr, "upperRoman")
    _add_section_break(front_marker, body_sect_pr, None)
    _set_page_number_format(document.sections[-1], "decimal", 1)

    stage = output_docx.with_suffix(".sections.docx")
    document.save(stage)
    document = Document(stage)
    if len(document.sections) != 4:
        stage.unlink(missing_ok=True)
        raise BuildError(f"预期生成 4 个 Word 分节，实际为 {len(document.sections)}。")

    for section in document.sections:
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)
        section.different_first_page_header_footer = False

    cover, abstracts, toc, body = document.sections
    for section in (cover, abstracts, toc, body):
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        _clear_story(section.header)
        _clear_story(section.footer)

    _set_page_number_format(cover, None)
    _set_page_number_format(abstracts, "upperRoman", 1)
    _set_page_number_format(toc, None)
    _set_page_number_format(body, "decimal", 1)

    front_footer = abstracts.footer.paragraphs[0]
    front_footer.style = document.styles["Footer"]
    _add_page_field(front_footer)
    for run in front_footer.runs:
        _set_run_fonts(run, "SimSun", "Times New Roman", 9)

    body_header = body.header.paragraphs[0]
    body_header.style = document.styles["Header"]
    body_header.text = title
    body_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_header_border(body_header)
    for run in body_header.runs:
        _set_run_fonts(run, "SimSun", "Times New Roman", 9)

    body_footer = body.footer.paragraphs[0]
    body_footer.style = document.styles["Footer"]
    _add_page_field(body_footer)
    for run in body_footer.runs:
        _set_run_fonts(run, "SimSun", "Times New Roman", 9)

    reusable_style_map = {
        "Normal": STYLE_BODY,
        "Body Text": STYLE_BODY,
        "Compact": STYLE_BODY,
        "First Paragraph": STYLE_FIRST_PARAGRAPH,
        "Heading 1": STYLE_HEADING_1,
        "Heading 2": STYLE_HEADING_2,
        "Heading 3": STYLE_HEADING_3,
        "Heading 4": STYLE_HEADING_4,
        "Yibin Heading 1": STYLE_HEADING_1,
        "Yibin Heading 2": STYLE_HEADING_2,
        "Yibin Heading 3": STYLE_HEADING_3,
        "Yibin Heading 4": STYLE_HEADING_4,
        "ChineseAbstract": STYLE_CHINESE_ABSTRACT,
        "EnglishAbstract": STYLE_ENGLISH_ABSTRACT,
        "Keywords": STYLE_KEYWORDS,
        "Bibliography": STYLE_BIBLIOGRAPHY,
        "Notes": STYLE_NOTES,
    }
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name
        target_style = reusable_style_map.get(style_name)
        if target_style:
            paragraph.style = document.styles[target_style]
        if style_name == "FrontTitle":
            paragraph.style = document.styles[
                STYLE_ENGLISH_ABSTRACT_TITLE
                if paragraph.text.strip() == "Abstract"
                else STYLE_FRONT_TITLE
            ]
        for run in paragraph.runs:
            if paragraph.style.name in reusable_style_map.values():
                _clear_direct_run_typography(run)

    heading_one = document.styles[STYLE_HEADING_1]
    if discipline == "science":
        heading_one.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_one.paragraph_format.first_line_indent = Pt(0)
        heading_one.paragraph_format.page_break_before = True
        heading_ppr = heading_one.element.get_or_add_pPr()
        heading_indent = heading_ppr.find(qn("w:ind"))
        if heading_indent is not None:
            heading_indent.set(qn("w:firstLineChars"), "0")
            heading_indent.set(qn("w:firstLine"), "0")
    else:
        heading_one.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading_one.paragraph_format.first_line_indent = Pt(24)
        heading_one.paragraph_format.page_break_before = False
        heading_ppr = heading_one.element.get_or_add_pPr()
        heading_indent = heading_ppr.find(qn("w:ind"))
        if heading_indent is None:
            heading_indent = OxmlElement("w:ind")
            heading_ppr.append(heading_indent)
        heading_indent.set(qn("w:firstLine"), "480")
        heading_indent.set(qn("w:firstLineChars"), "200")

    unnumbered_titles = {"结论", "注释", "参考文献", "附录", "致谢"}
    if discipline == "humanities":
        unnumbered_titles.add("绪论")
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name == STYLE_HEADING_1 and text in unnumbered_titles:
            paragraph.style = document.styles[STYLE_UNNUMBERED_HEADING]
        if re.match(r"^附录[A-ZＡ-Ｚ0-9一二三四五六七八九十]+(?:\s|　)", text):
            paragraph.style = document.styles[STYLE_APPENDIX_HEADING]

    _configure_heading_numbering(document, discipline)

    # All semantic section titles are style-driven.  Pandoc may leave direct
    # paragraph/run formatting on headings even after assigning a named style;
    # remove only those redundant overrides so a user's later style edit
    # changes every matching title consistently.
    title_styles = {
        STYLE_HEADING_1,
        STYLE_HEADING_2,
        STYLE_HEADING_3,
        STYLE_HEADING_4,
        STYLE_UNNUMBERED_HEADING,
        STYLE_FRONT_TITLE,
        STYLE_ENGLISH_ABSTRACT_TITLE,
        STYLE_TOC_TITLE,
        STYLE_APPENDIX_HEADING,
    }
    for paragraph in document.paragraphs:
        if paragraph.style.name not in title_styles:
            continue
        _clear_direct_paragraph_format(
            paragraph,
            "keepNext",
            "keepLines",
            "pageBreakBefore",
            "spacing",
            "ind",
            "jc",
        )
        for run in paragraph.runs:
            _clear_direct_run_typography(run)

    _format_tables(document, table_layouts, labels)
    _clamp_images(document)
    _format_images_and_captions(
        document,
        labels,
        discipline,
        figure_sequence=word_figure_sequence,
        table_sequence=word_table_sequence,
    )
    _replace_reference_markers(document, labels, citations, citation_mode)
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.core_properties.title = title
    document.core_properties.author = author
    subject = "宜宾学院本科毕业论文（非官方技术模板）"
    if english_title.strip():
        subject += " / " + english_title.strip()
    document.core_properties.subject = subject
    if citation_mode not in {"linked", "native"}:
        raise BuildError(f"不支持的 Word citation mode：{citation_mode}")
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    if citation_mode == "native":
        _inject_bibliography_sources(output_docx, citations, bibliography_items)
    final_normalized_pngs = _normalize_embedded_pngs(output_docx)
    if final_normalized_pngs:
        print(
            f"Word 图片兼容处理：最终文档中另有 {final_normalized_pngs} 个透明 PNG "
            "已转为白底 RGB PNG。"
        )
    stage.unlink(missing_ok=True)

    # Re-open once so corrupt OOXML fails in the build rather than on delivery.
    check = Document(output_docx)
    residual = {
        paragraph.text.strip()
        for paragraph in check.paragraphs
        if paragraph.text.strip().startswith("YIBIN_INTERNAL_")
    }
    if residual:
        raise BuildError("Word 后处理标记未清理：" + ", ".join(sorted(residual)))


def _resolve_argument_path(value: str | Path, project_root: Path) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        cwd_candidate = (Path.cwd() / path).resolve()
        project_candidate = (project_root / path).resolve()
        path = cwd_candidate if cwd_candidate.exists() else project_candidate
    return path.resolve()


def build(args: argparse.Namespace) -> Path:
    project_root = Path(__file__).resolve().parents[1]
    main_path = _resolve_argument_path(args.main, project_root)
    if not main_path.is_file():
        raise BuildError(f"主文件不存在：{main_path}")
    main_dir = main_path.parent
    allow_project_fallback = path_is_within(main_path, project_root)
    main_text = main_path.read_text(encoding="utf-8")
    events = parse_main_events(main_text)
    discipline = parse_discipline(main_text)

    metadata_path: Path | None = None
    if args.metadata:
        metadata_path = _resolve_argument_path(args.metadata, project_root)
    else:
        for event in events:
            if event.kind in {"input", "include"} and event.phase == "pre" and event.value:
                candidate = resolve_source(
                    event.value,
                    main_dir,
                    project_root,
                    allow_project_fallback=allow_project_fallback,
                )
                if "\\yibinsetup" in candidate.read_text(encoding="utf-8"):
                    metadata_path = candidate
                    break
        if metadata_path is None:
            candidates = [main_dir / "metadata.tex"]
            if allow_project_fallback:
                candidates.append(project_root / "metadata.tex")
            metadata_path = next((path.resolve() for path in candidates if path.is_file()), None)
    if metadata_path is None or not metadata_path.is_file():
        raise BuildError("未找到 metadata.tex；可用 --metadata 显式指定。")
    metadata = parse_metadata(metadata_path)

    pandoc = find_pandoc(args.pandoc, project_root)
    reference_doc = _resolve_argument_path(args.reference_doc, project_root)
    if not reference_doc.is_file():
        generator = project_root / "tools" / "build_reference_docx.py"
        run_checked(
            [sys.executable, str(generator), "--output", str(reference_doc)],
            cwd=project_root,
        )
    csl = _resolve_argument_path(args.csl, project_root)
    if not csl.is_file():
        raise BuildError(f"CSL 文件不存在：{csl}")

    bibliography: Path | None
    if args.bibliography:
        bibliography = _resolve_argument_path(args.bibliography, project_root)
    else:
        bib_match = re.search(
            r"\\addbibresource(?:\[[^]]*\])?\s*\{([^{}]+)\}",
            strip_tex_comments(main_text),
        )
        bib_name = bib_match.group(1).strip() if bib_match else "references.bib"
        bibliography = next(
            (
                candidate.resolve()
                for candidate in (
                    [main_dir / bib_name]
                    + ([project_root / bib_name] if allow_project_fallback else [])
                )
                if candidate.is_file()
            ),
            None,
        )
    if bibliography is None:
        raise BuildError("未找到 references.bib；可用 --bibliography 显式指定。")

    output = (
        _resolve_argument_path(args.output, project_root)
        if args.output
        else (main_dir / "build" / f"{main_path.stem}.docx").resolve()
    )
    if output.resolve() == reference_doc.resolve():
        raise BuildError("输出文件不能覆盖 word/reference.docx。")

    warnings: list[str] = []
    notes = NoteRegistry()
    counters = HeadingCounters()
    float_counters = FloatCounters()
    labels = LabelRegistry()
    citations = CitationRegistry()
    table_layouts: list[LatexTableLayout] = []
    markdown_parts: list[str] = []

    graphic_paths = collect_graphic_paths(
        main_text,
        main_dir=main_dir,
        project_root=project_root,
        allow_project_fallback=allow_project_fallback,
    )
    resource_candidates = [
        main_dir.resolve(),
        (main_dir / "assets").resolve(),
        (main_dir / "chapters").resolve(),
        *graphic_paths,
    ]
    if allow_project_fallback:
        resource_candidates.extend(
            [project_root.resolve(), (project_root / "assets").resolve()]
        )
    resource_candidates = list(dict.fromkeys(resource_candidates))

    def image_target_resolver(target: str) -> str:
        return resolve_word_image_target(target, resource_candidates)

    pre_events = [event for event in events if event.phase == "pre"]
    include_cover = any(event.kind == "makeyibincover" for event in pre_events)
    include_declarations = any(
        event.kind == "makeyibindeclarations" for event in pre_events
    )
    if not include_cover and not include_declarations:
        raise BuildError("主文件未调用 \\makeyibincover 或 \\makeyibindeclarations。")
    markdown_parts.append(styled_marker(SECTION_COVER_END))

    # Keep conversion intermediates beside the requested output.  Windows
    # installations commonly place TEMP on a small system drive; a thesis with
    # many raster figures can otherwise fail even when the workspace drive has
    # ample free space.
    temporary_root = output.parent / ".tmp"
    temporary_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix="yibinthesis-word-",
        dir=temporary_root,
    ) as temporary:
        temp_dir = Path(temporary)
        front_has_content = False
        abstract_section_closed = False
        front_events = [event for event in events if event.phase == "front"]
        for index, event in enumerate(front_events):
            if event.kind in {"frontmatter", "mainmatter", "backmatter"}:
                continue
            if event.kind in {"input", "include"} and event.value:
                source = resolve_source(
                    event.value,
                    main_dir,
                    project_root,
                    allow_project_fallback=allow_project_fallback,
                )
                if source.resolve() == metadata_path.resolve():
                    continue
                text = expand_nested_inputs(
                    source,
                    project_root,
                    allow_project_fallback=allow_project_fallback,
                )
                detect_unsupported(text, source, warnings)
                environment = (
                    "cnabstract"
                    if extract_environment(text, "cnabstract") is not None
                    else "enabstract"
                    if extract_environment(text, "enabstract") is not None
                    else None
                )
                if front_has_content:
                    markdown_parts.append(PAGE_BREAK)
                if environment:
                    markdown_parts.append(
                        abstract_markdown(
                            text,
                            environment=environment,
                            pandoc=pandoc,
                            cwd=main_dir,
                            temp_dir=temp_dir,
                            notes=notes,
                            discipline=discipline,
                            labels=labels,
                            citations=citations,
                            table_layouts=table_layouts,
                        )
                    )
                else:
                    chapter_md = latex_to_markdown(
                        normalize_latex(
                            text,
                            notes,
                            discipline,
                            labels,
                            citations,
                            table_layouts,
                        ),
                        pandoc=pandoc,
                        cwd=main_dir,
                        temp_dir=temp_dir,
                        name=f"front-{index}",
                    )
                    markdown_parts.append(chapter_md)
                front_has_content = True
            elif event.kind == "tableofcontents":
                markdown_parts.append(styled_marker(SECTION_ABSTRACT_END))
                markdown_parts.append(styled_marker(TOC_MARKER))
                abstract_section_closed = True
                front_has_content = True
        if not any(event.kind == "tableofcontents" for event in front_events):
            markdown_parts.append(styled_marker(SECTION_ABSTRACT_END))
            markdown_parts.append(styled_marker(TOC_MARKER))
            abstract_section_closed = True
        if not abstract_section_closed:
            markdown_parts.append(styled_marker(SECTION_ABSTRACT_END))
        markdown_parts.append(styled_marker(SECTION_FRONT_END))

        content_events = [event for event in events if event.phase in {"main", "back"}]
        bibliography_inserted = False
        for index, event in enumerate(content_events):
            if event.kind in {"mainmatter", "backmatter", "frontmatter"}:
                continue
            if event.kind in {"input", "include"} and event.value:
                source = resolve_source(
                    event.value,
                    main_dir,
                    project_root,
                    allow_project_fallback=allow_project_fallback,
                )
                text = expand_nested_inputs(
                    source,
                    project_root,
                    allow_project_fallback=allow_project_fallback,
                )
                detect_unsupported(text, source, warnings)
                chapter_before = counters.chapter
                converted = latex_to_markdown(
                    normalize_latex(
                        text,
                        notes,
                        discipline,
                        labels,
                        citations,
                        table_layouts,
                    ),
                    pandoc=pandoc,
                    cwd=main_dir,
                    temp_dir=temp_dir,
                    name=f"chapter-{index}",
                )
                converted = apply_heading_numbering(
                    converted,
                    discipline,
                    counters,
                    labels,
                )
                converted = apply_float_numbering(
                    converted,
                    discipline,
                    chapter_before,
                    float_counters,
                    labels,
                    image_target_resolver,
                )
                markdown_parts.append(converted)
            elif event.kind == "printyibinnotes":
                rendered_notes = notes_markdown(notes)
                if rendered_notes:
                    markdown_parts.append(rendered_notes)
            elif event.kind == "printyibinbibliography":
                seed = citation_seed_markdown(citations)
                if seed:
                    markdown_parts.append(seed)
                markdown_parts.append(bibliography_markdown())
                bibliography_inserted = True
        if not bibliography_inserted:
            seed = citation_seed_markdown(citations)
            if seed:
                markdown_parts.append(seed)
            markdown_parts.append(bibliography_markdown())

        assembled = "\n\n".join(part for part in markdown_parts if part.strip()) + "\n"
        assembled = labels.resolve(assembled)
        markdown_file = temp_dir / "assembled.md"
        markdown_file.write_text(assembled, encoding="utf-8", newline="\n")
        if args.keep_intermediate:
            kept = output.with_suffix(".pandoc.md")
            kept.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(markdown_file, kept)

        pandoc_docx = temp_dir / "pandoc.docx"
        resource_paths = os.pathsep.join(
            str(path) for path in resource_candidates
        )
        citeproc_bibliography = prepare_bibliography(
            bibliography,
            pandoc=pandoc,
            cwd=main_dir,
            temp_dir=temp_dir,
        )
        bibliography_items: list[dict[str, object]] = []
        if citeproc_bibliography.suffix.lower() == ".json":
            parsed_items = json.loads(citeproc_bibliography.read_text(encoding="utf-8"))
            if isinstance(parsed_items, list):
                bibliography_items = [item for item in parsed_items if isinstance(item, dict)]
        command = [
            str(pandoc),
            "--from=markdown+smart+raw_attribute+fenced_divs+bracketed_spans+citations+tex_math_dollars",
            "--to=docx",
            "--standalone",
            "--citeproc",
            f"--reference-doc={reference_doc}",
            f"--bibliography={citeproc_bibliography}",
            f"--csl={csl}",
            f"--resource-path={resource_paths}",
            "--metadata=lang:zh-CN",
            f"--output={pandoc_docx}",
            str(markdown_file),
        ]
        run_checked(command, cwd=main_dir)
        normalized_pngs = _normalize_embedded_pngs(pandoc_docx)
        if normalized_pngs:
            print(
                f"Word 图片兼容处理：已将 {normalized_pngs} 个透明 PNG 转为白底 RGB PNG。"
            )
        postprocess_docx(
            pandoc_docx,
            output,
            metadata=metadata,
            project_root=project_root,
            main_dir=main_dir,
            metadata_dir=metadata_path.parent,
            discipline=discipline,
            allow_project_fallback=allow_project_fallback,
            include_cover=include_cover,
            include_declarations=include_declarations,
            citation_mode=args.citation_mode,
            labels=labels,
            citations=citations,
            bibliography_items=bibliography_items,
            table_layouts=table_layouts,
            word_figure_sequence=args.word_figure_sequence,
            word_table_sequence=args.word_table_sequence,
            word_equation_sequence=args.word_equation_sequence,
        )

    print(f"Generated: {output}")
    print("Word 输出可编辑；目录和页码域将在 Microsoft Word 打开时自动更新。")
    print("说明：复杂 TikZ、任意自定义宏、复杂浮动体不保证与 PDF 同页或等版式转换。")
    for warning in dict.fromkeys(warnings):
        print(f"WARNING: {warning}", file=sys.stderr)
    return output


def make_parser() -> argparse.ArgumentParser:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--main",
        default=str(project_root / "main.tex"),
        help="LaTeX entry point; examples may provide their own main.tex.",
    )
    parser.add_argument("--metadata", help="Override metadata.tex path.")
    parser.add_argument("--bibliography", help="Override BibTeX database path.")
    parser.add_argument("--output", help="Output DOCX path.")
    parser.add_argument("--pandoc", help="Pandoc executable path.")
    parser.add_argument(
        "--reference-doc",
        default=str(project_root / "word" / "reference.docx"),
        help="Pandoc semantic reference DOCX.",
    )
    parser.add_argument(
        "--csl",
        default=str(project_root / "word" / "china-national-standard-gb-t-7714-2015-numeric.csl"),
        help="Citation Style Language file.",
    )
    parser.add_argument(
        "--keep-intermediate",
        action="store_true",
        help="Keep the assembled Pandoc Markdown beside the output DOCX.",
    )
    parser.add_argument(
        "--citation-mode",
        choices=("linked", "native"),
        default="linked",
        help=(
            "Word 引用模式；linked 保持 citeproc 编号并生成可跳转的参考文献书签，"
            "native 另写入 Word CITATION 域和 bibliography customXml。"
        ),
    )
    parser.add_argument(
        "--word-figure-sequence",
        default="图",
        help="Registered Word figure CaptionLabel name.",
    )
    parser.add_argument(
        "--word-table-sequence",
        default="表",
        help="Registered Word table CaptionLabel name.",
    )
    parser.add_argument(
        "--word-equation-sequence",
        default="公式",
        help="Registered Word equation CaptionLabel name.",
    )
    return parser


def main() -> int:
    try:
        build(make_parser().parse_args())
    except BuildError as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("ERROR: 构建已中止。", file=sys.stderr)
        return 130
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
